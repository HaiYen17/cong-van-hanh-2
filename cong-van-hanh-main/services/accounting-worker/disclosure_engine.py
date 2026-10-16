from __future__ import annotations

import math
from io import BytesIO
from copy import deepcopy
from datetime import date
from typing import Any, Callable

from report_engine import ReportEngineError, normalize_key


DISCLOSURE_REPORT_KEY = "B09_DN_TT99"
def _text(value: Any) -> str:
    return " ".join(str(value or "").strip().split())


def _number(value: Any) -> float:
    try:
        result = float(value or 0)
        return result if math.isfinite(result) else 0.0
    except (TypeError, ValueError):
        return 0.0


def _payload(record: dict[str, Any]) -> dict[str, Any]:
    value = record.get("payload")
    return value if isinstance(value, dict) else {}


def _numeric_columns(report: dict[str, Any]) -> list[tuple[int, dict[str, Any]]]:
    result: list[tuple[int, dict[str, Any]]] = []
    for index, column in enumerate(report.get("columns") or []):
        column_type = normalize_key(column.get("columnType"))
        if column_type in ("LABEL", "ROW_CODE", "NOTE_REFERENCE"):
            continue
        result.append((index, column))
    return result


def _row(report: dict[str, Any], row_code: Any) -> dict[str, Any] | None:
    expected = _text(row_code).upper()
    return next(
        (
            row
            for row in report.get("rows") or []
            if _text(row.get("rowCode")).upper() == expected
        ),
        None,
    )


def _evidence_with_path(
    evidence: dict[str, Any],
    *,
    note_code: str,
    report_name: str,
    row_code: str,
    row_label: str,
) -> dict[str, Any]:
    result = deepcopy(evidence or {})
    prefix = f"{note_code} → {report_name} → Dòng {row_code} - {row_label}"
    components = []
    for component in result.get("components") or []:
        item = deepcopy(component)
        child = _text(item.get("tracePath") or item.get("label"))
        item["tracePath"] = f"{prefix} → {child}" if child else prefix
        components.append(item)
    result["components"] = components
    return result


def build_disclosure_output(
    *,
    records: list[dict[str, Any]],
    period_from: date,
    period_to: date,
    selection: dict[str, Any],
    resolve_report: Callable[[str], dict[str, Any]],
    selected_item_keys: list[str] | None = None,
) -> dict[str, Any]:
    """Build B09 from the shared report-formula model and keep legacy links for controls."""

    templates = [
        record
        for record in records
        if record.get("recordType") == "DISCLOSURE_TEMPLATE"
        and normalize_key(_payload(record).get("templateKey") or record.get("entityKey"))
        == DISCLOSURE_REPORT_KEY
    ]
    if not templates:
        raise ReportEngineError(
            "Chưa có nền B09-DN và bản đồ nghiệp vụ trong phiên bản cấu hình đang chọn."
        )
    template = _payload(templates[0])
    sections = sorted(
        [
            record
            for record in records
            if record.get("recordType") == "DISCLOSURE_SECTION"
            and normalize_key(_payload(record).get("templateKey"))
            == DISCLOSURE_REPORT_KEY
        ],
        key=lambda record: (_number(_payload(record).get("startRow")), _number(record.get("sortOrder"))),
    )
    items = sorted(
        [
            record
            for record in records
            if record.get("recordType") == "DISCLOSURE_ITEM"
            and normalize_key(_payload(record).get("templateKey"))
            == DISCLOSURE_REPORT_KEY
        ],
        key=lambda record: (_number(_payload(record).get("startRow")), _number(record.get("sortOrder"))),
    )
    requested = {normalize_key(value) for value in (selected_item_keys or []) if _text(value)}
    if requested:
        items = [
            record
            for record in items
            if normalize_key(_payload(record).get("itemKey") or record.get("entityKey"))
            in requested
        ]
        missing = requested - {
            normalize_key(_payload(record).get("itemKey") or record.get("entityKey"))
            for record in items
        }
        if missing:
            raise ReportEngineError(
                "Không tìm thấy khoản mục thuyết minh: " + ", ".join(sorted(missing))
            )
    links_by_item: dict[str, list[dict[str, Any]]] = {}
    for record in records:
        if record.get("recordType") != "DISCLOSURE_LINK":
            continue
        payload = _payload(record)
        item_key = normalize_key(payload.get("disclosureItemKey"))
        if item_key:
            links_by_item.setdefault(item_key, []).append(record)

    report_cache: dict[str, dict[str, Any]] = {}
    warnings: list[str] = []
    warning_seen: set[str] = set()

    def warning(message: str) -> None:
        cleaned = _text(message)
        if cleaned and cleaned not in warning_seen:
            warning_seen.add(cleaned)
            warnings.append(cleaned)

    def report_output(report_key: str) -> dict[str, Any]:
        normalized = normalize_key(report_key)
        if normalized not in report_cache:
            report_cache[normalized] = resolve_report(normalized)
        return report_cache[normalized]

    has_formula_report = any(
        record.get("recordType") == "REPORT"
        and normalize_key(record.get("entityKey") or _payload(record).get("reportKey"))
        == DISCLOSURE_REPORT_KEY
        for record in records
    )
    formula_report = report_output(DISCLOSURE_REPORT_KEY) if has_formula_report else None
    formula_rows = {
        _text(row.get("rowCode")).upper(): row
        for row in ((formula_report or {}).get("rows") or [])
        if _text(row.get("rowCode"))
    }
    formula_columns = _numeric_columns(formula_report or {})

    output_items: list[dict[str, Any]] = []
    ready_count = 0
    review_count = 0
    manual_count = 0
    failed_control_count = 0

    for item_record in items:
        item = _payload(item_record)
        item_key = normalize_key(item.get("itemKey") or item_record.get("entityKey"))
        note_code = _text(item.get("noteCode"))
        automation_level = normalize_key(item.get("automationLevel") or "MANUAL")
        item_links = sorted(
            links_by_item.get(item_key, []),
            key=lambda record: (
                normalize_key(_payload(record).get("reportKey")),
                _text(_payload(record).get("reportRowCode")),
            ),
        )
        link_outputs: list[dict[str, Any]] = []
        aggregate: dict[str, dict[str, Any]] = {}
        reference_groups: dict[str, list[dict[str, Any]]] = {}
        item_controls: list[dict[str, Any]] = []
        missing_link = False
        source_report_needs_review = False

        for link_record in item_links:
            link = _payload(link_record)
            report_key = normalize_key(link.get("reportKey"))
            row_code = _text(link.get("reportRowCode"))
            link_role = normalize_key(link.get("linkRole") or "RECONCILIATION_ONLY")
            explicit_aggregation_role = normalize_key(link.get("aggregationRole"))
            aggregation_role = explicit_aggregation_role or (
                "ADDITIVE" if link_role == "PRIMARY_NOTE_REFERENCE" else "REFERENCE_ONLY"
            )
            report = report_output(report_key)
            report_meta = report.get("report") or {}
            report_name = _text(report_meta.get("reportName") or report_key)
            report_row = _row(report, row_code)
            if not report_row:
                missing_link = True
                warning(f"{note_code}: không tìm thấy dòng {report_key}:{row_code} khi sinh số.")
                link_outputs.append(
                    {
                        "reportKey": report_key,
                        "reportName": report_name,
                        "rowCode": row_code,
                        "rowLabel": _text(link.get("reportRowLabel")),
                        "linkRole": link_role,
                        "aggregationRole": aggregation_role,
                        "status": "MISSING_REPORT_ROW",
                        "values": [],
                    }
                )
                continue
            source_report_needs_review = source_report_needs_review or report.get("controlStatus") != "READY"
            row_label = _text(report_row.get("rowLabel") or link.get("reportRowLabel"))
            # Tương thích các bản đồ Phase 2 đã lưu trước khi có
            # aggregationRole. Dòng "Trong đó" là tiểu mục thuyết minh,
            # được hiển thị nhưng không cộng thêm vào dòng cha.
            if (
                not explicit_aggregation_role
                and link_role == "PRIMARY_NOTE_REFERENCE"
                and normalize_key(row_label).startswith("TRONG_DO")
            ):
                aggregation_role = "SUBSET_DETAIL"
            values: list[dict[str, Any]] = []
            for column_index, column in _numeric_columns(report):
                cells = report_row.get("cells") or []
                if column_index >= len(cells):
                    continue
                cell = cells[column_index] or {}
                value = _number(cell.get("value"))
                column_key = normalize_key(column.get("columnKey") or column.get("columnLabel"))
                column_type = normalize_key(column.get("columnType"))
                # Khi cộng số phải giữ nguyên từng cột nguồn. Hai cột Nợ/Có
                # có thể cùng columnType nhưng tuyệt đối không được gộp vào nhau.
                aggregate_key = column_key or column_type
                # Khi đối chiếu hai báo cáo khác nhau thì columnKey có thể
                # khác; columnType là nghĩa chung để ghép đúng cột cùng kỳ.
                reference_key = column_type or column_key
                evidence = _evidence_with_path(
                    cell.get("evidence") or {},
                    note_code=note_code,
                    report_name=report_name,
                    row_code=row_code,
                    row_label=row_label,
                )
                value_output = {
                    "columnKey": column_key,
                    "columnLabel": _text(column.get("columnLabel")),
                    "columnType": column_type,
                    "value": value,
                    "evidence": evidence,
                }
                values.append(value_output)
                if aggregation_role == "ADDITIVE":
                    total = aggregate.setdefault(
                        aggregate_key,
                        {
                            "columnKey": column_key,
                            "columnLabel": _text(column.get("columnLabel")),
                            "columnType": column_type,
                            "value": 0.0,
                            "components": [],
                            "evidenceComplete": True,
                            "sourceRows": [],
                        },
                    )
                    total["value"] += value
                    total["components"].extend(evidence.get("components") or [])
                    total["evidenceComplete"] = bool(total["evidenceComplete"]) and bool(
                        evidence.get("isComplete", False)
                    )
                    total["sourceRows"].append(f"{report_key}:{row_code}")
                elif aggregation_role == "REFERENCE_ONLY":
                    reference_groups.setdefault(reference_key, []).append(
                        {
                            "reportKey": report_key,
                            "reportName": report_name,
                            "rowCode": row_code,
                            "rowLabel": row_label,
                            "value": value,
                            "columnKey": column_key,
                            "columnLabel": _text(column.get("columnLabel")),
                            "columnType": column_type,
                            "evidence": evidence,
                        }
                    )
            link_outputs.append(
                {
                    "reportKey": report_key,
                    "reportName": report_name,
                    "rowCode": row_code,
                    "rowLabel": row_label,
                    "linkRole": link_role,
                    "aggregationRole": aggregation_role,
                    "status": "RESOLVED",
                    "values": values,
                }
            )

        totals: list[dict[str, Any]] = []
        for semantic_key, total in aggregate.items():
            component_total = sum(_number(component.get("signedValue")) for component in total["components"])
            difference = _number(total["value"]) - component_total
            passed = bool(total["evidenceComplete"]) and abs(difference) < 0.5
            item_controls.append(
                {
                    "checkKey": f"{item_key}:{semantic_key}:EVIDENCE_TIE",
                    "label": "Chi tiết chứng minh cộng khớp số trên báo cáo chính",
                    "columnKey": total["columnKey"],
                    "columnLabel": total["columnLabel"],
                    "status": "PASS" if passed else "FAIL",
                    "reportValue": total["value"],
                    "evidenceValue": component_total,
                    "difference": difference,
                }
            )
            totals.append(
                {
                    "columnKey": total["columnKey"],
                    "columnLabel": total["columnLabel"],
                    "columnType": total["columnType"],
                    "value": total["value"],
                    "sourceRows": total["sourceRows"],
                    "evidence": {
                        "value": total["value"],
                        "components": total["components"],
                        "isComplete": passed,
                        "difference": difference,
                        "warning": "" if passed else "Chi tiết nguồn chưa cộng khớp; hệ thống không tự tạo số bù.",
                    },
                }
            )

        for semantic_key, references in reference_groups.items():
            if len(references) < 2:
                continue
            values = [_number(reference.get("value")) for reference in references]
            difference = max(values) - min(values)
            passed = abs(difference) < 0.5 and all(
                bool((reference.get("evidence") or {}).get("isComplete", False))
                for reference in references
            )
            item_controls.append(
                {
                    "checkKey": f"{item_key}:{semantic_key}:REFERENCE_COMPARE",
                    "label": "Các báo cáo tham chiếu cho cùng khoản mục phải bằng nhau",
                    "columnKey": references[0].get("columnKey"),
                    "columnLabel": references[0].get("columnLabel"),
                    "status": "PASS" if passed else "FAIL",
                    "reportValue": values[0],
                    "evidenceValue": values[-1],
                    "difference": difference,
                    "references": [
                        {
                            "reportKey": reference.get("reportKey"),
                            "rowCode": reference.get("rowCode"),
                            "value": reference.get("value"),
                        }
                        for reference in references
                    ],
                }
            )

        configured_row = formula_rows.get(note_code.upper())
        configured_formula = False
        missing_formula = False
        if configured_row:
            configured_totals: list[dict[str, Any]] = []
            configured_cells = configured_row.get("cells") or []
            for column_index, column in formula_columns:
                if column_index >= len(configured_cells):
                    continue
                cell = configured_cells[column_index] or {}
                if normalize_key(cell.get("kind")) != "VALUE":
                    continue
                if normalize_key(cell.get("formulaStatus")) == "MISSING":
                    missing_formula = True
                    continue
                configured_formula = True
                evidence = deepcopy(cell.get("evidence") or {})
                source_rows = list(
                    dict.fromkeys(
                        _text(component.get("tracePath") or component.get("label"))
                        for component in evidence.get("components") or []
                        if _text(component.get("tracePath") or component.get("label"))
                    )
                )
                configured_totals.append(
                    {
                        "columnKey": normalize_key(column.get("columnKey")),
                        "columnLabel": _text(column.get("columnLabel")),
                        "columnType": normalize_key(column.get("columnType")),
                        "value": _number(cell.get("value")),
                        "sourceRows": source_rows,
                        "evidence": evidence,
                    }
                )
            # Công thức trong trình dựng chung là nguồn tính chính. Các
            # DISCLOSURE_LINK cũ chỉ còn làm đường gợi ý/đối chiếu, không được
            # cộng thêm lần nữa.
            if configured_formula or missing_formula:
                totals = configured_totals
            if missing_formula:
                warning(f"{note_code}: còn ô B09 chưa thiết lập công thức.")

        failed_controls = [control for control in item_controls if control.get("status") == "FAIL"]
        failed_control_count += len(failed_controls)
        if missing_formula:
            status = "NEEDS_REVIEW"
            review_count += 1
        elif configured_formula:
            status = (
                "READY"
                if (formula_report or {}).get("controlStatus") == "READY"
                and not failed_controls
                else "NEEDS_REVIEW"
            )
            if status == "READY":
                ready_count += 1
            else:
                review_count += 1
        elif automation_level == "MANUAL":
            status = "MANUAL_INPUT"
            manual_count += 1
        elif (
            automation_level == "AUTOMATIC"
            and not missing_link
            and not failed_controls
            and not source_report_needs_review
            and bool(item_links)
        ):
            status = "READY"
            ready_count += 1
        else:
            status = "NEEDS_REVIEW"
            review_count += 1
        if failed_controls:
            warning(f"{note_code}: có {len(failed_controls)} kiểm tra cộng khớp chưa đạt.")
        if source_report_needs_review:
            warning(f"{note_code}: báo cáo nguồn còn kiểm tra kế toán chưa đạt.")

        output_items.append(
            {
                "itemKey": item_key,
                "noteCode": note_code,
                "noteTitle": _text(item.get("noteTitle")),
                "sectionKey": normalize_key(item.get("sectionKey")),
                "automationLevel": automation_level,
                "status": status,
                "sourceStrategy": normalize_key(item.get("sourceStrategy")),
                "sourceDatasets": item.get("sourceDatasets") or [],
                "periodBasis": item.get("periodBasis") or [],
                "dimensions": item.get("dimensions") or [],
                "formulaPrinciple": _text(item.get("formulaPrinciple")),
                "reviewRequirements": item.get("reviewRequirements") or [],
                "totals": totals,
                "links": link_outputs,
                "controlChecks": item_controls,
            }
        )

    section_output = []
    item_keys = {item["itemKey"] for item in output_items}
    for section_record in sections:
        section = _payload(section_record)
        section_key = normalize_key(section.get("sectionKey"))
        section_items = [
            item["itemKey"] for item in output_items if item.get("sectionKey") == section_key
        ]
        if requested and not section_items:
            continue
        section_output.append(
            {
                "sectionKey": section_key,
                "sectionCode": _text(section.get("sectionCode")),
                "sectionTitle": _text(section.get("sectionTitle")),
                "itemKeys": [key for key in section_items if key in item_keys],
            }
        )

    control_status = "READY" if output_items and all(item["status"] == "READY" for item in output_items) else "NEEDS_REVIEW"
    return {
        "ok": True,
        "schemaVersion": "accounting-disclosure-output.v1",
        "report": {
            "reportKey": DISCLOSURE_REPORT_KEY,
            "reportName": _text(template.get("templateName") or "B09-DN · Bản thuyết minh Báo cáo tài chính"),
            "displayTitle": _text(template.get("displayTitle") or template.get("templateName") or "BẢN THUYẾT MINH BÁO CÁO TÀI CHÍNH"),
        },
        "periodFrom": period_from.isoformat(),
        "periodTo": period_to.isoformat(),
        "selection": selection,
        "sections": section_output,
        "items": output_items,
        "summary": {
            "itemCount": len(output_items),
            "readyCount": ready_count,
            "reviewCount": review_count,
            "manualCount": manual_count,
            "failedControlCount": failed_control_count,
            "sourceReportCount": len(
                [key for key in report_cache if key != DISCLOSURE_REPORT_KEY]
            ),
        },
        "controlStatus": control_status,
        "warnings": warnings[:100],
    }


def disclosure_output_to_xlsx(output: dict[str, Any]) -> bytes:
    """Xuất B09 dạng bảng làm việc dễ đọc, giữ đủ khoản mục kể cả chưa có số."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "B09-DN"
    report = output.get("report") if isinstance(output.get("report"), dict) else {}
    title = _text(report.get("displayTitle") or report.get("reportName")) or "BẢN THUYẾT MINH BÁO CÁO TÀI CHÍNH"
    period = f"Kỳ {output.get('periodFrom') or ''} đến {output.get('periodTo') or ''}"
    headers = ["Mã", "Khoản mục thuyết minh", "Chỉ tiêu số liệu", "Giá trị", "Trạng thái", "Nguồn đối chiếu", "Kết quả kiểm tra", "Nguyên tắc lấy số"]
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    sheet.cell(1, 1, title)
    sheet.cell(1, 1).font = Font(name="Arial", size=16, bold=True, color="114C41")
    sheet.cell(1, 1).alignment = Alignment(horizontal="center", vertical="center")
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))
    sheet.cell(2, 1, period)
    sheet.cell(2, 1).font = Font(name="Arial", size=10, italic=True, color="5D6F6A")
    sheet.cell(2, 1).alignment = Alignment(horizontal="center")
    header_row = 4
    for column_index, header in enumerate(headers, start=1):
        cell = sheet.cell(header_row, column_index, header)
        cell.font = Font(name="Arial", size=10, bold=True, color="114C41")
        cell.fill = PatternFill("solid", fgColor="E8F3EE")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    thin = Side(style="thin", color="C9D7D3")
    section_fill = PatternFill("solid", fgColor="DDECE6")
    item_map = {
        normalize_key(item.get("itemKey")): item
        for item in output.get("items") or []
        if isinstance(item, dict)
    }
    status_labels = {"READY": "Sẵn sàng", "NEEDS_REVIEW": "Cần rà soát", "MANUAL_INPUT": "Kế toán nhập"}
    row_index = header_row + 1
    item_count = 0
    for section in output.get("sections") or []:
        if not isinstance(section, dict):
            continue
        section_label = " · ".join(filter(None, [_text(section.get("sectionCode")), _text(section.get("sectionTitle"))])) or "Phần thuyết minh"
        sheet.merge_cells(start_row=row_index, start_column=1, end_row=row_index, end_column=len(headers))
        section_cell = sheet.cell(row_index, 1, section_label)
        section_cell.font = Font(name="Arial", size=10, bold=True, color="114C41")
        section_cell.fill = section_fill
        section_cell.alignment = Alignment(vertical="center")
        row_index += 1
        for item_key in section.get("itemKeys") or []:
            item = item_map.get(normalize_key(item_key))
            if not item:
                continue
            item_count += 1
            totals = [total for total in item.get("totals") or [] if isinstance(total, dict)] or [None]
            controls = [control for control in item.get("controlChecks") or [] if isinstance(control, dict)]
            control_text = " · ".join(
                f"{'Đạt' if control.get('status') == 'PASS' else 'Chưa khớp'}: {_text(control.get('label'))}"
                + (f" ({_text(control.get('columnLabel'))})" if control.get("columnLabel") else "")
                for control in controls
            ) or "Chưa có phép kiểm tra tự động"
            for total_index, total in enumerate(totals):
                source_rows = [] if total is None else list(dict.fromkeys(_text(value) for value in total.get("sourceRows") or [] if _text(value)))
                if not source_rows:
                    for link in item.get("links") or []:
                        if not isinstance(link, dict):
                            continue
                        label = _text(link.get("reportName") or link.get("reportKey"))
                        if link.get("rowCode"):
                            label += f" · dòng {_text(link.get('rowCode'))}"
                        if label and label not in source_rows:
                            source_rows.append(label)
                values = [
                    _text(item.get("noteCode")) if total_index == 0 else "",
                    _text(item.get("noteTitle")) if total_index == 0 else "",
                    _text(total.get("columnLabel")) if total else "Chưa có số tự động",
                    _number(total.get("value")) if total else None,
                    status_labels.get(normalize_key(item.get("status")), _text(item.get("status"))) if total_index == 0 else "",
                    " · ".join(source_rows) or "Chưa có dòng báo cáo tự động",
                    control_text if total_index == 0 else "",
                    _text(item.get("formulaPrinciple")) if total_index == 0 else "",
                ]
                for column_index, value in enumerate(values, start=1):
                    cell = sheet.cell(row_index, column_index, value)
                    cell.font = Font(name="Arial", size=9, bold=column_index in (1, 2) and total_index == 0)
                    cell.alignment = Alignment(horizontal="right" if column_index == 4 else "left", vertical="top", wrap_text=True)
                    cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
                    if column_index == 4 and value is not None:
                        cell.number_format = '#,##0;[Red](#,##0);-'
                row_index += 1

    widths = [10, 34, 22, 18, 16, 38, 40, 46]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A5"
    sheet.row_dimensions[1].height = 26
    sheet.row_dimensions[4].height = 34
    sheet.sheet_view.showGridLines = False
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_title_rows = "1:4"
    sheet.auto_filter.ref = f"A4:H{max(4, row_index - 1)}"
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()
