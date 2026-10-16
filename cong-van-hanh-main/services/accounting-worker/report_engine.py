from __future__ import annotations

import io
import math
import re
import unicodedata
from copy import deepcopy
from datetime import date
from functools import lru_cache
from typing import Any, Callable, Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


DATA_CHECK_KEY = "DATA_CHECK"
AMOUNT_FIELDS = {
    "OPENING_DEBIT": "opening_debit",
    "OPENING_CREDIT": "opening_credit",
    "PERIOD_DEBIT": "period_debit",
    "PERIOD_CREDIT": "period_credit",
    "YTD_DEBIT": "ytd_debit",
    "YTD_CREDIT": "ytd_credit",
    "CLOSING_DEBIT": "closing_debit",
    "CLOSING_CREDIT": "closing_credit",
}


class ReportEngineError(ValueError):
    pass


_WHITESPACE_RE = re.compile(r"\s+")


def _text(value: Any) -> str:
    text = str(value or "").strip()
    return _WHITESPACE_RE.sub(" ", text) if _WHITESPACE_RE.search(text) else text


def normalize_key(value: Any) -> str:
    plain = "".join(
        character
        for character in unicodedata.normalize("NFD", _text(value).replace("đ", "d").replace("Đ", "D"))
        if unicodedata.category(character) != "Mn"
    )
    return re.sub(r"^_+|_+$", "", re.sub(r"[^A-Z0-9._:-]+", "_", plain.upper()))


_normalize_key_uncached = normalize_key


@lru_cache(maxsize=16_384)
def _normalize_key_cached(value: str) -> str:
    return _normalize_key_uncached(value)


def normalize_key(value: Any) -> str:
    return _normalize_key_cached(_text(value))


def _number(value: Any) -> float:
    try:
        number = float(value or 0)
        return number if math.isfinite(number) else 0.0
    except (TypeError, ValueError):
        return 0.0


def _safe_color(value: Any, fallback: str) -> str:
    normalized = str(value or "").strip().lstrip("#")
    return normalized.upper() if re.fullmatch(r"[0-9A-Fa-f]{6}", normalized) else fallback


def _field(row: dict[str, Any], snake_name: str) -> Any:
    parts = snake_name.split("_")
    camel_name = parts[0] + "".join(part.title() for part in parts[1:])
    value = row.get(snake_name)
    return row.get(camel_name) if value in (None, "") else value


def shift_year(value: date, offset: int) -> date:
    target_year = value.year + max(-20, min(20, int(offset)))
    month_after = date(target_year + (value.month == 12), 1 if value.month == 12 else value.month + 1, 1)
    last_day = (month_after - date.resolution).day
    return date(target_year, value.month, min(value.day, last_day))


def column_period_offset(payload: dict[str, Any]) -> int:
    period_selector = normalize_key(payload.get("periodSelector"))
    value_basis = normalize_key(payload.get("valueBasis"))
    if period_selector in ("CURRENT_RANGE", "OPENING_OF_CURRENT_RANGE", "END_OF_PRIOR_YEAR"):
        return 0
    if period_selector == "SAME_RANGE_PRIOR_YEAR":
        return -1
    # Opening balances come from the current snapshot. Shifting the complete
    # range one year back was a V2 shortcut and makes January openings empty.
    if value_basis in ("OPENING", "OPENING_BALANCE", "OPENING_OF_CURRENT_RANGE"):
        return 0
    configured = payload.get("periodOffset")
    if configured not in (None, ""):
        try:
            return max(-20, min(20, int(float(configured))))
        except (TypeError, ValueError):
            pass
    return -1 if normalize_key(payload.get("columnType")) == "PREVIOUS_PERIOD" else 0


def required_offsets(records: Iterable[dict[str, Any]], report_key: str) -> list[int]:
    normalized_report = normalize_key(report_key or DATA_CHECK_KEY)
    if normalized_report == DATA_CHECK_KEY:
        return [0]
    offsets = [0]
    columns = [
        record
        for record in records
        if record.get("recordType") == "REPORT_COLUMN"
        and normalize_key((record.get("payload") or {}).get("reportKey")) == normalized_report
    ]
    for record in columns:
        offset = column_period_offset(record.get("payload") or {})
        if offset not in offsets:
            offsets.append(offset)
    return offsets


def _selected_note_ids(
    records: Iterable[dict[str, Any]],
    report_key: str,
    selected_note_item_ids: Iterable[Any],
) -> tuple[list[str], list[str]]:
    normalized_report = normalize_key(report_key)
    required: list[str] = []
    selected: list[str] = []
    for record in records:
        payload = record.get("payload") or {}
        item_id = _text(payload.get("noteItemId")).lower()
        if (
            record.get("recordType") == "REPORT_ROW"
            and normalize_key(payload.get("reportKey")) == normalized_report
            and item_id
            and payload.get("noteRequired") is True
            and item_id not in required
        ):
            required.append(item_id)
    for value in selected_note_item_ids:
        item_id = _text(value).lower()
        if item_id and item_id not in selected:
            selected.append(item_id)
    return required, required + [item_id for item_id in selected if item_id not in required]


def note_required_offsets(
    records: list[dict[str, Any]],
    report_key: str,
    note_items: list[dict[str, Any]],
    selected_note_item_ids: Iterable[Any],
) -> list[int]:
    _, selected = _selected_note_ids(records, report_key, selected_note_item_ids)
    selected_set = set(selected)
    offsets = {0}
    for item in note_items:
        if _text(item.get("itemId")).lower() not in selected_set:
            continue
        formula = item.get("formula") if isinstance(item.get("formula"), list) else []
        if any(
            "NAMTRUOC" in normalize_key(term.get("operand"))
            for term in formula
            if isinstance(term, dict)
        ):
            offsets.add(-1)
    return sorted(offsets)


def _metric_value(
    row: dict[str, Any],
    metric_key: Any,
    rule: dict[str, Any],
    period_basis_override: str = "",
) -> float:
    normalized_metric = normalize_key(metric_key)
    if normalized_metric in AMOUNT_FIELDS and not period_basis_override:
        return _number(_field(row, AMOUNT_FIELDS[normalized_metric]))
    operator = normalize_key(rule.get("operator"))
    default_basis = (
        "DELTA"
        if operator == "DELTA_BALANCE"
        else "PERIOD"
        if operator in (
            "MOVEMENT",
            "MOVEMENT_DETAIL",
            "CASH_MOVEMENT",
            "PROFIT_LOSS_ACTIVITY",
        )
        else "ENDING"
    )
    period_basis = normalize_key(period_basis_override or rule.get("periodBasis") or default_basis)
    side = normalize_key(rule.get("side") or "NET")
    if period_basis == "OPENING":
        debit = _number(_field(row, "opening_debit"))
        credit = _number(_field(row, "opening_credit"))
    elif period_basis == "PERIOD":
        debit = _number(_field(row, "period_debit"))
        credit = _number(_field(row, "period_credit"))
    elif period_basis == "YTD":
        debit = _number(_field(row, "ytd_debit"))
        credit = _number(_field(row, "ytd_credit"))
    elif period_basis == "DELTA":
        debit = _number(_field(row, "closing_debit")) - _number(_field(row, "opening_debit"))
        credit = _number(_field(row, "closing_credit")) - _number(_field(row, "opening_credit"))
    else:
        debit = _number(_field(row, "closing_debit"))
        credit = _number(_field(row, "closing_credit"))
    if side == "DEBIT":
        return debit
    if side == "CREDIT":
        return credit
    if side in ("DEBIT_BALANCE", "DEBIT_NET"):
        return debit - credit
    if side in ("CREDIT_BALANCE", "CREDIT_NET"):
        return credit - debit
    return debit - credit


def _member_active(payload: dict[str, Any], company_key: str, period_to: date) -> bool:
    if normalize_key(payload.get("status") or "ACTIVE") == "INACTIVE":
        return False
    member_company = normalize_key(payload.get("companyKey"))
    if member_company and member_company != normalize_key(company_key):
        return False
    valid_from = _text(payload.get("validFrom"))
    valid_to = _text(payload.get("validTo"))
    return not (valid_from and valid_from > period_to.isoformat()) and not (
        valid_to and valid_to < period_to.isoformat()
    )


def _account_matches(account_number: Any, member: dict[str, Any], row: dict[str, Any]) -> bool:
    account = re.sub(r"\s+", "", str(account_number or "").strip())
    raw_value = str(member.get("matchValue") or member.get("accountCode") or "").strip()
    match_type = normalize_key(member.get("matchType") or "EXACT")
    if not account or not raw_value:
        return False

    def matches_one(value: str, rule_type: str) -> bool:
        normalized = re.sub(r"\s+", "", value or "")
        normalized = re.sub(r"^TK", "", normalized, flags=re.IGNORECASE)
        if not normalized:
            return False
        if rule_type == "PREFIX":
            is_leaf = _field(row, "is_leaf")
            return account.startswith(normalized) and is_leaf is not False
        return account == normalized

    if match_type == "MIXED":
        parts = [part.strip() for part in re.split(r"[,;\n]+", raw_value) if part.strip()]
        return any(
            matches_one(part[:-1], "PREFIX") if part.rstrip().endswith("*") else matches_one(part, "EXACT")
            for part in parts
        )
    if match_type == "PREFIX":
        value = re.sub(r"\s+", "", raw_value)
        is_leaf = _field(row, "is_leaf")
        return account.startswith(value) and is_leaf is not False
    return matches_one(raw_value, "EXACT")


def _rows_for_set(
    balance_rows: list[dict[str, Any]],
    account_set_key: Any,
    members: list[dict[str, Any]],
    period_to: date,
) -> list[dict[str, Any]]:
    set_key = normalize_key(account_set_key)
    relevant = [
        record
        for record in members
        if normalize_key((record.get("payload") or {}).get("accountSetKey") or (record.get("payload") or {}).get("setKey"))
        == set_key
    ]
    if not set_key or not relevant:
        return []
    result: list[dict[str, Any]] = []
    for row in balance_rows:
        company_key = _text(_field(row, "company_key"))
        account_number = _field(row, "account_number")
        included = any(
            normalize_key((record.get("payload") or {}).get("role") or "INCLUDE") == "INCLUDE"
            and _member_active(record.get("payload") or {}, company_key, period_to)
            and _account_matches(account_number, record.get("payload") or {}, row)
            for record in relevant
        )
        excluded = any(
            normalize_key((record.get("payload") or {}).get("role")) == "EXCLUDE"
            and _member_active(record.get("payload") or {}, company_key, period_to)
            and _account_matches(account_number, record.get("payload") or {}, row)
            for record in relevant
        )
        if included and not excluded:
            result.append(row)
    return result


def _balance_sheet_overlap_messages(
    *,
    row_records: list[dict[str, Any]],
    rules: dict[str, dict[str, Any]],
    members: list[dict[str, Any]],
    source_rows: list[dict[str, Any]],
    period_to: date,
    column_payload: dict[str, Any],
) -> list[str]:
    """Find non-zero source balances assigned to more than one leaf line.

    A balance account may legitimately appear once on the debit side and once on
    the credit side (for example receivables/advances).  It must not, however,
    feed two debit presentation lines or two credit presentation lines in the
    same column.  That is the common failure mode when a broad account prefix is
    used for both current and non-current classifications.
    """

    allocations: dict[tuple[str, str, str], dict[str, Any]] = {}
    excluded_row_types = {
        "FORMULA",
        "TOTAL",
        "SUBTOTAL",
        "SECTION",
        "HEADER",
        "TITLE",
        "SPACER",
    }
    column_basis = _column_period_basis(column_payload)
    for record in row_records:
        row_payload = record.get("payload") or {}
        if normalize_key(row_payload.get("rowType")) in excluded_row_types:
            continue
        row_code = _text(row_payload.get("rowCode") or record.get("displayCode"))
        if not row_code:
            continue
        linked_rule = rules.get(normalize_key(row_payload.get("ruleKey")), {})
        cells = row_payload.get("cells") if isinstance(row_payload.get("cells"), dict) else {}
        column_key = normalize_key(column_payload.get("columnKey"))
        cell = cells.get(column_key) or cells.get(column_payload.get("columnKey")) or {}
        has_inline_rule = any(
            cell.get(key) not in (None, "", [])
            for key in ("accountSetKey", "operator", "terms")
        )
        rule = cell if has_inline_rule else linked_rule
        terms = rule.get("terms") if isinstance(rule.get("terms"), list) else [rule]
        for term in terms:
            operator = normalize_key(term.get("operator") or "BALANCE")
            set_key = normalize_key(term.get("accountSetKey"))
            if operator not in ("BALANCE", "DELTA_BALANCE") or not set_key:
                continue
            side = normalize_key(term.get("side") or "NET")
            if side in ("DEBIT", "DEBIT_BALANCE", "DEBIT_NET", "NET"):
                direction = "NỢ"
            elif side in ("CREDIT", "CREDIT_BALANCE", "CREDIT_NET"):
                direction = "CÓ"
            else:
                continue
            period_basis = "" if term.get("useColumnBasis") is False else column_basis
            if operator == "DELTA_BALANCE":
                period_basis = "DELTA"
            for source_row in _rows_for_set(source_rows, set_key, members, period_to):
                value = _metric_value(
                    source_row,
                    term.get("metricKey"),
                    term,
                    period_basis,
                )
                if abs(value) < 0.5:
                    continue
                account_number = _text(_field(source_row, "account_number"))
                company_key = normalize_key(_field(source_row, "company_key"))
                if not account_number:
                    continue
                key = (company_key, account_number, direction)
                allocation = allocations.setdefault(
                    key,
                    {
                        "rowCodes": set(),
                        "accountName": _text(_field(source_row, "account_name")),
                    },
                )
                allocation["rowCodes"].add(row_code)

    messages: list[str] = []
    column_label = _text(
        column_payload.get("columnLabel") or column_payload.get("columnKey")
    )
    for (company_key, account_number, direction), allocation in sorted(allocations.items()):
        row_codes = sorted(allocation["rowCodes"])
        if len(row_codes) < 2:
            continue
        account_name = allocation.get("accountName") or ""
        account_label = f"TK {account_number}"
        if account_name:
            account_label += f" - {account_name}"
        company_label = f" của {company_key}" if company_key else ""
        messages.append(
            f"{column_label}: {account_label}{company_label} đang được tính bên {direction} "
            f"đồng thời vào các dòng {', '.join(row_codes)}. "
            "Hãy gán tài khoản này vào đúng một dòng ngắn hạn hoặc dài hạn trước khi xuất chính thức."
        )
    return messages


def _column_period_basis(column_payload: dict[str, Any]) -> str:
    selector = normalize_key(column_payload.get("periodSelector"))
    if selector in ("OPENING_OF_CURRENT_RANGE", "END_OF_PRIOR_YEAR"):
        return "OPENING"
    value_basis = normalize_key(column_payload.get("valueBasis"))
    aliases = {
        "OPENING_BALANCE": "OPENING",
        "CLOSING": "ENDING",
        "CLOSING_BALANCE": "ENDING",
        "CURRENT_MOVEMENT": "PERIOD",
        "YEAR_TO_DATE": "YTD",
    }
    value_basis = aliases.get(value_basis, value_basis)
    return value_basis if value_basis in ("OPENING", "PERIOD", "YTD", "ENDING", "DELTA") else ""


def _path_value(value: Any, path: Any) -> Any:
    current = value
    for part in [item for item in re.split(r"[.\[\]]+", str(path or "")) if item]:
        if not isinstance(current, dict):
            return None
        expected = part.lower()
        key = next((key for key in current if str(key).lower() == expected), None)
        if key is None:
            return None
        current = current[key]
    return current


def _canonical_value(
    row: dict[str, Any],
    field_name: str,
    field_mappings: list[dict[str, Any]],
) -> Any:
    direct = _field(row, field_name)
    if direct not in (None, ""):
        return direct
    canonical = row.get("canonical_data") or row.get("canonicalData") or {}
    if not isinstance(canonical, dict):
        return None
    direct = _path_value(canonical, field_name)
    if direct not in (None, ""):
        return direct
    company_key = normalize_key(_field(row, "company_key"))
    for record in field_mappings:
        payload = record.get("payload") or {}
        if _text(payload.get("canonicalField")).lower() != field_name.lower():
            continue
        mapping_source = payload.get("sourceKey")
        mapping_company = normalize_key(payload.get("companyKey"))
        if mapping_source and not _source_matches(row, mapping_source):
            continue
        if mapping_company and mapping_company != company_key:
            continue
        mapped = _path_value(canonical, payload.get("sourceField"))
        if mapped not in (None, ""):
            return mapped
    return None


_SOURCE_KEY_ALIAS_GROUPS = (
    frozenset(
        {
            "AMIS_SO_CHI_TIET",
            "AMIS_GENERAL_LEDGER",
            # Khối tạo bằng cách nạp Excel phải dùng cùng công thức nghiệp vụ
            # với khối lấy trực tiếp từ MISA. Phần hậu tố chỉ định đúng kết nối
            # vẫn được giữ ở source_key; base_source_key xác định loại dữ liệu.
            "FILE_GENERAL_LEDGER",
        }
    ),
    frozenset(
        {
            "AMIS_CAN_DOI_THOI_GIAN_THUC",
            "AMIS_TRIAL_BALANCE",
            "FILE_TRIAL_BALANCE",
        }
    ),
)


def _source_key_aliases(source_key: Any) -> set[str]:
    """Return equivalent stable source keys across legacy and V3 profiles.

    The report templates intentionally keep their historical source keys so an
    approved formula does not have to be rewritten whenever the connection UI
    creates a newer profile.  A concrete connection key (the key with a hash
    suffix) still remains exact; only the stable base keys are aliased.
    """

    normalized = normalize_key(source_key)
    if not normalized:
        return set()
    for aliases in _SOURCE_KEY_ALIAS_GROUPS:
        if normalized in aliases:
            return set(aliases)
    return {normalized}


def _source_matches(row: dict[str, Any], source_key: Any) -> bool:
    expected = _source_key_aliases(source_key)
    if not expected:
        return True
    actual: set[str] = set()
    for value in (
        _field(row, "source_key"),
        _field(row, "base_source_key"),
    ):
        actual.update(_source_key_aliases(value))
    return bool(expected & actual)


def _account_in_set(
    account_number: Any,
    account_set_key: Any,
    members: list[dict[str, Any]],
    company_key: str,
    period_to: date,
    member_cache: dict[tuple[Any, ...], Any] | None = None,
) -> bool:
    set_key = normalize_key(account_set_key)
    if not set_key:
        return False
    normalized_company = normalize_key(company_key)
    normalized_account = re.sub(r"\s+", "", str(account_number or "").strip())
    match_cache_key = (
        "MATCH",
        set_key,
        normalized_company,
        period_to.isoformat(),
        normalized_account,
    )
    if member_cache is not None and match_cache_key in member_cache:
        return bool(member_cache[match_cache_key])
    cache_key = ("MEMBERS", set_key, normalized_company, period_to.isoformat())
    relevant = member_cache.get(cache_key) if member_cache is not None else None
    if relevant is None:
        relevant = [
            record
            for record in members
            if normalize_key((record.get("payload") or {}).get("accountSetKey") or (record.get("payload") or {}).get("setKey"))
            == set_key
            and _member_active(record.get("payload") or {}, company_key, period_to)
        ]
        if member_cache is not None:
            member_cache[cache_key] = relevant
    if not relevant:
        return False
    fake_row = {"is_leaf": True}
    included = any(
        normalize_key((record.get("payload") or {}).get("role") or "INCLUDE") == "INCLUDE"
        and _account_matches(account_number, record.get("payload") or {}, fake_row)
        for record in relevant
    )
    excluded = any(
        normalize_key((record.get("payload") or {}).get("role")) == "EXCLUDE"
        and _account_matches(account_number, record.get("payload") or {}, fake_row)
        for record in relevant
    )
    result = included and not excluded
    if member_cache is not None:
        member_cache[match_cache_key] = result
    return result


def _truthy(value: Any) -> bool:
    return value is True or normalize_key(value) in ("1", "TRUE", "YES", "Y", "CO")


def _filters_match(
    row: dict[str, Any],
    filters: Any,
    field_mappings: list[dict[str, Any]],
) -> bool:
    if not isinstance(filters, dict):
        return True

    def comparison_text(value: Any) -> str:
        return "".join(
            character
            for character in unicodedata.normalize("NFD", _text(value).replace("đ", "d").replace("Đ", "D"))
            if unicodedata.category(character) != "Mn"
        ).casefold()

    def actual_value(field: Any) -> Any:
        normalized_field = normalize_key(field)
        if normalized_field in ("SEARCH_TEXT", "DOCUMENT_TEXT", "NOI_DUNG_CHUNG_TU"):
            # Trường ảo này giúp cấu hình nghiệp vụ theo nội dung chứng từ mà
            # không khóa công thức vào tên một cá nhân/nhà cung cấp cụ thể.
            return " | ".join(
                _text(_canonical_value(row, candidate, field_mappings))
                for candidate in (
                    "description",
                    "journal_memo",
                    "expense_item_name",
                    "reference_no",
                )
                if _text(_canonical_value(row, candidate, field_mappings))
            )
        return _canonical_value(row, str(field), field_mappings)

    def expected_values(value: Any) -> list[Any]:
        if isinstance(value, list):
            return value
        return [
            item.strip()
            for item in re.split(r"[;,|\n]+", _text(value))
            if item.strip()
        ]

    def condition_matches(condition: Any) -> bool:
        if not isinstance(condition, dict):
            return True
        actual = actual_value(condition.get("field"))
        operator = normalize_key(condition.get("operator") or "EQUALS_ANY")
        values = expected_values(condition.get("value"))
        actual_text = comparison_text(actual)
        value_texts = [comparison_text(item) for item in values]
        if operator in ("IS_EMPTY", "EMPTY"):
            return not actual_text
        if operator in ("NOT_EMPTY", "IS_NOT_EMPTY"):
            return bool(actual_text)
        if operator in ("EQUALS", "EQUALS_ANY", "IN"):
            return actual_text in set(value_texts)
        if operator in ("NOT_EQUALS", "NOT_EQUALS_ANY", "NOT_IN"):
            return actual_text not in set(value_texts)
        if operator in ("CONTAINS", "CONTAINS_ANY"):
            return any(value and value in actual_text for value in value_texts)
        if operator in ("NOT_CONTAINS", "NOT_CONTAINS_ANY"):
            return all(not value or value not in actual_text for value in value_texts)
        if operator in ("STARTS_WITH", "STARTS_WITH_ANY", "PREFIX"):
            return any(value and actual_text.startswith(value) for value in value_texts)
        return False

    structured_conditions = filters.get("conditions")
    if isinstance(structured_conditions, list):
        results = [condition_matches(condition) for condition in structured_conditions]
        logic = normalize_key(filters.get("logic") or "ALL")
        structured_match = any(results) if logic in ("ANY", "OR") else all(results)
        if not structured_match:
            return False

    for field, expected in filters.items():
        if field in ("conditions", "logic"):
            continue
        actual = actual_value(field)
        if isinstance(expected, list):
            if str(actual) not in {str(item) for item in expected}:
                return False
        elif isinstance(expected, dict) and "not" in expected:
            if str(actual) == str(expected["not"]):
                return False
        elif str(actual) != str(expected):
            return False
    return True


_JOURNAL_DETAIL_FILTER_FIELDS = {
    "SEARCH_TEXT",
    "DOCUMENT_TEXT",
    "NOI_DUNG_CHUNG_TU",
    "DESCRIPTION",
    "JOURNAL_MEMO",
    "REFERENCE_NO",
    "REFERENCE_ID",
}


def _filters_require_journal_detail(filters: Any) -> bool:
    if not isinstance(filters, dict):
        return False
    for condition in filters.get("conditions") or []:
        if not isinstance(condition, dict):
            continue
        if normalize_key(condition.get("field")) in _JOURNAL_DETAIL_FILTER_FIELDS:
            return True
    return any(
        normalize_key(field) in _JOURNAL_DETAIL_FILTER_FIELDS
        for field in filters
        if field not in ("conditions", "logic")
    )


def _source_rows_for_term(
    balance_rows: list[dict[str, Any]],
    term: dict[str, Any],
    field_mappings: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    dataset_type = normalize_key(term.get("datasetType"))
    candidates = [
        row
        for row in balance_rows
        if _source_matches(row, term.get("sourceKey"))
        and (not dataset_type or normalize_key(_field(row, "dataset_type")) == dataset_type)
    ]
    granularities = {
        normalize_key(_canonical_value(row, "journal_granularity", field_mappings))
        for row in candidates
    }
    if _filters_require_journal_detail(term.get("filters")) and "DETAIL" in granularities:
        return [
            row
            for row in candidates
            if normalize_key(_canonical_value(row, "journal_granularity", field_mappings)) == "DETAIL"
        ]
    if "MONTHLY" in granularities:
        return [
            row
            for row in candidates
            if normalize_key(_canonical_value(row, "journal_granularity", field_mappings)) == "MONTHLY"
        ]
    return candidates


def _cash_flow_condition_matches(
    row: dict[str, Any],
    condition_key: Any,
    field_mappings: list[dict[str, Any]],
    condition_requirement: Any = "",
) -> bool:
    """Use MISA's cash-flow code and only fall back when the rule allows it.

    Legacy rules keep the configured debit/credit account-pair fallback so historical
    snapshots continue to work. High-risk cash-flow lines can set
    ``conditionRequirement=REQUIRED`` (or the longer policy name used by V5) to reject
    unclassified entries instead of silently assigning them to the wrong line.
    """

    expected = normalize_key(condition_key)
    if not expected:
        return True
    actual_value = _canonical_value(row, "cash_flow_code", field_mappings)
    actual_codes = [
        normalize_key(item)
        for item in re.split(r"[,;|\n]+", _text(actual_value))
        if normalize_key(item)
    ]
    if not actual_codes:
        requirement = normalize_key(condition_requirement)
        return requirement not in (
            "REQUIRED",
            "REQUIRE_CASH_FLOW_CODE",
            "REQUIRE_CASH_FLOW_CODE_OR_APPROVED_ADJUSTMENT",
        )
    expected_match = re.search(r"(?:DONG_TIEN|CASH_FLOW)_([A-Z0-9]+)$", expected)
    expected_code = expected_match.group(1) if expected_match else expected

    def equivalent(actual: str) -> bool:
        if actual == expected:
            return True
        actual_match = re.search(r"(?:DONG_TIEN|CASH_FLOW)_([A-Z0-9]+)$", actual)
        actual_code = actual_match.group(1) if actual_match else actual
        if actual_code == expected_code:
            return True
        if actual_code.isdigit() and expected_code.isdigit():
            return int(actual_code) == int(expected_code)
        return False

    return any(equivalent(actual) for actual in actual_codes)


def _detail_row_value(
    row: dict[str, Any],
    term: dict[str, Any],
    members: list[dict[str, Any]],
    field_mappings: list[dict[str, Any]],
    period_to: date,
    member_cache: dict[tuple[Any, ...], Any] | None = None,
) -> float | None:
    company_key = _text(_field(row, "company_key"))
    debit_account = _canonical_value(row, "debit_account", field_mappings)
    credit_account = _canonical_value(row, "credit_account", field_mappings)
    if not debit_account and not credit_account:
        return None
    primary_set = term.get("accountSetKey")
    debit_primary = _account_in_set(
        debit_account, primary_set, members, company_key, period_to, member_cache
    )
    credit_primary = _account_in_set(
        credit_account, primary_set, members, company_key, period_to, member_cache
    )
    if not debit_primary and not credit_primary:
        return 0.0
    if (
        normalize_key(term.get("operator")) == "CASH_MOVEMENT"
        and not _cash_flow_condition_matches(
            row,
            term.get("conditionKey"),
            field_mappings,
            term.get("conditionRequirement") or term.get("classificationPolicy"),
        )
    ):
        return 0.0
    counter_set = term.get("counterAccountSetKey")
    if counter_set:
        debit_primary = debit_primary and _account_in_set(
            credit_account, counter_set, members, company_key, period_to, member_cache
        )
        credit_primary = credit_primary and _account_in_set(
            debit_account, counter_set, members, company_key, period_to, member_cache
        )
        if not debit_primary and not credit_primary:
            return 0.0
    intercompany_mode = normalize_key(term.get("intercompanyMode"))
    is_intercompany = bool(_text(_canonical_value(row, "intercompany_partner", field_mappings)))
    if intercompany_mode in ("EXCLUDE", "EXCLUDE_INTERCOMPANY") and is_intercompany:
        return 0.0
    if intercompany_mode in ("ONLY", "INCLUDE_ONLY", "ONLY_INTERCOMPANY") and not is_intercompany:
        return 0.0
    if term.get("excludeClosingEntries", True) is not False:
        entry_type = normalize_key(_canonical_value(row, "entry_type", field_mappings))
        closing_flag = _canonical_value(row, "is_closing_entry", field_mappings)
        if _truthy(closing_flag) or entry_type in ("CLOSING", "CLOSE", "KET_CHUYEN"):
            return 0.0
        if str(debit_account or "").startswith("911") or str(credit_account or "").startswith("911"):
            return 0.0
    if term.get("excludeReversals") is True:
        reversal = _canonical_value(row, "is_reversal", field_mappings)
        entry_type = normalize_key(_canonical_value(row, "entry_type", field_mappings))
        if _truthy(reversal) or entry_type in ("REVERSAL", "REVERSE", "DAO"):
            return 0.0
    maturity_class = normalize_key(term.get("maturityClass"))
    if maturity_class:
        current_flag = _canonical_value(row, "current_flag", field_mappings)
        if current_flag in (None, ""):
            return 0.0
        is_current = _truthy(current_flag) or normalize_key(current_flag) in ("SHORT_TERM", "CURRENT")
        if maturity_class in ("SHORT_TERM", "CURRENT") and not is_current:
            return 0.0
        if maturity_class in ("LONG_TERM", "NON_CURRENT") and is_current:
            return 0.0
    if not _filters_match(row, term.get("filters"), field_mappings):
        return 0.0
    # Bút toán điều chỉnh âm phải giữ nguyên dấu. Dùng trị tuyệt đối sẽ biến
    # bút toán đảo thành tăng doanh thu/chi phí và làm sai báo cáo kỳ.
    amount = _number(_canonical_value(row, "amount", field_mappings))
    side = normalize_key(term.get("side") or "NET")
    if side == "DEBIT":
        return amount if debit_primary else 0.0
    if side == "CREDIT":
        return amount if credit_primary else 0.0
    return (amount if debit_primary else 0.0) - (amount if credit_primary else 0.0)


def _term_value(
    term: dict[str, Any],
    balance_rows: list[dict[str, Any]],
    members: list[dict[str, Any]],
    field_mappings: list[dict[str, Any]],
    period_to: date,
    column_payload: dict[str, Any],
    warning: Callable[[str], None],
    runtime_cache: dict[str, Any] | None = None,
    blocking_issue: Callable[[str], None] | None = None,
) -> float:
    operator = normalize_key(term.get("operator") or "BALANCE")
    sign = -1 if _number(term.get("sign")) < 0 else 1
    if operator == "MANUAL_ADJUSTMENT":
        if term.get("manualValue") not in (None, ""):
            message = (
                "Công thức có manualValue bị từ chối vì đây là số ghi cứng. "
                "Chỉ khoản điều chỉnh đã lưu, có lý do và lịch sử phê duyệt mới được sử dụng."
            )
            (blocking_issue or warning)(message)
            return 0.0
        expected_column = normalize_key(term.get("_columnKey"))
        manual_rows = [
            row
            for row in balance_rows
            if _source_matches(row, term.get("sourceKey"))
            and normalize_key(_field(row, "dataset_type")) == "MANUAL_ADJUSTMENT"
            and _filters_match(row, term.get("filters"), field_mappings)
            and (
                not normalize_key(_canonical_value(row, "column_key", field_mappings))
                or not expected_column
                or normalize_key(_canonical_value(row, "column_key", field_mappings))
                == expected_column
            )
        ]
        return sign * sum(
            _number(_canonical_value(row, "amount", field_mappings)) for row in manual_rows
        )
    set_key = normalize_key(term.get("accountSetKey"))
    if not set_key:
        warning("Còn thành phần công thức chưa chọn danh mục tài khoản.")
        return 0.0
    dataset_type = normalize_key(term.get("datasetType"))
    runtime_cache = runtime_cache if runtime_cache is not None else {}
    source_cache = runtime_cache.setdefault("source_rows", {})
    source_cache_key = (
        id(balance_rows),
        normalize_key(term.get("sourceKey")),
        dataset_type,
        "DETAIL" if _filters_require_journal_detail(term.get("filters")) else "SUMMARY",
    )
    source_rows = source_cache.get(source_cache_key)
    if source_rows is None:
        source_rows = _source_rows_for_term(balance_rows, term, field_mappings)
        source_cache[source_cache_key] = source_rows
    member_cache = runtime_cache.setdefault("account_membership", {})
    condition_requirement = normalize_key(
        term.get("conditionRequirement") or term.get("classificationPolicy")
    )
    classification_policy = normalize_key(term.get("classificationPolicy"))
    condition_required = condition_requirement in (
        "REQUIRED",
        "REQUIRE_CASH_FLOW_CODE",
        "REQUIRE_CASH_FLOW_CODE_OR_APPROVED_ADJUSTMENT",
    )
    if operator == "CASH_MOVEMENT" and condition_required and term.get("conditionKey"):
        missing_entry_count = 0
        fallback_term = {**term, "conditionKey": "", "conditionRequirement": ""}
        for row in source_rows:
            if _text(_canonical_value(row, "cash_flow_code", field_mappings)):
                continue
            candidate = _detail_row_value(
                row,
                fallback_term,
                members,
                field_mappings,
                period_to,
                member_cache,
            )
            if candidate is None or abs(candidate) <= 0.0000001:
                continue
            missing_entry_count += max(
                1,
                int(_number(_canonical_value(row, "entry_count", field_mappings))),
            )
        approved_adjustment_exists = False
        if (
            missing_entry_count
            and classification_policy == "REQUIRE_CASH_FLOW_CODE_OR_APPROVED_ADJUSTMENT"
            and term.get("_reportKey")
            and term.get("_rowCode")
        ):
            expected_column = normalize_key(term.get("_columnKey"))
            approved_adjustment_exists = any(
                normalize_key(_field(row, "dataset_type")) == "MANUAL_ADJUSTMENT"
                and normalize_key(_canonical_value(row, "approval_status", field_mappings))
                == "APPROVED"
                and normalize_key(_canonical_value(row, "report_key", field_mappings))
                == normalize_key(term.get("_reportKey"))
                and normalize_key(_canonical_value(row, "row_code", field_mappings))
                == normalize_key(term.get("_rowCode"))
                and (
                    not normalize_key(_canonical_value(row, "column_key", field_mappings))
                    or not expected_column
                    or normalize_key(_canonical_value(row, "column_key", field_mappings))
                    == expected_column
                )
                for row in balance_rows
            )
        if missing_entry_count and approved_adjustment_exists:
            warning(
                f"{missing_entry_count} bút toán của {_text(term.get('conditionKey'))} chưa có mã dòng tiền; "
                "báo cáo dùng khoản điều chỉnh đã được duyệt và lưu lịch sử."
            )
        elif missing_entry_count:
            message = (
                f"Không đủ căn cứ phân loại {missing_entry_count} bút toán cho "
                f"{_text(term.get('conditionKey'))}: nguồn thiếu mã dòng tiền. "
                "Hệ thống không tự gán theo tài khoản đối ứng; cần bổ sung mã nguồn "
                "hoặc khoản điều chỉnh đã được duyệt và có lịch sử."
            )
            (blocking_issue or warning)(message)
    if operator == "PROFIT_LOSS_ACTIVITY":
        side = normalize_key(term.get("side") or "CREDIT")
        if side not in ("CREDIT", "INCOME", "DEBIT", "EXPENSE"):
            warning("Hoạt động lãi/lỗ phải chọn bên Thu nhập (Có) hoặc Chi phí (Nợ).")
            return 0.0
        closing_total = 0.0
        closing_count = 0
        movement_debit = 0.0
        movement_credit = 0.0
        for row in source_rows:
            company_key = _text(_field(row, "company_key"))
            debit_account = _canonical_value(row, "debit_account", field_mappings)
            credit_account = _canonical_value(row, "credit_account", field_mappings)
            debit_primary = _account_in_set(
                debit_account, set_key, members, company_key, period_to, member_cache
            )
            credit_primary = _account_in_set(
                credit_account, set_key, members, company_key, period_to, member_cache
            )
            if not debit_primary and not credit_primary:
                continue
            intercompany_mode = normalize_key(term.get("intercompanyMode"))
            is_intercompany = bool(
                _text(_canonical_value(row, "intercompany_partner", field_mappings))
            )
            if intercompany_mode in ("EXCLUDE", "EXCLUDE_INTERCOMPANY") and is_intercompany:
                continue
            if intercompany_mode in ("ONLY", "INCLUDE_ONLY", "ONLY_INTERCOMPANY") and not is_intercompany:
                continue
            entry_type = normalize_key(
                _canonical_value(row, "entry_type", field_mappings)
            )
            if term.get("excludeReversals") is True and (
                _truthy(_canonical_value(row, "is_reversal", field_mappings))
                or entry_type in ("REVERSAL", "REVERSE", "DAO")
            ):
                continue
            if not _filters_match(row, term.get("filters"), field_mappings):
                continue
            amount = _number(_canonical_value(row, "amount", field_mappings))
            is_closing = (
                _truthy(_canonical_value(row, "is_closing_entry", field_mappings))
                or entry_type in ("CLOSING", "CLOSE", "KET_CHUYEN")
                or str(debit_account or "").startswith("911")
                or str(credit_account or "").startswith("911")
            )
            if is_closing:
                is_income_closing = (
                    side in ("CREDIT", "INCOME")
                    and debit_primary
                    and str(credit_account or "").startswith("911")
                )
                is_expense_closing = (
                    side in ("DEBIT", "EXPENSE")
                    and credit_primary
                    and str(debit_account or "").startswith("911")
                )
                if is_income_closing or is_expense_closing:
                    closing_total += amount
                    closing_count += 1
                continue
            if debit_primary:
                movement_debit += amount
            if credit_primary:
                movement_credit += amount
        if closing_count:
            return sign * closing_total
        fallback = (
            movement_credit - movement_debit
            if side in ("CREDIT", "INCOME")
            else movement_debit - movement_credit
        )
        return sign * fallback
    detail_values = [
        value
        for row in source_rows
        if (
            value := _detail_row_value(
                row, term, members, field_mappings, period_to, member_cache
            )
        ) is not None
    ]
    needs_detail = operator == "CASH_MOVEMENT" or bool(term.get("counterAccountSetKey"))
    if detail_values and (needs_detail or operator in ("MOVEMENT", "MOVEMENT_DETAIL")):
        return sign * sum(detail_values)
    if needs_detail or operator == "MOVEMENT_DETAIL":
        warning(
            f"Danh mục {set_key} cần sổ chi tiết có tài khoản Nợ/Có đối ứng; "
            "dữ liệu đang chọn chưa có các trường này."
        )
        return 0.0
    rows = _rows_for_set(source_rows, set_key, members, period_to)
    if not rows:
        has_definition = any(
            normalize_key(
                (record.get("payload") or {}).get("accountSetKey")
                or (record.get("payload") or {}).get("setKey")
            )
            == set_key
            for record in members
        )
        if not has_definition:
            warning(f"Danh mục {set_key} chưa có tài khoản thành phần trong cấu hình.")
    period_basis = (
        ""
        if term.get("useColumnBasis") is False
        else _column_period_basis(column_payload)
    )
    if operator == "DELTA_BALANCE":
        period_basis = "DELTA"
    subtotal = sum(
        _metric_value(row, term.get("metricKey"), term, period_basis) for row in rows
    )
    return sign * subtotal


def _evidence_account_name(
    account_number: Any,
    company_key: Any,
    row: dict[str, Any],
    account_names: dict[tuple[str, str], str],
) -> str:
    direct = _text(_field(row, "account_name"))
    if direct:
        return direct
    code = _text(account_number)
    company = normalize_key(company_key)
    return account_names.get((company, code)) or account_names.get(("", code)) or ""


def _evidence_component(
    *,
    row: dict[str, Any],
    label: str,
    signed_value: float,
    account_number: Any = "",
    account_name: str = "",
    counter_account_number: Any = "",
    counter_account_name: str = "",
    entry_count: Any = 0,
    formula_sign: Any = None,
    source_name: str = "",
    trace_path: str = "",
) -> dict[str, Any]:
    dataset_type = _text(_field(row, "dataset_type"))
    friendly_sources = {
        "GENERAL_LEDGER": "Sổ nhật ký chung AMIS",
        "TRIAL_BALANCE": "Bảng cân đối tài khoản AMIS",
        "ACCOUNT_DETAIL": "Sổ chi tiết tài khoản AMIS",
        "ACCOUNT_LEDGER": "Sổ chi tiết tài khoản AMIS",
        "MANUAL_ADJUSTMENT": "Điều chỉnh kế toán đã lưu",
    }
    return {
        "label": label,
        "tracePath": trace_path or label,
        "signedValue": signed_value,
        "formulaSign": (
            -1
            if _number(formula_sign) < 0
            else 1
            if formula_sign not in (None, "")
            else -1
            if signed_value < 0
            else 1
        ),
        "accountNumber": _text(account_number),
        "accountName": _text(account_name),
        "counterAccountNumber": _text(counter_account_number),
        "counterAccountName": _text(counter_account_name),
        "companyKey": _text(_field(row, "company_key")),
        "dataSpaceName": _text(_field(row, "data_space_name")),
        "datasetType": dataset_type,
        "sourceName": _text(source_name)
        or friendly_sources.get(normalize_key(dataset_type), "Dữ liệu kế toán đã đồng bộ"),
        "sourceKey": _text(_field(row, "source_key")),
        "entryCount": max(0, int(_number(entry_count))),
        "periodFrom": _text(_field(row, "period_from")),
        "periodTo": _text(_field(row, "period_to")),
    }


def _term_evidence_components(
    term: dict[str, Any],
    balance_rows: list[dict[str, Any]],
    members: list[dict[str, Any]],
    field_mappings: list[dict[str, Any]],
    period_to: date,
    column_payload: dict[str, Any],
    account_names: dict[tuple[str, str], str],
    term_value: float,
    runtime_cache: dict[str, Any] | None = None,
) -> list[dict[str, Any]]:
    operator = normalize_key(term.get("operator") or "BALANCE")
    sign = -1 if _number(term.get("sign")) < 0 else 1
    if operator == "MANUAL_ADJUSTMENT":
        if term.get("manualValue") not in (None, ""):
            return [
                _evidence_component(
                    row={},
                    label="Số ghi cứng trong công thức đã bị hệ thống loại bỏ",
                    signed_value=0.0,
                    source_name="Kiểm soát chống ghi cứng số liệu",
                )
            ]
        expected_column = normalize_key(term.get("_columnKey"))
        rows = [
            row
            for row in balance_rows
            if _source_matches(row, term.get("sourceKey"))
            and normalize_key(_field(row, "dataset_type")) == "MANUAL_ADJUSTMENT"
            and _filters_match(row, term.get("filters"), field_mappings)
            and (
                not normalize_key(_canonical_value(row, "column_key", field_mappings))
                or not expected_column
                or normalize_key(_canonical_value(row, "column_key", field_mappings))
                == expected_column
            )
        ]
        return [
            _evidence_component(
                row=row,
                label=(
                    "Khoản điều chỉnh đã duyệt: "
                    + _text(_canonical_value(row, "reason", field_mappings))
                ).rstrip(": "),
                signed_value=sign * _number(_canonical_value(row, "amount", field_mappings)),
                source_name=(
                    "Lịch sử điều chỉnh kế toán · "
                    + _text(_canonical_value(row, "source_reference", field_mappings))
                ).rstrip(" ·"),
            )
            for row in rows
        ]

    set_key = normalize_key(term.get("accountSetKey"))
    dataset_type = normalize_key(term.get("datasetType"))
    source_rows = _source_rows_for_term(balance_rows, term, field_mappings)
    runtime_cache = runtime_cache if runtime_cache is not None else {}
    member_cache = runtime_cache.setdefault("account_membership", {})
    needs_detail = operator == "CASH_MOVEMENT" or bool(term.get("counterAccountSetKey"))

    if operator == "PROFIT_LOSS_ACTIVITY":
        side = normalize_key(term.get("side") or "CREDIT")
        closing_items: list[tuple[dict[str, Any], str, str, float, str]] = []
        movement_items: list[tuple[dict[str, Any], str, str, float, str]] = []
        for row in source_rows:
            company_key = _text(_field(row, "company_key"))
            debit_account = _text(_canonical_value(row, "debit_account", field_mappings))
            credit_account = _text(_canonical_value(row, "credit_account", field_mappings))
            debit_primary = _account_in_set(
                debit_account, set_key, members, company_key, period_to, member_cache
            )
            credit_primary = _account_in_set(
                credit_account, set_key, members, company_key, period_to, member_cache
            )
            if not debit_primary and not credit_primary:
                continue
            intercompany_mode = normalize_key(term.get("intercompanyMode"))
            is_intercompany = bool(
                _text(_canonical_value(row, "intercompany_partner", field_mappings))
            )
            if intercompany_mode in ("EXCLUDE", "EXCLUDE_INTERCOMPANY") and is_intercompany:
                continue
            if intercompany_mode in ("ONLY", "INCLUDE_ONLY", "ONLY_INTERCOMPANY") and not is_intercompany:
                continue
            entry_type = normalize_key(
                _canonical_value(row, "entry_type", field_mappings)
            )
            if term.get("excludeReversals") is True and (
                _truthy(_canonical_value(row, "is_reversal", field_mappings))
                or entry_type in ("REVERSAL", "REVERSE", "DAO")
            ):
                continue
            if not _filters_match(row, term.get("filters"), field_mappings):
                continue
            amount = _number(_canonical_value(row, "amount", field_mappings))
            is_closing = (
                _truthy(_canonical_value(row, "is_closing_entry", field_mappings))
                or entry_type in ("CLOSING", "CLOSE", "KET_CHUYEN")
                or debit_account.startswith("911")
                or credit_account.startswith("911")
            )
            if is_closing:
                if side in ("CREDIT", "INCOME") and debit_primary and credit_account.startswith("911"):
                    closing_items.append(
                        (row, debit_account, credit_account, sign * amount, f"Kết chuyển TK {debit_account} sang TK {credit_account}")
                    )
                elif side in ("DEBIT", "EXPENSE") and credit_primary and debit_account.startswith("911"):
                    closing_items.append(
                        (row, credit_account, debit_account, sign * amount, f"Kết chuyển TK {credit_account} sang TK {debit_account}")
                    )
                continue
            if side in ("CREDIT", "INCOME"):
                if credit_primary:
                    movement_items.append(
                        (row, credit_account, debit_account, sign * amount, f"Phát sinh Có TK {credit_account}; đối ứng Nợ TK {debit_account}")
                    )
                if debit_primary:
                    movement_items.append(
                        (row, debit_account, credit_account, -sign * amount, f"Giảm thu nhập: phát sinh Nợ TK {debit_account}; đối ứng Có TK {credit_account}")
                    )
            else:
                if debit_primary:
                    movement_items.append(
                        (row, debit_account, credit_account, sign * amount, f"Phát sinh Nợ TK {debit_account}; đối ứng Có TK {credit_account}")
                    )
                if credit_primary:
                    movement_items.append(
                        (row, credit_account, debit_account, -sign * amount, f"Giảm chi phí: phát sinh Có TK {credit_account}; đối ứng Nợ TK {debit_account}")
                    )

        selected_items = closing_items if closing_items else movement_items
        grouped: dict[tuple[str, ...], dict[str, Any]] = {}
        for row, primary_account, counter_account, signed_value, action_label in selected_items:
            company_key = _text(_field(row, "company_key"))
            primary_name = _evidence_account_name(
                primary_account, company_key, row, account_names
            )
            counter_name = _evidence_account_name(
                counter_account, company_key, row, account_names
            )
            label = (
                f"{action_label}{' - ' + primary_name if primary_name else ''}"
            )
            group_key = (
                _text(_field(row, "source_key")),
                company_key,
                _text(_field(row, "data_space_name")),
                primary_account,
                counter_account,
                label,
                _text(_field(row, "period_from")),
                _text(_field(row, "period_to")),
            )
            existing = grouped.get(group_key)
            row_entry_count = max(
                0,
                int(_number(_canonical_value(row, "entry_count", field_mappings))),
            )
            if existing:
                existing["signedValue"] = _number(existing.get("signedValue")) + signed_value
                existing["entryCount"] = int(_number(existing.get("entryCount"))) + row_entry_count
            else:
                grouped[group_key] = _evidence_component(
                    row=row,
                    label=label,
                    signed_value=signed_value,
                    account_number=primary_account,
                    account_name=primary_name,
                    counter_account_number=counter_account,
                    counter_account_name=counter_name,
                    entry_count=row_entry_count,
                )
        return list(grouped.values())

    if needs_detail or operator in ("MOVEMENT", "MOVEMENT_DETAIL"):
        grouped: dict[tuple[str, ...], dict[str, Any]] = {}
        for row in source_rows:
            value = _detail_row_value(
                row, term, members, field_mappings, period_to, member_cache
            )
            if value is None or abs(value) < 0.0000001:
                continue
            debit_account = _text(_canonical_value(row, "debit_account", field_mappings))
            credit_account = _text(_canonical_value(row, "credit_account", field_mappings))
            side = normalize_key(term.get("side") or "NET")
            primary_account = debit_account if side == "DEBIT" else credit_account if side == "CREDIT" else debit_account or credit_account
            counter_account = credit_account if side == "DEBIT" else debit_account if side == "CREDIT" else credit_account
            company_key = _text(_field(row, "company_key"))
            primary_name = _evidence_account_name(primary_account, company_key, row, account_names)
            counter_name = _evidence_account_name(counter_account, company_key, row, account_names)
            group_key = (
                _text(_field(row, "source_key")), company_key,
                _text(_field(row, "data_space_name")), primary_account,
                counter_account, _text(_field(row, "period_from")),
                _text(_field(row, "period_to")),
            )
            signed_value = sign * value
            entry_count = max(
                0,
                int(
                    _number(
                        _canonical_value(row, "entry_count", field_mappings)
                    )
                ),
            )
            label = (
                f"Ghi Nợ TK {primary_account}{' - ' + primary_name if primary_name else ''}; đối ứng Có TK {counter_account}{' - ' + counter_name if counter_name else ''}"
                if side == "DEBIT"
                else f"Ghi Có TK {primary_account}{' - ' + primary_name if primary_name else ''}; đối ứng Nợ TK {counter_account}{' - ' + counter_name if counter_name else ''}"
                if side == "CREDIT"
                else f"Bút toán Nợ TK {debit_account} / Có TK {credit_account}"
            )
            existing = grouped.get(group_key)
            if existing:
                existing["signedValue"] = _number(existing.get("signedValue")) + signed_value
                existing["entryCount"] = int(_number(existing.get("entryCount"))) + entry_count
            else:
                grouped[group_key] = _evidence_component(
                    row=row,
                    label=label,
                    signed_value=signed_value,
                    account_number=primary_account,
                    account_name=primary_name,
                    counter_account_number=counter_account,
                    counter_account_name=counter_name,
                    entry_count=entry_count,
                )
        if grouped:
            return list(grouped.values())
        if needs_detail or operator == "MOVEMENT_DETAIL":
            return []

    period_basis = "" if term.get("useColumnBasis") is False else _column_period_basis(column_payload)
    if operator == "DELTA_BALANCE":
        period_basis = "DELTA"
    basis = normalize_key(period_basis or term.get("periodBasis") or ("PERIOD" if operator == "MOVEMENT" else "ENDING"))
    side = normalize_key(term.get("side") or "NET")
    period_labels = {
        "OPENING": "đầu kỳ", "PERIOD": "trong kỳ", "YTD": "lũy kế từ đầu năm",
        "ENDING": "cuối kỳ", "DELTA": "tăng/giảm giữa đầu kỳ và cuối kỳ",
    }
    side_labels = {
        "DEBIT": "Nợ", "CREDIT": "Có", "DEBIT_BALANCE": "ròng bên Nợ",
        "CREDIT_BALANCE": "ròng bên Có", "NET": "Nợ trừ Có",
    }
    components: list[dict[str, Any]] = []
    for row in _rows_for_set(source_rows, set_key, members, period_to):
        value = sign * _metric_value(row, term.get("metricKey"), term, period_basis)
        account_number = _text(_field(row, "account_number"))
        account_name = _evidence_account_name(
            account_number, _field(row, "company_key"), row, account_names
        )
        label = f"{side_labels.get(side, side or 'Số')} {period_labels.get(basis, basis.lower())} TK {account_number}{' - ' + account_name if account_name else ''}"
        components.append(
            _evidence_component(
                row=row,
                label=label,
                signed_value=value,
                account_number=account_number,
                account_name=account_name,
            )
        )
    return components


# Mã dòng pháp lý có thể bắt đầu bằng số rồi kết thúc bằng chữ hoặc phần định
# danh, ví dụ 411a, 420b, 270_KHAC. resolve_line được ưu tiên trước nên số
# thuần vẫn có thể dùng làm hằng số khi nó không trùng mã dòng.
_TOKEN_RE = re.compile(r"[A-Za-z0-9_À-ỹ][A-Za-z0-9_.À-ỹ:]*|[()+\-*/]")


def evaluate_expression(expression: Any, resolve_line: Callable[[str], float | None]) -> float:
    source = re.sub(r"^\s*=\s*", "", str(expression or "")).strip()
    if not source:
        return 0.0
    tokens = _TOKEN_RE.findall(source)
    compact = re.sub(r"\s+", "", source)
    if "".join(tokens) != compact:
        raise ReportEngineError("Công thức có ký tự chưa được hỗ trợ.")
    position = 0

    def primary() -> float:
        nonlocal position
        if position >= len(tokens):
            raise ReportEngineError("Công thức kết thúc chưa đầy đủ.")
        token = tokens[position]
        position += 1
        if token == "(":
            nested = add_subtract()
            if position >= len(tokens) or tokens[position] != ")":
                raise ReportEngineError("Thiếu dấu đóng ngoặc trong công thức.")
            position += 1
            return nested
        if token == "+":
            return primary()
        if token == "-":
            return -primary()
        resolved = resolve_line(token)
        if resolved is not None:
            return _number(resolved)
        if re.fullmatch(r"\d+(?:\.\d+)?", token):
            return float(token)
        raise ReportEngineError(f"Không tìm thấy dòng tham chiếu “{token}”.")

    def multiply_divide() -> float:
        nonlocal position
        value = primary()
        while position < len(tokens) and tokens[position] in ("*", "/"):
            operator = tokens[position]
            position += 1
            right = primary()
            value = value * right if operator == "*" else (0.0 if right == 0 else value / right)
        return value

    def add_subtract() -> float:
        nonlocal position
        value = multiply_divide()
        while position < len(tokens) and tokens[position] in ("+", "-"):
            operator = tokens[position]
            position += 1
            right = multiply_divide()
            value = value + right if operator == "+" else value - right
        return value

    result = add_subtract()
    if position != len(tokens):
        raise ReportEngineError("Công thức có ký tự chưa được hỗ trợ.")
    return result


def decompose_linear_expression(
    expression: Any, is_line_reference: Callable[[str], bool]
) -> tuple[float, dict[str, float]]:
    """Tách công thức tuyến tính thành hằng số và hệ số của từng dòng.

    Hàm này chỉ phục vụ đường kiểm toán. Công thức phi tuyến giữa hai dòng
    (ví dụ 10*20) vẫn được tính bởi ``evaluate_expression`` nhưng không được
    giả vờ là đã chứng minh tới dữ liệu nguồn.
    """

    source = re.sub(r"^\s*=\s*", "", str(expression or "")).strip()
    if not source:
        return 0.0, {}
    tokens = _TOKEN_RE.findall(source)
    compact = re.sub(r"\s+", "", source)
    if "".join(tokens) != compact:
        raise ReportEngineError("Công thức có ký tự chưa được hỗ trợ.")
    position = 0

    def scaled(
        value: tuple[float, dict[str, float]], factor: float
    ) -> tuple[float, dict[str, float]]:
        constant, weights = value
        return constant * factor, {
            key: coefficient * factor for key, coefficient in weights.items()
        }

    def combined(
        left: tuple[float, dict[str, float]],
        right: tuple[float, dict[str, float]],
        factor: float,
    ) -> tuple[float, dict[str, float]]:
        constant = left[0] + factor * right[0]
        weights = dict(left[1])
        for key, coefficient in right[1].items():
            weights[key] = weights.get(key, 0.0) + factor * coefficient
            if abs(weights[key]) < 0.0000001:
                del weights[key]
        return constant, weights

    def primary() -> tuple[float, dict[str, float]]:
        nonlocal position
        if position >= len(tokens):
            raise ReportEngineError("Công thức kết thúc chưa đầy đủ.")
        token = tokens[position]
        position += 1
        if token == "(":
            nested = add_subtract()
            if position >= len(tokens) or tokens[position] != ")":
                raise ReportEngineError("Thiếu dấu đóng ngoặc trong công thức.")
            position += 1
            return nested
        if token == "+":
            return primary()
        if token == "-":
            return scaled(primary(), -1)
        if is_line_reference(token):
            return 0.0, {token: 1.0}
        if re.fullmatch(r"\d+(?:\.\d+)?", token):
            return float(token), {}
        raise ReportEngineError(f"Không tìm thấy dòng tham chiếu “{token}”.")

    def multiply_divide() -> tuple[float, dict[str, float]]:
        nonlocal position
        value = primary()
        while position < len(tokens) and tokens[position] in ("*", "/"):
            operator = tokens[position]
            position += 1
            right = primary()
            if operator == "*":
                if value[1] and right[1]:
                    raise ReportEngineError(
                        "Công thức nhân hai dòng chưa thể bung thành số liệu nguồn."
                    )
                value = scaled(value, right[0]) if not right[1] else scaled(right, value[0])
            else:
                if right[1] or abs(right[0]) < 0.0000001:
                    raise ReportEngineError(
                        "Công thức chia cho một dòng hoặc cho 0 chưa thể bung thành số liệu nguồn."
                    )
                value = scaled(value, 1 / right[0])
        return value

    def add_subtract() -> tuple[float, dict[str, float]]:
        nonlocal position
        value = multiply_divide()
        while position < len(tokens) and tokens[position] in ("+", "-"):
            operator = tokens[position]
            position += 1
            value = combined(value, multiply_divide(), 1 if operator == "+" else -1)
        return value

    result = add_subtract()
    if position != len(tokens):
        raise ReportEngineError("Công thức có ký tự chưa được hỗ trợ.")
    return result


def _payload_map(records: Iterable[dict[str, Any]]) -> dict[str, dict[str, Any]]:
    return {normalize_key(record.get("entityKey")): record.get("payload") or {} for record in records}


def _style(
    row_payload: dict[str, Any],
    cell_payload: dict[str, Any],
    styles: dict[str, dict[str, Any]],
    column_payload: dict[str, Any],
) -> dict[str, Any]:
    result = deepcopy(styles.get(normalize_key(row_payload.get("styleKey")), {}))
    for key in (
        "bold",
        "italic",
        "uppercase",
        "align",
        "numberFormat",
        "indent",
        "fontSize",
        "textColor",
        "fillColor",
        "borderPreset",
        "borderWidth",
    ):
        if row_payload.get(key) not in (None, ""):
            result[key] = row_payload[key]
        if cell_payload.get(key) not in (None, ""):
            result[key] = cell_payload[key]
    result.setdefault("align", column_payload.get("align"))
    result.setdefault("numberFormat", column_payload.get("numberFormat"))
    return result


def _header_rows(columns: list[dict[str, Any]]) -> list[dict[str, Any]]:
    if not any(_text(column.get("headerGroup")) for column in columns):
        return [{"cells": [{"label": _text(column.get("columnLabel")), "columnSpan": 1, "rowSpan": 1} for column in columns]}]
    top_cells: list[dict[str, Any]] = []
    bottom_cells: list[dict[str, Any]] = []
    index = 0
    while index < len(columns):
        column = columns[index]
        group = _text(column.get("headerGroup"))
        if not group:
            top_cells.append({"label": _text(column.get("columnLabel")), "columnSpan": 1, "rowSpan": 2})
            index += 1
            continue
        span = 1
        while index + span < len(columns) and _text(columns[index + span].get("headerGroup")) == group:
            span += 1
        top_cells.append({"label": group, "columnSpan": span, "rowSpan": 1})
        bottom_cells.extend(
            {"label": _text(columns[index + offset].get("columnLabel")), "columnSpan": 1, "rowSpan": 1}
            for offset in range(span)
        )
        index += span
    return [{"cells": top_cells}, {"cells": bottom_cells}]


def _data_check_output(
    period_from: date,
    period_to: date,
    rows: list[dict[str, Any]],
    selection: dict[str, list[str]],
) -> dict[str, Any]:
    definitions = [
        ("ACCOUNT_NUMBER", "Mã tài khoản", "LABEL", 14, ""),
        ("ACCOUNT_NAME", "Tên tài khoản", "LABEL", 34, ""),
        ("OPENING_DEBIT", "Nợ", "CURRENT_PERIOD", 15, "Đầu kỳ"),
        ("OPENING_CREDIT", "Có", "CURRENT_PERIOD", 15, "Đầu kỳ"),
        ("PERIOD_DEBIT", "Nợ", "CURRENT_PERIOD", 15, "Phát sinh"),
        ("PERIOD_CREDIT", "Có", "CURRENT_PERIOD", 15, "Phát sinh"),
        ("YTD_DEBIT", "Nợ", "CURRENT_PERIOD", 15, "Lũy kế"),
        ("YTD_CREDIT", "Có", "CURRENT_PERIOD", 15, "Lũy kế"),
        ("CLOSING_DEBIT", "Nợ", "CURRENT_PERIOD", 15, "Cuối kỳ"),
        ("CLOSING_CREDIT", "Có", "CURRENT_PERIOD", 15, "Cuối kỳ"),
    ]
    columns = [
        {
            "columnKey": key,
            "columnLabel": label,
            "columnType": column_type,
            "width": width,
            "headerGroup": group,
            "align": "LEFT" if column_type == "LABEL" else "RIGHT",
            "numberFormat": "#,##0;(#,##0);-",
        }
        for key, label, column_type, width, group in definitions
    ]
    output_rows = []
    for index, row in enumerate(rows):
        account_number = _text(_field(row, "account_number"))
        account_name = _text(_field(row, "account_name"))
        values: dict[str, Any] = {
            "ACCOUNT_NUMBER": account_number,
            "ACCOUNT_NAME": account_name,
            **{key: _number(_field(row, field)) for key, field in AMOUNT_FIELDS.items()},
        }
        bold = bool(_field(row, "is_bold")) or _number(row.get("grade")) <= 1
        cells = []
        for column in columns:
            is_label = column["columnKey"] in ("ACCOUNT_NUMBER", "ACCOUNT_NAME")
            cells.append(
                {
                    "columnKey": column["columnKey"],
                    "kind": "TEXT" if is_label else "VALUE",
                    "text": values[column["columnKey"]] if is_label else "",
                    "value": None if is_label else values[column["columnKey"]],
                    "style": {
                        "bold": bold,
                        "align": "LEFT" if is_label else "RIGHT",
                        "numberFormat": "#,##0;(#,##0);-",
                    },
                    "mergeColumnSpan": 1,
                    "mergeRowSpan": 1,
                }
            )
        output_rows.append(
            {
                "rowCode": account_number or str(index + 1),
                "rowLabel": account_name,
                "rowType": "DATA",
                "cells": cells,
            }
        )
    return {
        "ok": True,
        "report": {
            "reportKey": DATA_CHECK_KEY,
            "reportName": "Bảng cân đối tài khoản - dữ liệu đã đồng bộ",
            "displayTitle": "BẢNG CÂN ĐỐI TÀI KHOẢN",
            "hideColumnHeader": False,
        },
        "columns": columns,
        "headerRows": _header_rows(columns),
        "rows": output_rows,
        "warnings": [],
        "balanceRowCount": len(rows),
        "nonZeroCellCount": sum(
            1 for row in output_rows for cell in row["cells"] if cell["value"] is not None and cell["value"] != 0
        ),
        "periodFrom": period_from.isoformat(),
        "periodTo": period_to.isoformat(),
        "selection": selection,
    }


def _note_account_prefixes(value: Any) -> list[str]:
    return list(dict.fromkeys(re.findall(r"\d{2,}", str(value or ""))))


def _note_rows_for_accounts(
    balance_rows: list[dict[str, Any]],
    account_number: Any,
) -> list[dict[str, Any]]:
    prefixes = _note_account_prefixes(account_number)
    if not prefixes:
        return []
    matched = [
        row
        for row in balance_rows
        if _field(row, "is_leaf") is not False
        and any(_text(_field(row, "account_number")).startswith(prefix) for prefix in prefixes)
    ]
    if matched:
        return matched
    return [
        row
        for row in balance_rows
        if any(_text(_field(row, "account_number")) == prefix for prefix in prefixes)
    ]


def _note_term_metric(operand: Any) -> tuple[str, int, str] | None:
    normalized = normalize_key(operand)
    offset = -1 if "NAMTRUOC" in normalized else 0
    if "DOITUONG" in normalized or "BAN_BDSDT" in normalized or "THANHLYTSCD" in normalized:
        return None
    supported_details = ("", "_NAMTRUOC", "_CHITIETTHEOTK")
    mapping = (
        ("DUNODK", "opening_debit", "opening"),
        ("DUCODK", "opening_credit", "opening"),
        ("DUNOCK", "closing_debit", "closing"),
        ("DUCOCK", "closing_credit", "closing"),
        ("PHATSINHNO", "period_debit", "movement"),
        ("PHATSINHCO", "period_credit", "movement"),
    )
    for prefix, metric, bucket in mapping:
        if normalized == prefix or any(normalized == prefix + suffix for suffix in supported_details[1:]):
            return metric, offset, bucket
    return None


def _note_term_value(
    term: dict[str, Any],
    rows_by_offset: dict[int, list[dict[str, Any]]],
    warning: Callable[[str], None],
    item_code: str,
) -> tuple[str, float] | None:
    operand = term.get("operand")
    metric = _note_term_metric(operand)
    if not metric:
        warning(
            f"Thuyết minh {item_code or '(không mã)'} có thành phần “{_text(operand)}” "
            "cần dữ liệu chi tiết/điều chỉnh chưa có trong khối cân đối tài khoản."
        )
        return None
    if _text(term.get("co_account_number")):
        warning(
            f"Thuyết minh {item_code or '(không mã)'} cần tài khoản đối ứng; "
            "khối cân đối hiện tại chưa đủ dữ liệu để tính thành phần này."
        )
        return None
    metric_name, offset, bucket = metric
    rows = _note_rows_for_accounts(
        rows_by_offset.get(offset, []),
        term.get("account_number"),
    )
    value = sum(_number(_field(row, metric_name)) for row in rows)
    sign = -1 if _number(term.get("operation_sign")) < 0 else 1
    return bucket, sign * value


def _build_note_output(
    *,
    records: list[dict[str, Any]],
    report_key: str,
    note_items: list[dict[str, Any]],
    selected_note_item_ids: Iterable[Any],
    rows_by_offset: dict[int, list[dict[str, Any]]],
    warning: Callable[[str], None],
) -> dict[str, Any]:
    required_ids, selected_ids = _selected_note_ids(
        records,
        report_key,
        selected_note_item_ids,
    )
    item_by_id = {
        _text(item.get("itemId")).lower(): item
        for item in note_items
        if _text(item.get("itemId"))
    }
    for item_id in selected_ids:
        if item_id not in item_by_id:
            warning(
                f"Không tìm thấy dòng thuyết minh {item_id} trong danh mục MISA đã lưu."
            )
    selected_set = set(selected_ids)
    ordered = sorted(
        [
            item
            for item in note_items
            if _text(item.get("itemId")).lower() in selected_set
        ],
        key=lambda item: (
            _text(item.get("partInTab")),
            _text(item.get("part")),
            _number(item.get("itemIndex")),
            _text(item.get("itemCode")),
        ),
    )
    rows: list[dict[str, Any]] = []
    for item in ordered:
        item_id = _text(item.get("itemId")).lower()
        formula = item.get("formula") if isinstance(item.get("formula"), list) else []
        values: dict[str, float | None] = {
            "openingValue": None,
            "movementValue": None,
            "closingValue": None,
        }
        supported_count = 0
        for term in formula:
            if not isinstance(term, dict):
                continue
            evaluated = _note_term_value(
                term,
                rows_by_offset,
                warning,
                _text(item.get("itemCode")),
            )
            if not evaluated:
                continue
            bucket, value = evaluated
            field = {
                "opening": "openingValue",
                "movement": "movementValue",
                "closing": "closingValue",
            }[bucket]
            values[field] = _number(values[field]) + value
            supported_count += 1
        rows.append(
            {
                "itemId": item_id,
                "itemCode": item.get("itemCode") or "",
                "itemName": item.get("itemName") or "",
                "part": item.get("part") or "",
                "partInTab": item.get("partInTab") or "",
                "isBold": bool(item.get("isBold")),
                "isItalic": bool(item.get("isItalic")),
                "required": item_id in required_ids,
                "formulaTermCount": len(formula),
                "supportedTermCount": supported_count,
                **values,
            }
        )
    return {
        "title": "THUYẾT MINH BÁO CÁO TÀI CHÍNH",
        "columns": [
            {"key": "itemCode", "label": "Mã", "width": 14},
            {"key": "itemName", "label": "Nội dung thuyết minh", "width": 52},
            {"key": "openingValue", "label": "Đầu kỳ", "width": 16},
            {"key": "movementValue", "label": "Phát sinh", "width": 16},
            {"key": "closingValue", "label": "Cuối kỳ", "width": 16},
        ],
        "rows": rows,
        "requiredItemIds": required_ids,
        "selectedItemIds": selected_ids,
        "optionalItemIds": [
            item_id for item_id in selected_ids if item_id not in required_ids
        ],
    }


def build_output(
    *,
    records: list[dict[str, Any]],
    report_key: str,
    period_from: date,
    period_to: date,
    rows_by_offset: dict[int, list[dict[str, Any]]],
    selection: dict[str, list[str]],
    note_items: list[dict[str, Any]] | None = None,
    selected_note_item_ids: Iterable[Any] = (),
    cross_report_resolver: Callable[
        [str, str, dict[str, Any], dict[str, Any]], float | None
    ]
    | None = None,
    include_evidence: bool = False,
    cell_adjustments: list[dict[str, Any]] | None = None,
) -> dict[str, Any]:
    normalized_report = normalize_key(report_key or DATA_CHECK_KEY)
    if normalized_report == DATA_CHECK_KEY:
        return _data_check_output(period_from, period_to, rows_by_offset.get(0, []), selection)
    report_record = next(
        (
            record
            for record in records
            if record.get("recordType") == "REPORT" and normalize_key(record.get("entityKey")) == normalized_report
        ),
        None,
    )
    if not report_record:
        raise ReportEngineError("Mẫu báo cáo không còn tồn tại trong phiên bản đang chọn.")
    columns = sorted(
        [
            record
            for record in records
            if record.get("recordType") == "REPORT_COLUMN"
            and normalize_key((record.get("payload") or {}).get("reportKey")) == normalized_report
        ],
        key=lambda record: _number(record.get("sortOrder")),
    )
    row_records = sorted(
        [
            record
            for record in records
            if record.get("recordType") == "REPORT_ROW"
            and normalize_key((record.get("payload") or {}).get("reportKey")) == normalized_report
        ],
        key=lambda record: _number(record.get("sortOrder")),
    )
    if not columns:
        raise ReportEngineError("Mẫu báo cáo chưa có cột.")
    if not row_records:
        raise ReportEngineError("Mẫu báo cáo chưa có dòng.")
    rules = _payload_map(record for record in records if record.get("recordType") == "RULE")
    styles = _payload_map(record for record in records if record.get("recordType") == "STYLE")
    members = [record for record in records if record.get("recordType") == "ACCOUNT_SET_MEMBER"]
    field_mappings = [record for record in records if record.get("recordType") == "SOURCE_FIELD_MAPPING"]
    account_names: dict[tuple[str, str], str] = {}
    for account_record in records:
        if account_record.get("recordType") != "ACCOUNT":
            continue
        payload = account_record.get("payload") or {}
        account_number = _text(
            payload.get("accountNumber")
            or payload.get("accountCode")
            or account_record.get("displayCode")
            or account_record.get("entityKey")
        )
        if not account_number:
            continue
        company_key = normalize_key(payload.get("companyKey") or account_record.get("companyKey"))
        account_names[(company_key, account_number)] = _text(
            payload.get("accountName") or account_record.get("displayName")
        )
    warnings: list[str] = []
    warning_seen: set[str] = set()
    blocking_issues: list[str] = []
    blocking_issue_seen: set[str] = set()
    note_items = note_items or []
    _, included_note_ids = _selected_note_ids(
        records,
        normalized_report,
        selected_note_item_ids,
    )
    included_note_set = set(included_note_ids)
    note_item_by_id = {
        _text(item.get("itemId")).lower(): item
        for item in note_items
        if _text(item.get("itemId"))
    }

    def warning(message: str) -> None:
        cleaned = _text(message)
        if cleaned and cleaned not in warning_seen:
            warning_seen.add(cleaned)
            warnings.append(cleaned)

    def blocking_issue(message: str) -> None:
        cleaned = _text(message)
        if cleaned and cleaned not in blocking_issue_seen:
            blocking_issue_seen.add(cleaned)
            blocking_issues.append(cleaned)
        warning(cleaned)

    for offset in required_offsets(records, normalized_report):
        if offset and not rows_by_offset.get(offset):
            warning(
                "Chưa có snapshot cho cột kỳ trước "
                f"({shift_year(period_from, offset).isoformat()} → {shift_year(period_to, offset).isoformat()}), "
                "hệ thống tạm hiển thị 0."
            )
    row_by_code = {
        _text((record.get("payload") or {}).get("rowCode") or record.get("displayCode")): record
        for record in row_records
    }
    value_cache: dict[str, float] = {}
    stack: set[str] = set()
    # Các dòng/cột thường dùng cùng một nguồn và danh mục. Giữ cache trong một
    # lần dựng báo cáo để không quét lại hàng chục nghìn bút toán cho mỗi ô.
    term_runtime_cache: dict[str, Any] = {}
    normalized_adjustments: list[dict[str, Any]] = []
    for raw_adjustment in cell_adjustments or []:
        row_code = _text(raw_adjustment.get("rowCode") or raw_adjustment.get("row_code")).upper()
        if not row_code:
            continue
        normalized_adjustments.append(
            {
                **raw_adjustment,
                "rowCode": row_code,
                "columnKey": normalize_key(
                    raw_adjustment.get("columnKey") or raw_adjustment.get("column_key")
                ),
                "amount": _number(raw_adjustment.get("amount")),
            }
        )

    def adjustments_for_cell(
        row_payload: dict[str, Any], column_payload: dict[str, Any]
    ) -> list[dict[str, Any]]:
        row_code = _text(row_payload.get("rowCode")).upper()
        column_key = normalize_key(column_payload.get("columnKey"))
        column_type = normalize_key(column_payload.get("columnType"))
        offset = column_period_offset(column_payload)
        return [
            adjustment
            for adjustment in normalized_adjustments
            if adjustment["rowCode"] == row_code
            and (
                adjustment["columnKey"] == column_key
                or (
                    not adjustment["columnKey"]
                    and offset == 0
                    and column_type == "CURRENT_PERIOD"
                )
            )
        ]

    def contextual_term(
        term: dict[str, Any],
        row_payload: dict[str, Any],
        column_payload: dict[str, Any],
    ) -> dict[str, Any]:
        return {
            **term,
            "_reportKey": normalized_report,
            "_rowCode": _text(row_payload.get("rowCode")),
            "_columnKey": _text(column_payload.get("columnKey")),
        }

    def cell_and_rule(
        record: dict[str, Any], column_record: dict[str, Any]
    ) -> tuple[dict[str, Any], dict[str, Any]]:
        row_payload = record.get("payload") or {}
        column_payload = column_record.get("payload") or {}
        cells = row_payload.get("cells") if isinstance(row_payload.get("cells"), dict) else {}
        column_key = normalize_key(column_payload.get("columnKey"))
        cell = cells.get(column_key) or cells.get(column_payload.get("columnKey")) or {}
        linked_rule = rules.get(normalize_key(row_payload.get("ruleKey")), {})
        has_inline_rule = any(
            cell.get(key) not in (None, "", [])
            for key in (
                "formulaRequired",
                "contentKind",
                "accountSetKey",
                "operator",
                "expression",
                "terms",
            )
        )
        return cell, cell if has_inline_rule else linked_rule

    def rule_has_usable_formula(rule: dict[str, Any]) -> bool:
        terms = rule.get("terms") if isinstance(rule.get("terms"), list) else []
        if terms:
            return any(rule_has_usable_formula(term) for term in terms)
        operator = normalize_key(rule.get("operator"))
        if operator == "MANUAL_ADJUSTMENT":
            # Đây là một công thức hợp lệ dù kỳ hiện tại chưa có khoản điều chỉnh
            # được duyệt. Khi đó kết quả đúng là 0, không phải thiếu công thức.
            return True
        if operator == "CROSS_REPORT":
            return bool(_text(rule.get("targetReportKey")) and _text(rule.get("targetRowCode")))
        if operator in ("LINE_REFERENCE", "REPORT_LINE"):
            return bool(_text(rule.get("targetRowCode") or rule.get("lineRef") or rule.get("rowCode")))
        if operator == "LINE_EXPRESSION":
            return bool(_text(rule.get("expression")))
        if _text(rule.get("accountSetKey")):
            return True
        if _text(rule.get("expression")):
            return True
        line_refs = rule.get("lineRefs") or []
        if not isinstance(line_refs, list):
            line_refs = [
                item.strip()
                for item in re.split(r"[\n,;]+", str(line_refs))
                if item.strip()
            ]
        return bool(
            line_refs
            and operator in ("SUM_LINES", "SUBTRACT_LINES", "LINE_EXPRESSION")
        )

    def cell_is_missing_formula(
        record: dict[str, Any], column_record: dict[str, Any]
    ) -> bool:
        row_payload = record.get("payload") or {}
        if normalize_key(row_payload.get("rowType")) not in (
            "DATA",
            "FORMULA",
            "SUBTOTAL",
            "TOTAL",
            "CROSS_REPORT",
            "MANUAL_ADJUSTMENT",
        ):
            return False
        cell, rule = cell_and_rule(record, column_record)
        return cell.get("formulaRequired") is True or not rule_has_usable_formula(rule)

    def evaluate_formula_term(
        term: dict[str, Any],
        column_record: dict[str, Any],
        row_payload: dict[str, Any] | None = None,
    ) -> float:
        column_payload = column_record.get("payload") or {}
        runtime_term = contextual_term(term, row_payload or {}, column_payload)
        operator = normalize_key(term.get("operator") or "BALANCE")
        sign = -1 if _number(term.get("sign")) < 0 else 1
        if operator == "CROSS_REPORT":
            target_report = normalize_key(term.get("targetReportKey"))
            target_row = _text(term.get("targetRowCode"))
            if not target_report or not target_row:
                warning("Còn thành phần liên kết báo cáo chưa chọn đủ báo cáo và dòng nguồn.")
                return 0.0
            if cross_report_resolver is None:
                warning(f"Chưa thể lấy dòng {target_report}:{target_row} từ báo cáo liên kết.")
                return 0.0
            resolved = cross_report_resolver(target_report, target_row, column_payload, term)
            if resolved is None:
                warning(f"Không tìm thấy giá trị của dòng {target_report}:{target_row}.")
                return 0.0
            return sign * _number(resolved)
        if operator in ("LINE_REFERENCE", "REPORT_LINE"):
            target_row = _text(
                term.get("targetRowCode") or term.get("lineRef") or term.get("rowCode")
            ).split(":")[-1]
            reference = row_by_code.get(target_row)
            if not reference:
                warning(f"Không tìm thấy dòng tham chiếu {target_row or '(chưa chọn)' }.")
                return 0.0
            return sign * evaluate_row(reference, column_record)
        if operator == "LINE_EXPRESSION":
            expression = _text(term.get("expression"))
            return sign * evaluate_expression(
                expression,
                lambda token: evaluate_row(row_by_code[token], column_record)
                if token in row_by_code
                else None,
            )
        return _term_value(
            runtime_term,
            rows_by_offset.get(column_period_offset(column_payload), []),
            members,
            field_mappings,
            shift_year(period_to, column_period_offset(column_payload)),
            column_payload,
            warning,
            term_runtime_cache,
            blocking_issue,
        )

    def evaluate_row(record: dict[str, Any], column_record: dict[str, Any]) -> float:
        row_payload = record.get("payload") or {}
        column_payload = column_record.get("payload") or {}
        cache_key = f"{record.get('recordId')}::{column_record.get('recordId')}"
        if cache_key in value_cache:
            return value_cache[cache_key]
        if cache_key in stack:
            warning(f"Phát hiện công thức tham chiếu vòng tại dòng {row_payload.get('rowCode') or record.get('displayCode')}.")
            return 0.0
        stack.add(cache_key)
        cell, rule = cell_and_rule(record, column_record)
        value = 0.0
        try:
            if cell.get("formulaRequired"):
                warning(f"Dòng {row_payload.get('rowCode') or record.get('displayCode')} còn ô chưa thiết lập công thức.")
            elif isinstance(rule.get("terms"), list) and rule["terms"]:
                value = sum(
                    evaluate_formula_term(term, column_record, row_payload)
                    for term in rule["terms"]
                )
            elif normalize_key(rule.get("operator")) == "CROSS_REPORT":
                target_report = normalize_key(rule.get("targetReportKey"))
                target_row = _text(rule.get("targetRowCode"))
                if not target_report or not target_row:
                    warning(
                        f"Dòng {row_payload.get('rowCode') or record.get('displayCode')} chưa chọn báo cáo và dòng nguồn."
                    )
                elif cross_report_resolver is None:
                    warning(f"Chưa thể lấy dòng {target_report}:{target_row} từ báo cáo liên kết.")
                else:
                    resolved = cross_report_resolver(target_report, target_row, column_payload, rule)
                    if resolved is None:
                        warning(f"Không tìm thấy giá trị của dòng {target_report}:{target_row}.")
                    else:
                        value = (-1 if _number(rule.get("sign")) < 0 else 1) * _number(resolved)
            elif normalize_key(rule.get("operator")) == "MANUAL_ADJUSTMENT":
                offset = column_period_offset(column_payload)
                value = _term_value(
                    contextual_term(rule, row_payload, column_payload),
                    rows_by_offset.get(offset, []),
                    members,
                    field_mappings,
                    shift_year(period_to, offset),
                    column_payload,
                    warning,
                    term_runtime_cache,
                    blocking_issue,
                )
            elif rule.get("accountSetKey"):
                offset = column_period_offset(column_payload)
                value = _term_value(
                    contextual_term(rule, row_payload, column_payload),
                    rows_by_offset.get(offset, []),
                    members,
                    field_mappings,
                    shift_year(period_to, offset),
                    column_payload,
                    warning,
                    term_runtime_cache,
                    blocking_issue,
                )
            elif _text(rule.get("expression")):
                value = evaluate_expression(
                    rule["expression"],
                    lambda token: evaluate_row(row_by_code[token], column_record) if token in row_by_code else None,
                )
            else:
                line_refs = rule.get("lineRefs") or []
                if not isinstance(line_refs, list):
                    line_refs = [item.strip() for item in re.split(r"[\n,;]+", str(line_refs)) if item.strip()]
                operator = normalize_key(rule.get("operator"))
                if line_refs and operator in ("SUM_LINES", "SUBTRACT_LINES", "LINE_EXPRESSION"):
                    line_values = []
                    for line_ref in line_refs:
                        local_code = str(line_ref).split(":")[-1]
                        reference = row_by_code.get(local_code)
                        if not reference:
                            warning(f"Không tìm thấy dòng tham chiếu {line_ref}.")
                            line_values.append(0.0)
                        else:
                            line_values.append(evaluate_row(reference, column_record))
                    if operator == "SUBTRACT_LINES":
                        value = line_values[0] - sum(line_values[1:]) if line_values else 0.0
                    else:
                        value = sum(line_values)
                elif normalize_key(row_payload.get("rowType")) in (
                    "DATA",
                    "FORMULA",
                    "SUBTOTAL",
                    "TOTAL",
                    "CROSS_REPORT",
                ):
                    warning(f"Dòng {row_payload.get('rowCode') or record.get('displayCode')} chưa có công thức dùng được.")
        except ReportEngineError as error:
            warning(f"Công thức dòng {row_payload.get('rowCode') or record.get('displayCode')}: {error}")
        finally:
            stack.discard(cache_key)
        value += sum(
            adjustment["amount"]
            for adjustment in adjustments_for_cell(row_payload, column_payload)
        )
        value_cache[cache_key] = value if math.isfinite(value) else 0.0
        return value_cache[cache_key]

    def evidence_rule_for_cell(
        record: dict[str, Any], column_record: dict[str, Any]
    ) -> dict[str, Any]:
        return cell_and_rule(record, column_record)[1]

    evidence_cache: dict[str, dict[str, Any]] = {}
    evidence_stack: set[str] = set()

    def row_evidence_title(record: dict[str, Any]) -> str:
        payload = record.get("payload") or {}
        code = _text(payload.get("rowCode") or record.get("displayCode"))
        label = _text(payload.get("rowLabel") or record.get("displayName"))
        return f"Dòng {code} - {label}" if label else f"Dòng {code}"

    def reference_evidence_component(
        label: str,
        signed_value: float,
        *,
        formula_sign: float | None = None,
        source_name: str = "Cấu hình mẫu báo cáo",
        trace_path: str = "",
    ) -> dict[str, Any]:
        return _evidence_component(
            row={},
            label=label,
            signed_value=signed_value,
            formula_sign=formula_sign,
            source_name=source_name,
            trace_path=trace_path or label,
        )

    def scaled_component(
        component: dict[str, Any], factor: float, prefix: str = ""
    ) -> dict[str, Any]:
        result = deepcopy(component)
        result["signedValue"] = factor * _number(component.get("signedValue"))
        component_sign = -1 if _number(component.get("formulaSign")) < 0 else 1
        result["formulaSign"] = component_sign * (-1 if factor < 0 else 1)
        child_path = _text(component.get("tracePath") or component.get("label"))
        result["tracePath"] = f"{prefix} → {child_path}" if prefix else child_path
        return result

    def build_cell_evidence_internal(
        record: dict[str, Any], column_record: dict[str, Any]
    ) -> dict[str, Any]:
        cache_key = f"{record.get('recordId')}::{column_record.get('recordId')}"
        if cache_key in evidence_cache:
            return deepcopy(evidence_cache[cache_key])
        value = evaluate_row(record, column_record)
        current_title = row_evidence_title(record)
        if cell_is_missing_formula(record, column_record):
            result = {
                "value": None,
                "components": [],
                "isComplete": False,
                "difference": None,
                "warning": f"{current_title} chưa được thiết lập công thức.",
            }
            evidence_cache[cache_key] = deepcopy(result)
            return result
        if cache_key in evidence_stack:
            return {
                "value": value,
                "components": [],
                "isComplete": False,
                "difference": value,
                "warning": f"Không thể truy nguồn vì {current_title.lower()} tham chiếu vòng.",
            }

        evidence_stack.add(cache_key)
        rule = evidence_rule_for_cell(record, column_record)
        column_payload = column_record.get("payload") or {}
        offset = column_period_offset(column_payload)
        source_rows = rows_by_offset.get(offset, [])
        effective_period_to = shift_year(period_to, offset)
        components: list[dict[str, Any]] = []
        issue_messages: list[str] = []

        def issue(message: str) -> None:
            cleaned = _text(message)
            if cleaned and cleaned not in issue_messages:
                issue_messages.append(cleaned)

        for adjustment in adjustments_for_cell(record.get("payload") or {}, column_payload):
            label = _text(adjustment.get("reason")) or "Bút toán điều chỉnh khi xuất báo cáo"
            source_reference = _text(adjustment.get("sourceReference") or adjustment.get("source_reference"))
            if source_reference:
                label += f" · {source_reference}"
            components.append(
                reference_evidence_component(
                    label,
                    adjustment["amount"],
                    source_name="Bút toán điều chỉnh được chọn",
                    trace_path=f"{current_title} → {label}",
                )
            )

        def account_set_criteria(account_set_key: Any) -> str:
            set_key = normalize_key(account_set_key)
            relevant = [
                record.get("payload") or {}
                for record in members
                if normalize_key(
                    (record.get("payload") or {}).get("accountSetKey")
                    or (record.get("payload") or {}).get("setKey")
                )
                == set_key
                and normalize_key((record.get("payload") or {}).get("status") or "ACTIVE")
                != "INACTIVE"
            ]
            included: list[str] = []
            excluded: list[str] = []
            for payload in relevant:
                valid_from = _text(payload.get("validFrom"))
                valid_to = _text(payload.get("validTo"))
                if valid_from and valid_from > effective_period_to.isoformat():
                    continue
                if valid_to and valid_to < effective_period_to.isoformat():
                    continue
                raw_value = _text(payload.get("matchValue") or payload.get("accountCode"))
                match_type = normalize_key(payload.get("matchType") or "EXACT")
                raw_parts = [
                    item.strip()
                    for item in re.split(r"[,;\n]+", raw_value)
                    if item.strip()
                ]
                labels: list[str] = []
                for raw_part in raw_parts:
                    value = re.sub(r"^TK\s*", "", raw_part, flags=re.IGNORECASE).strip()
                    if not value:
                        continue
                    if match_type == "PREFIX" and not value.endswith("*"):
                        value += "*"
                    company_key = _text(payload.get("companyKey"))
                    labels.append(
                        f"TK {value}{f' ({company_key})' if company_key else ''}"
                    )
                target = excluded if normalize_key(payload.get("role")) == "EXCLUDE" else included
                for label in labels:
                    if label not in target:
                        target.append(label)
            if not included and not excluded:
                return f"nhóm tài khoản {set_key or 'chưa chọn'}"
            parts = []
            if included:
                parts.append("gồm " + ", ".join(included))
            if excluded:
                parts.append("loại trừ " + ", ".join(excluded))
            return "; ".join(parts)

        def configured_zero_component(term: dict[str, Any]) -> dict[str, Any]:
            operator = normalize_key(term.get("operator") or rule.get("operator") or "BALANCE")
            sign = -1.0 if _number(term.get("sign") or rule.get("sign")) < 0 else 1.0
            set_key = term.get("accountSetKey") or rule.get("accountSetKey")
            criteria = account_set_criteria(set_key)
            counter_key = term.get("counterAccountSetKey") or rule.get("counterAccountSetKey")
            counter_criteria = account_set_criteria(counter_key) if counter_key else ""
            period_basis = (
                ""
                if term.get("useColumnBasis") is False
                else _column_period_basis(column_payload)
            )
            if operator == "DELTA_BALANCE":
                period_basis = "DELTA"
            basis = normalize_key(
                period_basis
                or term.get("periodBasis")
                or rule.get("periodBasis")
                or ("PERIOD" if operator in ("MOVEMENT", "MOVEMENT_DETAIL", "CASH_MOVEMENT", "PROFIT_LOSS_ACTIVITY") else "ENDING")
            )
            side = normalize_key(term.get("side") or rule.get("side") or "NET")
            period_labels = {
                "OPENING": "đầu kỳ",
                "PERIOD": "trong kỳ",
                "YTD": "lũy kế từ đầu năm",
                "ENDING": "cuối kỳ",
                "DELTA": "chênh lệch đầu kỳ và cuối kỳ",
            }
            side_labels = {
                "DEBIT": "Nợ",
                "CREDIT": "Có",
                "DEBIT_BALANCE": "số dư Nợ",
                "DEBIT_NET": "số dư Nợ",
                "CREDIT_BALANCE": "số dư Có",
                "CREDIT_NET": "số dư Có",
                "NET": "Nợ trừ Có",
            }
            if operator == "MANUAL_ADJUSTMENT":
                action = "lấy khoản điều chỉnh kế toán đã lưu"
            elif operator == "PROFIT_LOSS_ACTIVITY":
                action = (
                    f"lấy phát sinh {side_labels.get(side, side or 'ròng')} {period_labels.get(basis, 'trong kỳ')} "
                    f"của {criteria}, ưu tiên bút toán kết chuyển"
                )
            elif operator in ("MOVEMENT", "MOVEMENT_DETAIL", "CASH_MOVEMENT"):
                action = (
                    f"lấy phát sinh {side_labels.get(side, side or 'ròng')} "
                    f"{period_labels.get(basis, 'trong kỳ')} của {criteria}"
                )
                if counter_criteria:
                    action += f", với tài khoản đối ứng {counter_criteria}"
            else:
                action = (
                    f"lấy {side_labels.get(side, side or 'số')} "
                    f"{period_labels.get(basis, basis.lower())} của {criteria}"
                )
            label = f"Công thức vẫn áp dụng: {action}; kỳ này không phát sinh"
            return reference_evidence_component(
                label,
                0.0,
                formula_sign=sign,
                source_name="Cấu hình công thức - kỳ này không phát sinh",
                trace_path=f"{current_title} → {label}",
            )

        def add_source_components(items: list[dict[str, Any]]) -> None:
            for component in items:
                components.append(scaled_component(component, 1, current_title))

        def add_data_term_components(term: dict[str, Any], term_value: float) -> None:
            runtime_term = contextual_term(term, row_payload, column_payload)
            term_components = _term_evidence_components(
                runtime_term,
                source_rows,
                members,
                field_mappings,
                effective_period_to,
                column_payload,
                account_names,
                term_value,
                term_runtime_cache,
            )
            non_zero_components = [
                component
                for component in term_components
                if abs(_number(component.get("signedValue"))) > 0.0000001
            ]
            if non_zero_components or (
                normalize_key(term.get("operator")) == "MANUAL_ADJUSTMENT"
                and term_components
            ):
                add_source_components(non_zero_components or term_components)
            elif abs(term_value) <= 0.0000001:
                components.append(configured_zero_component(term))
            else:
                issue(
                    "Công thức có giá trị nhưng chưa truy được dữ liệu nguồn chi tiết cho thành phần này."
                )

        def add_reference(target_code: str, factor: float) -> None:
            local_code = _text(target_code).split(":")[-1]
            target = row_by_code.get(local_code)
            if not target:
                issue(f"Không tìm thấy dòng {local_code or '(chưa chọn)'} để truy nguồn.")
                return
            child = build_cell_evidence_internal(target, column_record)
            if not child.get("isComplete"):
                issue(
                    _text(child.get("warning"))
                    or f"{row_evidence_title(target)} chưa truy đủ dữ liệu nguồn."
                )
            child_components = child.get("components") or []
            if not child_components and abs(_number(child.get("value"))) > 0.0001:
                issue(f"{row_evidence_title(target)} có số nhưng chưa có dữ liệu nguồn chi tiết.")
                return
            for component in child_components:
                components.append(scaled_component(component, factor, current_title))

        def add_linear_expression(expression: Any, factor: float = 1.0) -> None:
            try:
                constant, weights = decompose_linear_expression(
                    expression,
                    lambda token: token in row_by_code
                    or _text(token).split(":")[-1] in row_by_code,
                )
            except ReportEngineError as error:
                issue(f"Không thể bung công thức “{_text(expression)}” đến dữ liệu nguồn: {error}")
                return
            for token, coefficient in weights.items():
                add_reference(token, factor * coefficient)
            if abs(constant) > 0.0000001:
                label = f"Hằng số được khai báo trực tiếp trong công thức: {constant:g}"
                components.append(
                    reference_evidence_component(
                        label,
                        factor * constant,
                        trace_path=f"{current_title} → {label}",
                    )
                )

        def add_term(term: dict[str, Any]) -> None:
            operator = normalize_key(term.get("operator") or rule.get("operator") or "BALANCE")
            sign = -1.0 if _number(term.get("sign")) < 0 else 1.0
            if operator == "CROSS_REPORT":
                target_report = _text(term.get("targetReportKey") or rule.get("targetReportKey"))
                target_row = _text(term.get("targetRowCode") or rule.get("targetRowCode"))
                label = (
                    f"Dòng {target_row or 'chưa chọn'} của báo cáo "
                    f"{target_report or 'chưa chọn'}"
                )
                components.append(
                    reference_evidence_component(
                        label,
                        evaluate_formula_term(term, column_record, row_payload),
                        source_name="Báo cáo liên kết",
                        trace_path=f"{current_title} → {label}",
                    )
                )
                issue(
                    f"{label} mới có giá trị tổng; chưa kèm dữ liệu nguồn của báo cáo liên kết."
                )
                return
            if operator in ("LINE_REFERENCE", "REPORT_LINE"):
                add_reference(
                    _text(
                        term.get("targetRowCode")
                        or term.get("lineRef")
                        or term.get("rowCode")
                    ),
                    sign,
                )
                return
            if operator == "LINE_EXPRESSION":
                add_linear_expression(term.get("expression"), sign)
                return
            term_value = evaluate_formula_term(term, column_record, row_payload)
            add_data_term_components(term, term_value)

        try:
            terms = rule.get("terms") if isinstance(rule.get("terms"), list) else []
            if terms:
                for term in terms:
                    add_term(term)
            elif normalize_key(rule.get("operator")) in (
                "CROSS_REPORT",
                "LINE_REFERENCE",
                "REPORT_LINE",
                "LINE_EXPRESSION",
            ):
                add_term(rule)
            elif normalize_key(rule.get("operator")) == "MANUAL_ADJUSTMENT" or rule.get("accountSetKey"):
                term_value = _term_value(
                    contextual_term(rule, row_payload, column_payload),
                    source_rows,
                    members,
                    field_mappings,
                    effective_period_to,
                    column_payload,
                    warning,
                    term_runtime_cache,
                    blocking_issue,
                )
                add_data_term_components(rule, term_value)
            else:
                line_refs = rule.get("lineRefs") or []
                if not isinstance(line_refs, list):
                    line_refs = [
                        item.strip()
                        for item in re.split(r"[\n,;]+", str(line_refs))
                        if item.strip()
                    ]
                operator = normalize_key(rule.get("operator"))
                if line_refs and operator in ("SUM_LINES", "SUBTRACT_LINES"):
                    for index, line_ref in enumerate(line_refs):
                        add_reference(
                            str(line_ref),
                            -1.0 if operator == "SUBTRACT_LINES" and index else 1.0,
                        )
                elif _text(rule.get("expression")):
                    add_linear_expression(rule.get("expression"))
        finally:
            evidence_stack.discard(cache_key)

        component_total = sum(_number(item.get("signedValue")) for item in components)
        difference = value - component_total
        is_complete = not issue_messages and abs(difference) <= 0.0001
        if abs(difference) > 0.0001:
            issue(
                "Các nguồn đã truy được chưa khớp kết quả; "
                f"còn chênh lệch {difference:,.4f}. Hệ thống không tự tạo số bù."
            )
            is_complete = False
        if not components and abs(value) <= 0.0001 and not issue_messages:
            components.append(
                reference_evidence_component(
                    "Không phát sinh số liệu theo công thức đã chọn",
                    0.0,
                    trace_path=f"{current_title} → Không phát sinh số liệu",
                )
            )
        elif not components and abs(value) > 0.0001:
            issue("Ô báo cáo có số nhưng chưa truy được thành phần dữ liệu nguồn.")
            is_complete = False
        result = {
            "value": value,
            "components": components,
            "isComplete": is_complete,
            "difference": difference,
            "warning": " ".join(issue_messages),
        }
        evidence_cache[cache_key] = deepcopy(result)
        return result

    def build_cell_evidence(
        record: dict[str, Any], column_record: dict[str, Any]
    ) -> dict[str, Any]:
        return build_cell_evidence_internal(record, column_record)

    output_columns = []
    for record in columns:
        payload = record.get("payload") or {}
        column_type = payload.get("columnType") or "CURRENT_PERIOD"
        output_columns.append(
            {
                "recordId": record.get("recordId"),
                "columnKey": payload.get("columnKey"),
                "columnLabel": payload.get("columnLabel") or record.get("displayName"),
                "headerGroup": _text(payload.get("headerGroup")),
                "columnType": column_type,
                "width": _number(payload.get("width") or (44 if column_type == "LABEL" else 14)),
                "align": payload.get("align") or ("LEFT" if column_type == "LABEL" else "RIGHT"),
                "numberFormat": payload.get("numberFormat") or "#,##0;(#,##0);-",
                "periodOffset": column_period_offset(payload),
            }
        )
    output_rows = []
    for record in row_records:
        row_payload = record.get("payload") or {}
        row_type = normalize_key(row_payload.get("rowType") or "DATA")
        cells = row_payload.get("cells") if isinstance(row_payload.get("cells"), dict) else {}
        output_cells = []
        for column_record in columns:
            column_payload = column_record.get("payload") or {}
            column_key = normalize_key(column_payload.get("columnKey"))
            cell = cells.get(column_key) or cells.get(column_payload.get("columnKey")) or {}
            column_type = normalize_key(column_payload.get("columnType") or "CURRENT_PERIOD")
            kind = normalize_key(cell.get("contentKind"))
            is_label = (
                column_type in ("LABEL", "ROW_CODE", "NOTE_REFERENCE")
                or kind in ("TEXT", "TITLE", "SPACER")
                or row_type in ("TITLE", "SPACER")
            )
            display_text = _text(cell.get("text"))
            if not display_text and column_type == "ROW_CODE":
                display_text = _text(row_payload.get("displayCode") or row_payload.get("rowCode"))
            elif not display_text and column_type == "NOTE_REFERENCE":
                linked_note_id = _text(row_payload.get("noteItemId")).lower()
                if linked_note_id:
                    linked_note = note_item_by_id.get(linked_note_id) or {}
                    display_text = (
                        _text(linked_note.get("itemCode") or row_payload.get("noteReference"))
                        if linked_note_id in included_note_set
                        else ""
                    )
                else:
                    display_text = _text(row_payload.get("noteReference"))
            elif not display_text and (column_type == "LABEL" or row_type == "TITLE"):
                display_text = _text(row_payload.get("rowLabel") or record.get("displayName"))
            evaluated_value = None if is_label else evaluate_row(record, column_record)
            missing_formula = not is_label and cell_is_missing_formula(record, column_record)
            output_cell = {
                    "columnKey": column_payload.get("columnKey"),
                    "kind": "SPACER" if row_type == "SPACER" else "TEXT" if is_label else "VALUE",
                    "text": display_text if is_label else "",
                    "value": None if is_label or missing_formula else evaluated_value,
                    "formulaStatus": "MISSING" if missing_formula else "CONFIGURED" if not is_label else "",
                    "style": _style(row_payload, cell, styles, column_payload),
                    "mergeColumnSpan": max(
                        1,
                        int(
                            _number(
                                cell.get("mergeColumnSpan")
                                or (
                                    row_payload.get("mergeColumnSpan")
                                    if column_payload.get("columnKey") == row_payload.get("mergeColumnKey")
                                    else 1
                                )
                            )
                        ),
                    ),
                    "mergeRowSpan": max(
                        1,
                        int(
                            _number(
                                cell.get("mergeRowSpan")
                                or (
                                    row_payload.get("mergeRowSpan")
                                    if column_payload.get("columnKey") == row_payload.get("mergeColumnKey")
                                    else 1
                                )
                            )
                        ),
                    ),
                }
            if include_evidence and not is_label:
                output_cell["evidence"] = build_cell_evidence(record, column_record)
            output_cells.append(output_cell)
        output_rows.append(
            {
                "recordId": record.get("recordId"),
                "rowCode": row_payload.get("rowCode") or record.get("displayCode"),
                "rowLabel": row_payload.get("rowLabel") or record.get("displayName"),
                "rowType": row_type,
                "rowHeight": _number(row_payload.get("rowHeight")),
                "hideOnPrint": row_payload.get("hideOnPrint") is True,
                "cells": output_cells,
            }
        )
    note_output = _build_note_output(
        records=records,
        report_key=normalized_report,
        note_items=note_items,
        selected_note_item_ids=selected_note_item_ids,
        rows_by_offset=rows_by_offset,
        warning=warning,
    )
    report_payload = report_record.get("payload") or {}
    control_checks: list[dict[str, Any]] = []

    def output_value(row_code: str, column_index: int) -> float | None:
        row = next(
            (item for item in output_rows if _text(item.get("rowCode")).upper() == row_code.upper()),
            None,
        )
        if not row or column_index >= len(row.get("cells") or []):
            return None
        value = row["cells"][column_index].get("value")
        return None if value is None else _number(value)

    def source_cash_closing_value(column_record: dict[str, Any]) -> float | None:
        """Resolve cash ending balance independently from line 70's presentation formula."""
        cash_row = row_by_code.get("60")
        if not cash_row:
            return None
        cash_row_payload = cash_row.get("payload") or {}
        cash_rule = rules.get(normalize_key(cash_row_payload.get("ruleKey")), {})
        terms = cash_rule.get("terms") if isinstance(cash_rule.get("terms"), list) else []
        source_terms = [
            {
                **term,
                "periodBasis": "ENDING",
                "useColumnBasis": False,
            }
            for term in terms
            if normalize_key(term.get("operator")) != "MANUAL_ADJUSTMENT"
            and term.get("accountSetKey")
        ]
        if not source_terms:
            return None
        column_payload = column_record.get("payload") or {}
        offset = column_period_offset(column_payload)
        source_rows = rows_by_offset.get(offset, [])
        if not source_rows:
            return None
        return sum(
            _term_value(
                term,
                source_rows,
                members,
                field_mappings,
                shift_year(period_to, offset),
                column_payload,
                warning,
                term_runtime_cache,
                blocking_issue,
            )
            for term in source_terms
        )

    for column_index, column_record in enumerate(columns):
        column_payload = column_record.get("payload") or {}
        column_type = normalize_key(column_payload.get("columnType"))
        if column_type in ("LABEL", "ROW_CODE", "NOTE_REFERENCE"):
            continue
        column_key = _text(column_payload.get("columnKey"))
        column_label = _text(column_payload.get("columnLabel") or column_record.get("displayName"))
        if normalized_report == "BAO_CAO_TINH_HINH_TAI_CHINH":
            assets = output_value("270_TONG_TS", column_index)
            sources = output_value("440", column_index)
            if assets is not None and sources is not None:
                difference = assets - sources
                control_checks.append(
                    {
                        "checkKey": "BALANCE_SHEET_BALANCES",
                        "label": "Tổng tài sản bằng tổng nguồn vốn",
                        "columnKey": column_key,
                        "columnLabel": column_label,
                        "status": "PASS" if abs(difference) < 0.5 else "FAIL",
                        "leftLabel": "Tổng tài sản",
                        "leftValue": assets,
                        "rightLabel": "Tổng nguồn vốn",
                        "rightValue": sources,
                        "difference": difference,
                    }
                )
            offset = column_period_offset(column_payload)
            for message in _balance_sheet_overlap_messages(
                row_records=row_records,
                rules=rules,
                members=members,
                source_rows=rows_by_offset.get(offset, []),
                period_to=shift_year(period_to, offset),
                column_payload=column_payload,
            ):
                blocking_issue(message)
        if normalized_report in (
            "LUU_CHUYEN_TIEN_TE_GIAN_TIEP",
            "LUU_CHUYEN_TIEN_TE_TRUC_TIEP",
        ):
            net_cash_flow = output_value("50", column_index)
            opening_cash = output_value("60", column_index)
            exchange_effect = output_value("61", column_index)
            closing_cash = source_cash_closing_value(column_record)
            if None not in (net_cash_flow, opening_cash, exchange_effect, closing_cash):
                calculated_closing_cash = net_cash_flow + opening_cash + exchange_effect
                difference = calculated_closing_cash - closing_cash
                control_checks.append(
                    {
                        "checkKey": "CASH_FLOW_RECONCILES_TO_CASH_BALANCE",
                        "label": "Lưu chuyển tiền thuần khớp tiền và tương đương tiền đầu-cuối kỳ",
                        "columnKey": column_key,
                        "columnLabel": column_label,
                        "status": "PASS" if abs(difference) < 0.5 else "FAIL",
                        "leftLabel": "Tiền cuối kỳ theo lưu chuyển",
                        "leftValue": calculated_closing_cash,
                        "rightLabel": "Tiền và tương đương tiền cuối kỳ theo số dư nguồn",
                        "rightValue": closing_cash,
                        "difference": difference,
                    }
                )

    retained_row = next(
        (
            record
            for record in records
            if record.get("recordType") == "REPORT_ROW"
            and normalize_key((record.get("payload") or {}).get("reportKey"))
            == "BAO_CAO_TINH_HINH_TAI_CHINH"
            and _text((record.get("payload") or {}).get("rowCode")).upper() == "420A"
        ),
        None,
    )
    retained_rule = (
        rules.get(normalize_key((retained_row.get("payload") or {}).get("ruleKey")), {})
        if retained_row
        else {}
    )
    retained_set_key = retained_rule.get("accountSetKey")
    profit_row_code = (
        "60"
        if normalized_report == "BAO_CAO_KET_QUA_KINH_DOANH"
        else "420B"
        if normalized_report == "BAO_CAO_TINH_HINH_TAI_CHINH"
        else ""
    )
    if retained_set_key and profit_row_code:
        for column_index, column_record in enumerate(columns):
            column_payload = column_record.get("payload") or {}
            if normalize_key(column_payload.get("columnType")) != "CURRENT_PERIOD":
                continue
            offset = column_period_offset(column_payload)
            source_rows = rows_by_offset.get(offset, [])
            if not source_rows:
                continue
            profit = output_value(profit_row_code, column_index)
            retained_rows = _rows_for_set(source_rows, retained_set_key, members, period_to)
            retained_movement = sum(
                _metric_value(
                    row,
                    "",
                    {"operator": "MOVEMENT", "periodBasis": "PERIOD", "side": "CREDIT_BALANCE"},
                )
                for row in retained_rows
            )
            if profit is None or not retained_rows:
                continue
            difference = profit - retained_movement
            control_checks.append(
                {
                    "checkKey": "PROFIT_TO_RETAINED_EARNINGS",
                    "label": "LNST trên KQKD khớp biến động TK 421",
                    "columnKey": _text(column_payload.get("columnKey")),
                    "columnLabel": _text(column_payload.get("columnLabel") or column_record.get("displayName")),
                    "status": "PASS" if abs(difference) < 0.5 else "FAIL",
                    "leftLabel": "LNST trên báo cáo",
                    "leftValue": profit,
                    "rightLabel": "Biến động Có - Nợ TK 421",
                    "rightValue": retained_movement,
                    "difference": difference,
                }
            )

    for index, message in enumerate(blocking_issues, 1):
        control_checks.append(
            {
                "checkKey": f"SOURCE_CLASSIFICATION_COMPLETE_{index}",
                "label": message,
                "columnKey": "",
                "columnLabel": "",
                "status": "FAIL",
                "leftLabel": "Dữ liệu nguồn",
                "leftValue": None,
                "rightLabel": "Căn cứ phân loại bắt buộc",
                "rightValue": None,
                "difference": None,
            }
        )

    failed_checks = [check for check in control_checks if check["status"] == "FAIL"]
    for check in failed_checks:
        if check.get("difference") is None:
            warning(f"Kiểm tra kế toán chưa đạt: {check['label']}")
        else:
            warning(
                f"Kiểm tra kế toán chưa đạt: {check['label']} · chênh {check['difference']:,.0f}."
            )
    return {
        "ok": True,
        "report": {
            "reportKey": normalized_report,
            "reportName": report_payload.get("reportName") or report_record.get("displayName"),
            "displayTitle": report_payload.get("displayTitle")
            or report_payload.get("reportName")
            or report_record.get("displayName"),
            "hideColumnHeader": report_payload.get("hideColumnHeader") is True,
        },
        "columns": output_columns,
        "headerRows": _header_rows(output_columns),
        "rows": output_rows,
        "notes": note_output,
        "controlChecks": control_checks,
        "controlStatus": "NEEDS_REVIEW" if failed_checks else "READY",
        "blockingIssues": blocking_issues,
        "warnings": warnings[:80],
        "balanceRowCount": sum(len(rows) for rows in rows_by_offset.values()),
        "nonZeroCellCount": sum(
            1 for row in output_rows for cell in row["cells"] if cell["value"] is not None and cell["value"] != 0
        )
        + sum(
            1
            for row in note_output["rows"]
            for key in ("openingValue", "movementValue", "closingValue")
            if row.get(key) not in (None, 0)
        ),
        "periodFrom": period_from.isoformat(),
        "periodTo": period_to.isoformat(),
        "selection": selection,
        "adjustments": normalized_adjustments,
    }


def output_to_xlsx(output: dict[str, Any]) -> bytes:
    columns = output.get("columns") or []
    rows = [row for row in (output.get("rows") or []) if not row.get("hideOnPrint")]
    if not columns:
        raise ReportEngineError("Mẫu báo cáo chưa có cột để xuất.")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Báo cáo"
    sheet.sheet_view.showGridLines = False
    column_count = len(columns)
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=column_count)
    title = sheet.cell(1, 1, (output.get("report") or {}).get("displayTitle"))
    title.font = Font(bold=True, size=16, color="114C41")
    title.alignment = Alignment(horizontal="center", vertical="center")
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=column_count)
    period = sheet.cell(2, 1, f"Kỳ: {output.get('periodFrom')} đến {output.get('periodTo')}")
    period.font = Font(size=10, color="60736F")
    period.alignment = Alignment(horizontal="center")
    thin = Side(style="thin", color="D1DDDA")
    header_side = Side(style="thin", color="9DB7B0")
    header_row = 4
    report = output.get("report") or {}
    header_rows = [] if report.get("hideColumnHeader") else output.get("headerRows") or _header_rows(columns)
    occupied_header: set[tuple[int, int]] = set()
    for header_index, definition in enumerate(header_rows):
        cursor = 0
        for header_cell in definition.get("cells") or []:
            while (header_index, cursor) in occupied_header:
                cursor += 1
            if cursor >= column_count:
                break
            column_span = min(max(1, int(header_cell.get("columnSpan") or 1)), column_count - cursor)
            row_span = min(max(1, int(header_cell.get("rowSpan") or 1)), len(header_rows) - header_index)
            start_row = header_row + header_index
            start_column = cursor + 1
            if column_span > 1 or row_span > 1:
                sheet.merge_cells(
                    start_row=start_row,
                    start_column=start_column,
                    end_row=start_row + row_span - 1,
                    end_column=start_column + column_span - 1,
                )
            cell = sheet.cell(start_row, start_column, header_cell.get("label") or "")
            cell.font = Font(bold=True, color="114C41")
            cell.fill = PatternFill("solid", fgColor="EDF8F0")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = Border(left=header_side, right=header_side, top=header_side, bottom=header_side)
            for row_offset in range(row_span):
                for column_offset in range(column_span):
                    occupied_header.add((header_index + row_offset, cursor + column_offset))
            cursor += column_span
        sheet.row_dimensions[header_row + header_index].height = 28
    data_start = header_row if report.get("hideColumnHeader") else header_row + len(header_rows)
    occupied: set[tuple[int, int]] = set()
    for row_index, row in enumerate(rows):
        excel_row = data_start + row_index
        if row.get("rowHeight"):
            sheet.row_dimensions[excel_row].height = min(120, max(18, _number(row["rowHeight"])))
        for column_index, item in enumerate(row.get("cells") or []):
            if (row_index, column_index) in occupied:
                continue
            column_span = min(max(1, int(item.get("mergeColumnSpan") or 1)), column_count - column_index)
            row_span = min(max(1, int(item.get("mergeRowSpan") or 1)), len(rows) - row_index)
            if column_span > 1 or row_span > 1:
                sheet.merge_cells(
                    start_row=excel_row,
                    start_column=column_index + 1,
                    end_row=excel_row + row_span - 1,
                    end_column=column_index + column_span,
                )
                for row_offset in range(row_span):
                    for column_offset in range(column_span):
                        if row_offset or column_offset:
                            occupied.add((row_index + row_offset, column_index + column_offset))
            style = item.get("style") or {}
            value = item.get("value") if item.get("value") is not None else item.get("text") or ""
            if style.get("uppercase") and item.get("value") is None:
                value = str(value).upper()
            cell = sheet.cell(excel_row, column_index + 1, value)
            row_type = normalize_key(row.get("rowType"))
            cell.font = Font(
                bold=bool(style.get("bold") or row_type in ("TITLE", "SECTION", "SUBTOTAL", "TOTAL")),
                italic=bool(style.get("italic")),
                size=max(8, min(32, _number(style.get("fontSize")))) if style.get("fontSize") else None,
                color=_safe_color(style.get("textColor"), "183F37"),
            )
            cell.fill = PatternFill("solid", fgColor=_safe_color(style.get("fillColor"), "FFFFFF"))
            cell.alignment = Alignment(
                horizontal=str(style.get("align") or ("RIGHT" if item.get("value") is not None else "LEFT")).lower(),
                vertical="center",
                wrap_text=True,
                indent=max(0, min(15, int(_number(style.get("indent"))))),
            )
            if item.get("value") is not None:
                cell.number_format = style.get("numberFormat") or columns[column_index].get("numberFormat") or "#,##0;(#,##0);-"
            preset = normalize_key(style.get("borderPreset") or "DEFAULT")
            strong = Side(
                style="thick" if _number(style.get("borderWidth")) >= 4 else "medium",
                color="114C41",
            )
            cell.border = Border(
                left=strong if preset == "OUTLINE" else thin,
                right=strong if preset == "OUTLINE" else thin,
                top=strong if preset in ("TOP", "TOP_BOTTOM", "OUTLINE") else thin,
                bottom=strong if preset in ("BOTTOM", "TOP_BOTTOM", "OUTLINE") else thin,
            )
    for index, column in enumerate(columns, 1):
        sheet.column_dimensions[get_column_letter(index)].width = min(65, max(10, _number(column.get("width"))))
    sheet.freeze_panes = f"A{3 if report.get('hideColumnHeader') else header_row + len(header_rows)}"
    notes = output.get("notes") or {}
    note_rows = notes.get("rows") if isinstance(notes.get("rows"), list) else []
    if note_rows:
        note_sheet = workbook.create_sheet("Thuyết minh")
        note_sheet.sheet_view.showGridLines = False
        note_sheet.merge_cells("A1:E1")
        note_title = note_sheet["A1"]
        note_title.value = notes.get("title") or "THUYẾT MINH BÁO CÁO TÀI CHÍNH"
        note_title.font = Font(bold=True, size=15, color="114C41")
        note_title.alignment = Alignment(horizontal="center")
        note_sheet.merge_cells("A2:E2")
        note_sheet["A2"] = f"Kỳ: {output.get('periodFrom')} đến {output.get('periodTo')}"
        note_sheet["A2"].font = Font(size=10, color="60736F")
        note_sheet["A2"].alignment = Alignment(horizontal="center")
        note_headers = ["Mã", "Nội dung thuyết minh", "Đầu kỳ", "Phát sinh", "Cuối kỳ"]
        for index, label in enumerate(note_headers, 1):
            cell = note_sheet.cell(4, index, label)
            cell.font = Font(bold=True, color="114C41")
            cell.fill = PatternFill("solid", fgColor="EDF8F0")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = Border(left=header_side, right=header_side, top=header_side, bottom=header_side)
        for row_index, item in enumerate(note_rows, 5):
            values = [
                item.get("itemCode") or "",
                item.get("itemName") or "",
                item.get("openingValue"),
                item.get("movementValue"),
                item.get("closingValue"),
            ]
            for column_index, value in enumerate(values, 1):
                cell = note_sheet.cell(row_index, column_index, "" if value is None else value)
                cell.font = Font(
                    bold=bool(item.get("isBold") or item.get("required")),
                    italic=bool(item.get("isItalic")),
                    color="183F37",
                )
                cell.alignment = Alignment(
                    horizontal="right" if column_index >= 3 else "left",
                    vertical="center",
                    wrap_text=True,
                )
                cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
                if column_index >= 3 and value is not None:
                    cell.number_format = "#,##0;(#,##0);-"
        for index, width in enumerate((16, 62, 18, 18, 18), 1):
            note_sheet.column_dimensions[get_column_letter(index)].width = width
        note_sheet.freeze_panes = "A5"
    checks = output.get("controlChecks") or []
    if checks:
        check_sheet = workbook.create_sheet("Kiểm tra")
        check_sheet.sheet_view.showGridLines = False
        check_sheet.merge_cells("A1:H1")
        check_sheet["A1"] = "KIỂM TRA KẾ TOÁN TỰ ĐỘNG"
        check_sheet["A1"].font = Font(bold=True, size=15, color="114C41")
        check_sheet["A1"].alignment = Alignment(horizontal="center")
        check_sheet.merge_cells("A2:H2")
        check_sheet["A2"] = (
            "SẴN SÀNG PHÁT HÀNH"
            if output.get("controlStatus") == "READY"
            else "CẦN XỬ LÝ TRƯỚC KHI PHÁT HÀNH"
        )
        check_sheet["A2"].font = Font(
            bold=True,
            color="0D6B4E" if output.get("controlStatus") == "READY" else "A92F26",
        )
        check_sheet["A2"].alignment = Alignment(horizontal="center")
        check_headers = [
            "Trạng thái", "Phép kiểm tra", "Cột", "Vế trái", "Giá trị",
            "Vế phải", "Giá trị", "Chênh lệch",
        ]
        for index, label in enumerate(check_headers, 1):
            cell = check_sheet.cell(4, index, label)
            cell.font = Font(bold=True, color="114C41")
            cell.fill = PatternFill("solid", fgColor="EDF8F0")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = Border(left=header_side, right=header_side, top=header_side, bottom=header_side)
        for row_index, check in enumerate(checks, 5):
            passed = check.get("status") == "PASS"
            values = [
                "ĐẠT" if passed else "CHƯA ĐẠT",
                check.get("label") or "",
                check.get("columnLabel") or check.get("columnKey") or "",
                check.get("leftLabel") or "",
                check.get("leftValue"),
                check.get("rightLabel") or "",
                check.get("rightValue"),
                check.get("difference"),
            ]
            for column_index, value in enumerate(values, 1):
                cell = check_sheet.cell(row_index, column_index, "" if value is None else value)
                cell.font = Font(
                    bold=column_index in (1, 8),
                    color="0D6B4E" if passed else "A92F26" if column_index in (1, 8) else "183F37",
                )
                cell.fill = PatternFill("solid", fgColor="EFF9F3" if passed else "FFF1EF")
                cell.alignment = Alignment(
                    horizontal="right" if column_index in (5, 7, 8) else "left",
                    vertical="center",
                    wrap_text=True,
                )
                cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
                if column_index in (5, 7, 8) and value is not None:
                    cell.number_format = "#,##0;(#,##0);-"
        for index, width in enumerate((14, 42, 18, 28, 18, 28, 18, 18), 1):
            check_sheet.column_dimensions[get_column_letter(index)].width = width
        check_sheet.freeze_panes = "A5"
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
