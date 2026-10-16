from __future__ import annotations

import base64
import io
import json
from datetime import date
from decimal import Decimal

from fastapi import HTTPException
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

import worker


def _preview_workbook(workbook: Workbook, file_name: str):
    stream = io.BytesIO()
    workbook.save(stream)
    parsed = load_workbook(io.BytesIO(stream.getvalue()), read_only=True, data_only=True)
    try:
        return worker._accounting_file_preview(
            parsed,
            worker.AccountingFilePreviewRequest(
                fileName=file_name,
                fileDataUrl="unused",
                companyKey="TVD",
                companyLabel="TVD",
            ),
            file_name,
        )
    finally:
        parsed.close()


def test_excel_import_detects_single_row_general_journal_header() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "SỔ NHẬT KÝ CHUNG"
    sheet.append(["SỔ NHẬT KÝ CHUNG"])
    sheet.append(["Tháng 1 năm 2026"])
    sheet.append([])
    sheet.append([
        "Ngày hạch toán", "Ngày chứng từ", "Số chứng từ", "Diễn giải",
        "Tài khoản", "TK đối ứng", "Phát sinh Nợ", "Phát sinh Có",
    ])
    sheet.append([
        "01/01/2026", "01/01/2026", "CT001", "Bán hàng",
        "1111", "5111", 100, 0,
    ])
    sheet.append([
        "01/01/2026", "01/01/2026", "CT001", "Bán hàng",
        "5111", "1111", 0, 100,
    ])

    preview = _preview_workbook(workbook, "SO NHAT KY CHUNG T01.2026_DIGITAL.xlsx")

    assert preview["datasetType"] == "GENERAL_LEDGER"
    assert preview["headerRow"] == 4
    assert preview["headerDepth"] == 1
    assert preview["periodFrom"] == "2026-01-01"
    assert preview["periodTo"] == "2026-01-31"
    assert preview["rowCount"] == 2
    assert preview["proposedDataSpaceName"] == "TVD_2026"


def test_excel_import_detects_two_level_trial_balance_header() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "BẢNG CÂN ĐỐI TÀI KHOẢN"
    sheet.append(["CÔNG TY TNHH THUẦN VIỆT DIGITAL"])
    sheet.append([])
    sheet.append(["Tháng 1 năm 2026"])
    for _ in range(4):
        sheet.append([])
    sheet.append([
        "Số tài khoản", "Tên tài khoản", "Đầu kỳ", None,
        "Phát sinh", None, None, None, "Cuối kỳ", None,
    ])
    sheet.append([
        None, None, "Nợ", "Có", "Nợ", "Có",
        "Nợ lũy kế", "Có lũy kế", "Nợ", "Có",
    ])
    sheet.append(["111", "Tiền mặt", 10, 0, 5, 2, 5, 2, 13, 0])

    preview = _preview_workbook(workbook, "BANG CAN DOI TAI KHOAN T01.2026_DIGITAL.xlsx")

    assert preview["datasetType"] == "TRIAL_BALANCE"
    assert preview["headerRow"] == 8
    assert preview["headerDepth"] == 2
    assert preview["rowCount"] == 1
    assert {item["targetField"] for item in preview["mappings"]} >= {
        "account_number", "account_name", "opening_debit_amount",
        "opening_credit_amount", "debit_amount", "credit_amount",
        "accum_debit_amount", "accum_credit_amount",
        "closing_debit_amount", "closing_credit_amount",
    }


def test_liveness_does_not_require_database_or_secret() -> None:
    with TestClient(worker.app) as client:
        response = client.get("/healthz")
    assert response.status_code == 200
    assert response.json()["service"] == "accounting-worker"


def test_readiness_fails_closed_when_service_is_not_configured() -> None:
    with TestClient(worker.app) as client:
        response = client.get("/readyz")
    assert response.status_code == 503


def test_protected_route_uses_worker_secret(monkeypatch) -> None:
    monkeypatch.setattr(worker, "WORKER_SECRET", "a" * 40)
    with TestClient(worker.app) as client:
        denied = client.get("/health", headers={"Authorization": "Bearer wrong"})
    assert denied.status_code == 401


def test_misa_url_restriction_blocks_non_misa_and_userinfo() -> None:
    assert worker._allowed_misa_url("https://actapp.misa.vn/g1/api/report") == (
        "https://actapp.misa.vn/g1/api/report"
    )
    for url in (
        "http://actapp.misa.vn/g1/api/report",
        "https://misa.vn.example.com/api",
        "https://user:password@actapp.misa.vn/api",
        "https://actapp.misa.vn:8443/api",
    ):
        try:
            worker._allowed_misa_url(url)
        except Exception:
            pass
        else:
            raise AssertionError(f"URL nguy hiểm không bị chặn: {url}")


def test_runtime_schema_contains_all_snapshot_and_audit_tables() -> None:
    for table in (
        "accounting_source_connections",
        "accounting_snapshots",
        "accounting_sync_runs",
        "accounting_source_rows",
        "accounting_journal_entries",
        "accounting_journal_monthly",
        "accounting_data_controls",
        "accounting_active_snapshots",
        "accounting_change_events",
        "accounting_report_adjustments",
        "accounting_report_adjustment_events",
        "accounting_note_catalog_meta",
        "accounting_note_catalog_items",
        "accounting_dirty_periods",
        "accounting_cleanup_audit",
    ):
        assert table in worker.DDL
    for column in (
        "credential_status",
        "credential_last_checked_at",
        "credential_last_success_at",
        "credential_last_error",
    ):
        assert column in worker.DDL


def test_report_formula_text_condition_requires_detailed_journal_rows() -> None:
    records = [
        {
            "payload": {
                "cells": {
                    "CURRENT": {
                        "terms": [
                            {
                                "operator": "CASH_MOVEMENT",
                                "filters": {
                                    "logic": "ANY",
                                    "conditions": [
                                        {
                                            "field": "search_text",
                                            "operator": "CONTAINS_ANY",
                                            "value": ["lãi vay", "loan interest"],
                                        }
                                    ],
                                },
                            }
                        ]
                    }
                }
            }
        }
    ]

    assert worker._report_records_require_journal_detail(records) is True


def test_report_formula_dimension_condition_can_use_monthly_journal_rows() -> None:
    records = [
        {
            "payload": {
                "terms": [
                    {
                        "filters": {
                            "conditions": [
                                {
                                    "field": "cash_flow_code",
                                    "operator": "EQUALS_ANY",
                                    "value": ["14"],
                                }
                            ]
                        }
                    }
                ]
            }
        }
    ]

    assert worker._report_records_require_journal_detail(records) is False


def test_journal_detail_detection_is_scoped_to_the_requested_report() -> None:
    records = [
        {
            "recordType": "REPORT",
            "entityKey": "B01",
            "payload": {},
        },
        {
            "recordType": "REPORT_ROW",
            "entityKey": "B01_100",
            "payload": {
                "reportKey": "B01",
                "rowCode": "100",
                "cells": {"CURRENT": {"operator": "BALANCE"}},
            },
        },
        {
            "recordType": "REPORT",
            "entityKey": "B03",
            "payload": {},
        },
        {
            "recordType": "REPORT_ROW",
            "entityKey": "B03_14",
            "payload": {
                "reportKey": "B03",
                "rowCode": "14",
                "cells": {
                    "CURRENT": {
                        "operator": "CASH_MOVEMENT",
                        "filters": {
                            "conditions": [
                                {
                                    "field": "search_text",
                                    "operator": "CONTAINS_ANY",
                                    "value": ["lãi vay"],
                                }
                            ]
                        },
                    }
                },
            },
        },
    ]

    assert worker._report_records_require_journal_detail(records, "B01") is False
    assert worker._report_records_require_journal_detail(records, "B03") is True


def test_journal_detail_detection_follows_cross_report_dependencies() -> None:
    records = [
        {"recordType": "REPORT", "entityKey": "B01", "payload": {}},
        {"recordType": "REPORT", "entityKey": "B03", "payload": {}},
        {
            "recordType": "REPORT_ROW",
            "entityKey": "B01_420",
            "payload": {
                "reportKey": "B01",
                "rowCode": "420",
                "cells": {
                    "CURRENT": {
                        "operator": "CROSS_REPORT",
                        "targetReportKey": "B03",
                        "targetRowCode": "14",
                    }
                },
            },
        },
        {
            "recordType": "REPORT_ROW",
            "entityKey": "B03_14",
            "payload": {
                "reportKey": "B03",
                "rowCode": "14",
                "cells": {
                    "CURRENT": {
                        "operator": "CASH_MOVEMENT",
                        "filters": {
                            "conditions": [
                                {
                                    "field": "reference_no",
                                    "operator": "CONTAINS",
                                    "value": "UNC",
                                }
                            ]
                        },
                    }
                },
            },
        },
    ]

    assert worker._report_dependency_keys(records, "B01") == {"B01", "B03"}
    assert worker._report_records_require_journal_detail(records, "B01") is True


def test_disclosure_detail_detection_follows_only_its_linked_reports() -> None:
    records = [
        {
            "recordType": "DISCLOSURE_TEMPLATE",
            "entityKey": worker.DISCLOSURE_REPORT_KEY,
            "payload": {"templateKey": worker.DISCLOSURE_REPORT_KEY},
        },
        {
            "recordType": "DISCLOSURE_LINK",
            "entityKey": "VI_6:1",
            "payload": {"disclosureItemKey": "VI_6", "reportKey": "B03", "rowCode": "14"},
        },
        {"recordType": "REPORT", "entityKey": "B03", "payload": {}},
        {
            "recordType": "REPORT_ROW",
            "entityKey": "B03_14",
            "payload": {
                "reportKey": "B03",
                "rowCode": "14",
                "filters": {
                    "conditions": [
                        {"field": "search_text", "operator": "CONTAINS", "value": "lãi vay"}
                    ]
                },
            },
        },
        {
            "recordType": "REPORT_ROW",
            "entityKey": "UNRELATED_01",
            "payload": {
                "reportKey": "UNRELATED",
                "filters": {
                    "conditions": [
                        {"field": "reference_no", "operator": "CONTAINS", "value": "UNC"}
                    ]
                },
            },
        },
    ]

    assert worker._report_dependency_keys(records, worker.DISCLOSURE_REPORT_KEY) == {
        worker.DISCLOSURE_REPORT_KEY,
        "B03",
    }
    assert worker._report_records_require_journal_detail(
        records, worker.DISCLOSURE_REPORT_KEY
    ) is True


def test_approved_adjustment_is_exposed_as_auditable_report_source_row() -> None:
    row = worker._report_adjustment_source_row(
        {
            "adjustment_id": "11111111-1111-1111-1111-111111111111",
            "external_key": "TVG-2026-01-CF-14",
            "report_key": "CF",
            "row_code": "14",
            "column_key": "KY_NAY",
            "company_key": "TVG",
            "data_space_key": "SPACE",
            "data_space_name": "GLOBAL_2026",
            "period_from": "2026-01-01",
            "period_to": "2026-01-31",
            "amount": Decimal("-2007252"),
            "reason": "Chi phí đi vay đã trả",
            "source_reference": "4 BCTC_Global.xlsx!CF!F34",
            "supporting_data": {"manualCell": "F34"},
            "status": "APPROVED",
            "approved_by": "Kế toán trưởng",
            "approved_at": "2026-08-24T10:00:00+07:00",
        }
    )

    assert row["source_key"] == "BCTC_DIEU_CHINH"
    assert row["dataset_type"] == "MANUAL_ADJUSTMENT"
    assert row["canonical_data"]["amount"] == Decimal("-2007252")
    assert row["canonical_data"]["approval_status"] == "APPROVED"


def test_official_report_export_is_blocked_when_accounting_controls_fail() -> None:
    try:
        worker._assert_report_exportable(
            {
                "blockingIssues": ["TK 1311 đang được tính vào hai dòng."],
                "controlChecks": [
                    {
                        "status": "FAIL",
                        "label": "Tổng tài sản bằng tổng nguồn vốn",
                        "difference": 125,
                    }
                ],
            }
        )
    except HTTPException as exc:
        assert exc.status_code == 409
        assert "Chưa thể xuất báo cáo chính thức" in str(exc.detail)
        assert "TK 1311" in str(exc.detail)
    else:
        raise AssertionError("Báo cáo lỗi kiểm soát vẫn được xuất chính thức")


def test_official_report_export_accepts_a_clean_preview() -> None:
    worker._assert_report_exportable(
        {
            "blockingIssues": [],
            "controlChecks": [{"status": "PASS", "label": "Đã cân"}],
            "controlStatus": "PASS",
        }
    )


def test_jwt_expiry_is_only_informational_until_misa_rejects_it() -> None:
    expired_jwt = "e30.eyJleHAiOjF9.signature"
    credentials = worker._credentials_with_bearer({}, expired_jwt)
    assert credentials["bearer"] == expired_jwt


def _unsigned_test_jwt(payload: dict) -> str:
    encode = lambda value: base64.urlsafe_b64encode(
        json.dumps(value, separators=(",", ":")).encode("utf-8")
    ).decode("ascii").rstrip("=")
    return f"{encode({'alg': 'none'})}.{encode(payload)}.test"


def test_system_trial_balance_profile_survives_without_any_saved_source(monkeypatch) -> None:
    monkeypatch.setattr(worker, "MISA_BOOTSTRAP_TENANT_ID", "tenant-system")
    monkeypatch.setattr(worker, "MISA_BOOTSTRAP_TENANT_CODE", "TV")
    monkeypatch.setattr(worker, "MISA_BOOTSTRAP_DATABASE_ID", "database-anchor")
    monkeypatch.setattr(worker, "MISA_BOOTSTRAP_BRANCH_ID", "branch-anchor")
    monkeypatch.setattr(worker, "MISA_BOOTSTRAP_COMPANY_KEY", "TVG")
    monkeypatch.setattr(worker, "MISA_BOOTSTRAP_CONTEXTS", "")
    token = _unsigned_test_jwt({"uid": "user-1", "aut": "0", "exp": 2_000_000_000})

    source, credentials, profile = worker._misa_bootstrap_reference(token, company_key="TVG")
    context = json.loads(credentials["misaContext"])
    body = credentials["requestBody"]
    columns = json.loads(body["columns"])

    assert source["source_key"] == ""
    assert source["dataset_type"] == "TRIAL_BALANCE"
    assert source["base_source_key"] == "AMIS_TRIAL_BALANCE"
    assert profile["profileKey"] == "TRIAL_BALANCE"
    assert context["TenantId"] == "tenant-system"
    assert context["DatabaseId"] == "database-anchor"
    assert context["BranchId"] == "branch-anchor"
    assert context["UserId"] == "user-1"
    assert body["reportList"]["report_id"] == "AccumGLF01_DN"
    assert body["reportList"]["is_favorite"] is False
    assert body["reportList"]["is_disabled"] is False
    assert body["reportList"]["script_summary"] == worker.TRIAL_BALANCE_SCRIPT_SUMMARY
    assert body["reportList"]["preview_image"] == "BẢNG CÂN ĐỐI TÀI KHOẢN ( Mẫu quản trị ).png"
    assert body["reportList"]["is_created_from_old_db"] is False
    assert {item["field"] for item in columns} >= {
        "opening_debit_amount",
        "opening_credit_amount",
        "accum_debit_amount",
        "accum_credit_amount",
        "closing_debit_amount",
        "closing_credit_amount",
    }


def test_system_general_ledger_profile_is_ready_for_the_same_misa_space() -> None:
    profile = worker._misa_system_profile("GENERAL_LEDGER")
    body = profile["requestBody"]
    columns = json.loads(body["columns"])
    parameters = worker._decode_parameters(body["parameters"])

    assert profile["baseSourceKey"] == "AMIS_GENERAL_LEDGER"
    assert profile["datasetType"] == "GENERAL_LEDGER"
    assert profile["sourceName"] == "Sổ nhật ký chung"
    assert profile["pageSize"] == 100
    assert body["reportList"]["report_id"] == "GLGeneralDiary"
    assert body["reportList"]["function_report_name"] == "func_rpt_gl_general_diary_book"
    assert body["reportList"]["rp_function_name_async"] == "func_rpt_gl_general_diary_book_v4"
    assert body["reportList"]["version"] == 2
    assert base64.b64decode(body["report_id"]).decode("ascii") == "GLGeneralDiary"
    assert parameters["p_is_refresh"] is True
    assert parameters["p_is_similar_sum"] is False
    assert parameters["p_is_show_accum_amount"] is False
    assert "p_is_sum_same_entry" not in parameters
    assert "p_is_show_accum" not in parameters
    assert "p_language" not in parameters
    assert {item["field"] for item in columns} >= {
        "posted_date",
        "refno",
        "account_number",
        "corresponding_account_number",
        "debit_amount",
        "credit_amount",
    }


def test_all_builtin_misa_tables_are_created_together() -> None:
    profiles = worker._misa_system_profiles()
    assert [item["datasetType"] for item in profiles] == [
        "TRIAL_BALANCE",
        "GENERAL_LEDGER",
    ]


def test_blank_misa_failure_keeps_the_code_and_is_retryable() -> None:
    code, detail, readable = worker._misa_failure_detail(
        {"Success": False, "Code": 0, "ErrorsMessage": [], "Data": {"tenant": "ẩn"}}
    )
    assert code == 0
    assert "mã 0" in detail
    assert readable is False
    assert "tenant" not in detail
    try:
        worker._raise_misa_payload_failure(
            {"Success": False, "Code": 0, "ErrorsMessage": []}
        )
    except worker.MisaTransientError as error:
        assert "MISA" in str(error)
    else:
        raise AssertionError("Lỗi MISA rỗng phải được thử lại thay vì thất bại ngay")


def test_system_misa_profile_requires_server_tenant_not_accountant_json(monkeypatch) -> None:
    monkeypatch.setattr(worker, "MISA_BOOTSTRAP_TENANT_ID", "")
    monkeypatch.setattr(worker, "MISA_BOOTSTRAP_TENANT_CODE", "")
    monkeypatch.setattr(worker, "MISA_BOOTSTRAP_DATABASE_ID", "")
    monkeypatch.setattr(worker, "MISA_BOOTSTRAP_BRANCH_ID", "")
    monkeypatch.setattr(worker, "MISA_BOOTSTRAP_COMPANY_KEY", "")
    monkeypatch.setattr(worker, "MISA_BOOTSTRAP_CONTEXTS", "")
    token = _unsigned_test_jwt({"uid": "user-1"})
    try:
        worker._misa_bootstrap_reference(token)
    except HTTPException as error:
        assert error.status_code == 503
        assert "Kế toán không phải nhập JSON" in str(error.detail)
    else:
        raise AssertionError("Worker phải chặn khi cấu hình hệ thống chưa được thiết lập")


def test_system_misa_profile_selects_bootstrap_anchor_by_company(monkeypatch) -> None:
    monkeypatch.setattr(worker, "MISA_BOOTSTRAP_TENANT_ID", "")
    monkeypatch.setattr(worker, "MISA_BOOTSTRAP_TENANT_CODE", "")
    monkeypatch.setattr(worker, "MISA_BOOTSTRAP_DATABASE_ID", "")
    monkeypatch.setattr(worker, "MISA_BOOTSTRAP_BRANCH_ID", "")
    monkeypatch.setattr(worker, "MISA_BOOTSTRAP_COMPANY_KEY", "")
    monkeypatch.setattr(
        worker,
        "MISA_BOOTSTRAP_CONTEXTS",
        json.dumps({
            "TVG": {
                "TenantId": "tenant-tvg",
                "TenantCode": "TVG-MISA",
                "DatabaseId": "database-tvg",
                "BranchId": "branch-tvg",
            },
            "NTV": {
                "TenantId": "tenant-ntv",
                "TenantCode": "NTV-MISA",
                "DatabaseId": "database-ntv",
                "BranchId": "branch-ntv",
            },
        }),
    )
    token = _unsigned_test_jwt({"uid": "user-1", "aut": "0"})

    _, credentials, _ = worker._misa_bootstrap_reference(token, company_key="ntv")
    context = json.loads(credentials["misaContext"])

    assert context["TenantId"] == "tenant-ntv"
    assert context["DatabaseId"] == "database-ntv"
    assert context["BranchId"] == "branch-ntv"


def test_visible_databases_uses_current_misa_management_endpoint_first(monkeypatch) -> None:
    calls = []

    def fake_request(method, url, credentials, body=None):
        calls.append((method, url, body))
        return {
            "Success": True,
            "Data": [{
                "tenant_id": "tenant-1",
                "tenant_code": "TV",
                "database_id": "database-1",
                "database_name": "Dữ liệu thật",
                "user_status": 1,
            }],
        }

    monkeypatch.setattr(worker, "_misa_platform_request", fake_request)
    credentials = {
        "misaContext": json.dumps({
            "TenantId": "tenant-1",
            "DatabaseId": "database-1",
            "BranchId": "branch-1",
            "UserId": "user-1",
            "AuthType": 0,
        }),
    }
    source = {
        "endpoint_url": "https://actapp.misa.vn/g1/api/report/v1/report/dynamic/v2/paging_filter",
    }

    rows = worker._misa_visible_databases(source, credentials)

    assert len(calls) == 1
    assert calls[0][0] == "POST"
    assert calls[0][1].endswith("/g1/api/system/v1/database/databases_user_can_see")
    assert calls[0][2]["IsCheckConnectOtherApp"] is True
    assert rows[0]["databaseName"] == "Dữ liệu thật"
    assert rows[0]["isCurrent"] is True


def test_database_context_retries_misa_code_4_with_continue_confirmation(monkeypatch) -> None:
    calls = []

    def fake_request(method, url, credentials, body=None, allowed_failure_codes=None):
        calls.append((method, url, allowed_failure_codes))
        if url.endswith("isContinueAccessDBBackup=false"):
            assert allowed_failure_codes == {4}
            return {
                "Success": False,
                "Code": 4,
                "ErrorsMessage": [],
                "Data": {
                    "tenant": "Công ty mẫu",
                    "databaseId": "database-target",
                    "userId": "user-1",
                },
            }
        return {
            "Success": True,
            "Code": 0,
            "Data": {
                "Context": json.dumps({
                    "DatabaseId": "database-target",
                    "BranchId": "branch-1",
                }),
                "Env": "g1",
            },
        }

    monkeypatch.setattr(worker, "_misa_platform_request", fake_request)
    payload = worker._misa_database_context(
        {"endpoint_url": "https://actapp.misa.vn/g1/api/report/v1/report/dynamic/v2/paging_filter"},
        {"bearer": "token", "misaContext": "{}"},
        "database-target",
        "user-1",
    )

    assert payload["Success"] is True
    assert len(calls) == 2
    assert calls[0][1].endswith("?isContinueAccessDBBackup=false")
    assert calls[1][1].endswith("?isContinueAccessDBBackup=true")
    assert calls[1][2] is None


def test_database_context_does_not_retry_unexpected_misa_failure(monkeypatch) -> None:
    calls = []

    def fake_request(method, url, credentials, body=None, allowed_failure_codes=None):
        calls.append(url)
        return {"Success": False, "Code": 9, "ErrorsMessage": [], "Data": {}}

    monkeypatch.setattr(worker, "_misa_platform_request", fake_request)
    try:
        worker._misa_database_context(
            {"endpoint_url": "https://actapp.misa.vn/g1/api/report/v1/report/dynamic/v2/paging_filter"},
            {"bearer": "token", "misaContext": "{}"},
            "database-target",
            "user-1",
        )
    except worker.WorkerError as error:
        assert "chưa cho phép mở khối" in str(error)
    else:
        raise AssertionError("Không được tự xác nhận cho mã lỗi MISA ngoài Code 4")
    assert len(calls) == 1


def test_empty_snapshot_cannot_replace_a_non_empty_snapshot() -> None:
    message = worker._empty_snapshot_rejection_message(0, 30_064)
    assert "giữ nguyên phiên tốt" in message
    assert worker._empty_snapshot_rejection_message(0, 0) == ""
    assert worker._empty_snapshot_rejection_message(12, 30_064) == ""


def test_misa_report_building_is_not_mistaken_for_zero_rows() -> None:
    assert worker._misa_report_pending(
        {"Code": 210, "Data": {"Total": 0, "TableEmpty": False, "isCheckSessionStatus": False}},
        [],
    )
    assert worker._misa_report_pending(
        {"Data": {"Total": 0, "TableEmpty": True, "isCheckSessionStatus": True}},
        [],
    )
    assert worker._misa_report_pending(
        {"Code": 212, "Data": {"Total": 0, "TableEmpty": False, "isCheckSessionStatus": False}},
        [],
    )
    assert not worker._misa_report_pending(
        {"Data": {"Total": 0, "TableEmpty": True, "isCheckSessionStatus": False}},
        [],
    )
    assert not worker._misa_report_pending(
        {"Code": 211, "Data": {"Total": 0, "TableEmpty": False, "isCheckSessionStatus": False}},
        [],
    )
    assert not worker._misa_report_pending(
        {"Data": {"Total": 1, "TableEmpty": False}},
        [{"account_number": "111"}],
    )


def test_misa_page_integrity_retries_an_empty_middle_page() -> None:
    worker._validate_misa_page_row_count([{}] * 100, 36_408, 100, 218)
    worker._validate_misa_page_row_count([{}] * 8, 36_408, 100, 365)

    try:
        worker._validate_misa_page_row_count([], 36_408, 100, 219)
    except worker.MisaTransientError as error:
        assert "trang 219/365" in str(error)
        assert "0/100" in str(error)
    else:
        raise AssertionError("Trang giữa rỗng phải được giữ nguyên vị trí để tự thử lại")


def test_misa_page_integrity_retries_a_partial_page() -> None:
    try:
        worker._validate_misa_page_row_count([{}] * 99, 36_408, 100, 219)
    except worker.MisaTransientError as error:
        assert "99/100" in str(error)
    else:
        raise AssertionError("Trang giữa thiếu dòng không được đánh dấu hoàn tất")


def test_misa_report_session_id_accepts_report_response_variants() -> None:
    assert worker._misa_report_session_id(
        {"Code": 210, "Data": {"SessionId": "root-session"}}
    ) == "root-session"
    assert worker._misa_report_session_id(
        {"Code": 211, "Data": {"sessionId": "camel-session"}}
    ) == "camel-session"
    assert worker._misa_report_session_id(
        {
            "Code": 211,
            "Data": {
                "PageData": [
                    {"session_id": "journal-session", "account_number": "111"}
                ]
            },
        }
    ) == "journal-session"
    assert worker._misa_report_session_id(
        {
            "Code": 211,
            "Data": {
                "PageData": json.dumps([{"session_id": "json-session"}])
            },
        }
    ) == "json-session"
    assert worker._misa_report_session_id({"Code": 210, "Data": {"Total": 0}}) == ""


def test_misa_initial_request_uses_the_protocol_of_each_report_viewer() -> None:
    from datetime import date

    template = worker._general_ledger_request_template()
    journal_body = worker._request_body(
        template,
        {
            "dataset_type": "GENERAL_LEDGER",
            "period_from": date(2026, 1, 1),
            "period_to": date(2026, 1, 31),
            "page_size": 1000,
        },
        True,
    )
    trial_body = worker._request_body(
        worker._trial_balance_request_template(),
        {
            "dataset_type": "TRIAL_BALANCE",
            "period_from": date(2026, 1, 1),
            "period_to": date(2026, 1, 31),
            "page_size": 100,
        },
        True,
    )

    assert journal_body["isViewMCP"] is True
    assert journal_body["pageSize"] == 100
    assert trial_body["isViewMCP"] is True
    assert "sessionId" not in journal_body
    assert "sessionId" not in trial_body
    assert worker._decode_parameters(journal_body["parameters"])["p_is_refresh"] is True
    assert len(worker._decode_parameters(journal_body["parameters"])["p_session_key"]) == 64


def test_misa_first_page_uses_the_protocol_of_each_report_viewer() -> None:
    from datetime import date

    run = {
        "period_from": date(2026, 1, 1),
        "period_to": date(2026, 1, 31),
        "page_size": 100,
        "misa_session_id": "report-session",
    }
    journal_body = worker._request_body(
        worker._general_ledger_request_template(),
        {**run, "dataset_type": "GENERAL_LEDGER"},
        False,
        1,
    )
    trial_body = worker._request_body(
        worker._trial_balance_request_template(),
        {**run, "dataset_type": "TRIAL_BALANCE"},
        False,
        1,
    )

    assert journal_body["requestFrom"] == 1
    assert trial_body["requestFrom"] == 2
    assert journal_body["sessionId"] == "report-session"
    assert trial_body["sessionId"] == "report-session"
    assert journal_body["isViewMCP"] is False
    assert trial_body["isViewMCP"] is False


def test_misa_session_key_matches_the_viewer_javascript_rules() -> None:
    parameters = {
        "p_branch_id": "branch-1,",
        "p_include_dependent_branch": False,
        "p_from_date": "2025-12-31T17:00:00.000Z",
        "p_to_date": "2026-01-30T17:00:00.000Z",
        "p_is_similar_sum": False,
        "p_is_show_accum_amount": False,
        "p_is_management_book": False,
        "p_is_refresh": True,
        "p_session_key": "old-value-is-ignored",
    }
    material = "branch-1,2025-12-31T17:00:00.00Z2026-01-30T17:00:00.00Z"

    assert worker._misa_session_key(parameters) == worker.hashlib.sha256(material.encode("utf-8")).hexdigest()


def test_misa_journal_request_upgrades_a_saved_legacy_template() -> None:
    from datetime import date

    template = worker._general_ledger_request_template()
    parameters = worker._decode_parameters(template["parameters"])
    parameters.pop("p_is_similar_sum")
    parameters.pop("p_is_show_accum_amount")
    parameters["p_is_sum_same_entry"] = True
    parameters["p_is_show_accum"] = True
    parameters["p_is_similar_branch"] = False
    parameters["p_language"] = "vi"
    template["parameters"] = worker._encode_parameters(parameters)
    template["reportList"]["version"] = 0
    template["reportList"].pop("rp_function_name_async", None)
    template["reportList"]["timeout_interval"] = -1

    body = worker._request_body(
        template,
        {
            "dataset_type": "GENERAL_LEDGER",
            "period_from": date(2026, 1, 1),
            "period_to": date(2026, 1, 31),
            "page_size": 1000,
        },
        True,
    )
    upgraded = worker._decode_parameters(body["parameters"])

    assert upgraded["p_is_similar_sum"] is False
    assert upgraded["p_is_show_accum_amount"] is False
    assert "p_is_sum_same_entry" not in upgraded
    assert "p_is_show_accum" not in upgraded
    assert "p_is_similar_branch" not in upgraded
    assert "p_language" not in upgraded
    assert body["reportList"]["version"] == 2
    assert body["reportList"]["rp_function_name_async"] == "func_rpt_gl_general_diary_book_v4"
    assert "timeout_interval" not in body["reportList"]


def test_misa_trial_balance_request_upgrades_a_saved_legacy_template() -> None:
    from datetime import date

    template = worker._trial_balance_request_template()
    template["columns"] = "[]"
    template["report_id"] = "legacy"
    for field in (
        "is_favorite",
        "is_disabled",
        "script_summary",
        "preview_image",
        "is_created_from_old_db",
    ):
        template["reportList"].pop(field, None)

    body = worker._request_body(
        template,
        {
            "dataset_type": "TRIAL_BALANCE",
            "period_from": date(2026, 1, 1),
            "period_to": date(2026, 1, 31),
            "page_size": 100,
        },
        True,
    )

    assert body["report_id"] == "QWNjdW1HTEYwMV9ETg=="
    assert len(json.loads(body["columns"])) == 10
    assert body["reportList"]["is_favorite"] is False
    assert body["reportList"]["is_disabled"] is False
    assert body["reportList"]["script_summary"] == worker.TRIAL_BALANCE_SCRIPT_SUMMARY
    assert body["reportList"]["preview_image"] == "BẢNG CÂN ĐỐI TÀI KHOẢN ( Mẫu quản trị ).png"
    assert body["reportList"]["is_created_from_old_db"] is False


def test_misa_missing_session_diagnostic_only_exposes_response_shape() -> None:
    diagnostic = worker._misa_session_diagnostic(
        {"Code": 210, "Data": {"Total": 0, "tenant": "không được lộ"}}
    )
    assert "mã 210" in diagnostic
    assert "Total" in diagnostic
    assert "tenant" in diagnostic
    assert "không được lộ" not in diagnostic


def test_prisma_only_database_parameters_are_removed_for_psycopg() -> None:
    normalized = worker._psycopg_database_url(
        "postgresql://user:pass@localhost/db?schema=public&connection_limit=5&sslmode=require"
    )
    assert "schema=" not in normalized
    assert "connection_limit=" not in normalized
    assert "sslmode=require" in normalized


def test_complete_month_range_supports_month_and_year_reports() -> None:
    from datetime import date

    assert worker._is_complete_month_range(date(2026, 1, 1), date(2026, 1, 31))
    assert worker._is_complete_month_range(date(2026, 1, 1), date(2026, 12, 31))
    assert not worker._is_complete_month_range(date(2026, 1, 2), date(2026, 1, 31))


def test_report_selection_keeps_the_explicit_snapshot_ids() -> None:
    from datetime import date

    snapshot_id = "6ae6d9fe-9fb5-4705-b0f3-a9e9dbcf16e4"
    selection = worker._report_selection(
        worker.ReportOutputRequest(
            versionId="CFG-TEST",
            reportKey="B02",
            snapshotIds=[snapshot_id, snapshot_id],
            periodFrom=date(2026, 1, 1),
            periodTo=date(2026, 1, 31),
        )
    )

    assert selection["snapshotIds"] == [snapshot_id]
    assert selection["snapshotIdsByOffset"] == {"0": [snapshot_id]}
    assert selection["sourceKeys"] == []


def test_report_selection_keeps_comparison_snapshot_selection_closed() -> None:
    from datetime import date

    current_snapshot_id = "6ae6d9fe-9fb5-4705-b0f3-a9e9dbcf16e4"
    prior_snapshot_id = "90bc6db1-2ada-4192-a4fd-c518a40bb404"
    selection = worker._report_selection(
        worker.ReportOutputRequest(
            versionId="CFG-TEST",
            snapshotIds=[current_snapshot_id],
            snapshotIdsByOffset={"-1": [prior_snapshot_id]},
            periodFrom=date(2026, 1, 1),
            periodTo=date(2026, 1, 31),
        )
    )

    assert selection["snapshotIdsByOffset"] == {
        "-1": [prior_snapshot_id],
        "0": [current_snapshot_id],
    }


def test_report_selection_rejects_an_invalid_snapshot_id() -> None:
    from datetime import date
    from fastapi import HTTPException

    try:
        worker._report_selection(
            worker.ReportOutputRequest(
                versionId="CFG-TEST",
                snapshotIds=["not-a-snapshot"],
                periodFrom=date(2026, 1, 1),
                periodTo=date(2026, 1, 31),
            )
        )
    except HTTPException as error:
        assert error.status_code == 400
    else:
        raise AssertionError("Snapshot id không hợp lệ phải bị chặn")


def test_snapshot_selection_rejects_two_versions_of_one_source_slot() -> None:
    from fastapi import HTTPException

    snapshots = [
        {
            "snapshot_id": "old", "source_key": "MISA_TB", "company_key": "NTV",
            "data_space_key": "SPACE_2026", "dataset_type": "TRIAL_BALANCE",
            "period_from": "2026-01-01", "period_to": "2026-01-31",
        },
        {
            "snapshot_id": "new", "source_key": "MISA_TB", "company_key": "NTV",
            "data_space_key": "SPACE_2026", "dataset_type": "TRIAL_BALANCE",
            "period_from": "2026-01-01", "period_to": "2026-01-31",
        },
    ]
    try:
        worker._reject_overlapping_snapshot_versions(snapshots)
    except HTTPException as error:
        assert error.status_code == 409
    else:
        raise AssertionError("Hai phiên của cùng nguồn không được cộng chung")


def test_snapshot_selection_rejects_two_data_spaces_for_one_company_period() -> None:
    from fastapi import HTTPException

    snapshots = [
        {
            "snapshot_id": "trial", "source_key": "MISA_TB_TEST", "company_key": "TVG",
            "data_space_key": "GLOBAL_26_TEST", "dataset_type": "TRIAL_BALANCE",
            "period_from": "2026-01-01", "period_to": "2026-01-31",
        },
        {
            "snapshot_id": "journal", "source_key": "MISA_GL_REAL", "company_key": "TVG",
            "data_space_key": "GLOBAL_2026", "dataset_type": "GENERAL_LEDGER",
            "period_from": "2026-01-01", "period_to": "2026-01-31",
        },
    ]
    try:
        worker._reject_mixed_data_spaces_per_company(snapshots)
    except HTTPException as error:
        assert error.status_code == 409
        assert "nhiều khối dữ liệu" in str(error.detail)
    else:
        raise AssertionError("Một công ty không được lấy số từ hai khối dữ liệu cùng kỳ")


def test_snapshot_selection_allows_one_data_space_for_each_consolidated_company() -> None:
    snapshots = [
        {
            "snapshot_id": "tvg", "source_key": "MISA_TVG", "company_key": "TVG",
            "data_space_key": "GLOBAL_2026", "dataset_type": "TRIAL_BALANCE",
            "period_from": "2026-01-01", "period_to": "2026-01-31",
        },
        {
            "snapshot_id": "ntv", "source_key": "MISA_NTV", "company_key": "NTV",
            "data_space_key": "NTV_2026", "dataset_type": "TRIAL_BALANCE",
            "period_from": "2026-01-01", "period_to": "2026-01-31",
        },
    ]

    worker._reject_mixed_data_spaces_per_company(snapshots)


def test_misa_data_space_reuses_database_identity_across_environments() -> None:
    class Result:
        def fetchone(self):
            return {"data_space_key": "MISA_EXISTING_SPACE"}

    class Connection:
        params = None

        def execute(self, _query, params):
            self.params = params
            return Result()

    connection = Connection()
    resolved = worker._compatible_data_space_key(
        connection,
        "NTV",
        {
            "tenantId": "tenant-1",
            "databaseId": "database-1",
            "workingBook": 0,
            "dbType": 1,
            "environment": "g2",
        },
        "MISA_NEW_SPACE",
    )

    assert resolved == "MISA_EXISTING_SPACE"
    assert connection.params == ("NTV", "tenant-1", "database-1", "0", "1")


def test_journal_period_filters_use_vietnam_accounting_date() -> None:
    import inspect

    source = inspect.getsource(worker)
    assert "posted_at AT TIME ZONE 'Asia/Ho_Chi_Minh'" in source
    assert "posted_at::date BETWEEN" not in source


def test_cleanup_only_accepts_explicit_test_data_space_labels() -> None:
    assert worker._has_test_data_space_marker("Năm 2026 Test", "BCTC demo")
    assert worker._has_test_data_space_marker("Bản nháp tháng 01")
    assert worker._has_test_data_space_marker("Khối thử nghiệm nội bộ")
    assert not worker._has_test_data_space_marker("Dữ liệu kế toán năm 2026", "BCTC chính thức")


def test_cleanup_confirmation_names_every_selected_space() -> None:
    phrase = worker._cleanup_confirmation_phrase([
        {"dataSpaceName": "Năm 2025 Test", "companyKey": "NTV"},
        {"dataSpaceName": "Năm 2026 Demo", "companyKey": "GLOBAL"},
    ])
    assert phrase == "XÓA: Năm 2025 Test [NTV] | Năm 2026 Demo [GLOBAL]"


def test_cleanup_target_list_is_deduplicated() -> None:
    targets = worker._cleanup_targets([
        worker.TestDataSpaceTarget(companyKey="ntv", dataSpaceKey="space-test"),
        worker.TestDataSpaceTarget(companyKey="NTV", dataSpaceKey="SPACE-TEST"),
    ])
    assert targets == [{"companyKey": "NTV", "dataSpaceKey": "SPACE-TEST"}]


def test_cleanup_count_query_has_one_parameter_pair_per_subquery() -> None:
    import inspect

    source = inspect.getsource(worker._test_data_space_counts)
    assert source.count("%s") == 34
    assert "(company_key, data_space_key) * 17" in source
    assert "accounting_report_result_snapshots" in source


def test_cleanup_confirmation_supports_vietnamese_text() -> None:
    phrase = "XÓA: Năm 2026 Thử nghiệm [NTV]"
    assert worker._secure_text_equal(phrase, phrase)
    assert not worker._secure_text_equal(phrase, "XÓA: Năm 2026 Demo [NTV]")


def test_imported_prior_report_overlays_only_present_cells() -> None:
    output = {
        "columns": [
            {"columnKey": "LABEL", "columnType": "LABEL", "periodOffset": 0},
            {"columnKey": "CURRENT", "columnType": "CURRENT_PERIOD", "periodOffset": 0},
            {"columnKey": "PREVIOUS", "columnType": "PREVIOUS_PERIOD", "periodOffset": -1},
        ],
        "rows": [
            {"rowCode": "01", "cells": [{"text": "A"}, {"value": 200}, {"value": 100}]},
            {"rowCode": "02", "cells": [{"text": "B"}, {"value": 40}, {"value": 30}]},
        ],
    }
    imported = {
        "report_result_id": "11111111-1111-1111-1111-111111111111",
        "report_key": "BAO_CAO_KET_QUA_KINH_DOANH",
        "report_name": "KQKD",
        "company_key": "TVG",
        "data_space_key": "SPACE",
        "data_space_name": "TVD_2025",
        "period_from": date(2025, 1, 1),
        "period_to": date(2025, 1, 31),
        "source_file_name": "KQKD 2025.xlsx",
        "source_sheet_name": "KQKD",
        "output_data": {
            "columns": [
                {"columnKey": "LABEL", "columnType": "LABEL", "periodOffset": 0},
                {"columnKey": "CURRENT", "columnType": "CURRENT_PERIOD", "periodOffset": 0},
            ],
            "rows": [
                {"rowCode": "01", "cells": [{"text": "A"}, {"value": 135}]},
                {"rowCode": "02", "cells": [{"text": "B"}, {"value": None}]},
            ],
        },
    }
    result = worker._apply_imported_report_results(output, {-1: [imported]})
    assert result["rows"][0]["cells"][2]["value"] == 135
    assert result["rows"][0]["cells"][1]["value"] == 200
    assert result["rows"][1]["cells"][2]["value"] == 30
    assert result["rows"][0]["cells"][2]["formulaStatus"] == "IMPORTED_REPORT_RESULT"
    assert result["importedReportResults"][0]["targetPeriodOffset"] == -1


def test_b09_import_reads_exact_total_cells_from_disclosure_map(monkeypatch) -> None:
    report_key = worker.DISCLOSURE_REPORT_KEY
    records = [
        {
            "recordType": "REPORT",
            "entityKey": report_key,
            "displayName": "Bản thuyết minh Báo cáo tài chính",
            "sortOrder": 1,
            "payload": {
                "reportKey": report_key,
                "reportName": "Bản thuyết minh Báo cáo tài chính",
                "displayTitle": "BẢN THUYẾT MINH BÁO CÁO TÀI CHÍNH",
            },
        },
        {
            "recordType": "REPORT_COLUMN",
            "entityKey": f"{report_key}:LABEL",
            "displayName": "Khoản mục",
            "sortOrder": 10,
            "payload": {"reportKey": report_key, "columnKey": "LABEL", "columnType": "LABEL"},
        },
        {
            "recordType": "REPORT_COLUMN",
            "entityKey": f"{report_key}:CODE",
            "displayName": "Mã",
            "sortOrder": 20,
            "payload": {"reportKey": report_key, "columnKey": "CODE", "columnType": "ROW_CODE"},
        },
        {
            "recordType": "REPORT_COLUMN",
            "entityKey": f"{report_key}:CURRENT",
            "displayName": "Năm nay",
            "sortOrder": 30,
            "payload": {
                "reportKey": report_key,
                "columnKey": "CURRENT",
                "columnType": "CURRENT_PERIOD",
                "columnLabel": "Năm nay",
                "periodOffset": 0,
            },
        },
        {
            "recordType": "REPORT_COLUMN",
            "entityKey": f"{report_key}:PRIOR",
            "displayName": "Năm trước",
            "sortOrder": 40,
            "payload": {
                "reportKey": report_key,
                "columnKey": "PRIOR",
                "columnType": "PREVIOUS_PERIOD",
                "columnLabel": "Năm trước",
                "periodOffset": -1,
            },
        },
        {
            "recordType": "REPORT_ROW",
            "entityKey": f"{report_key}:V.1",
            "displayCode": "V.1",
            "displayName": "Tiền và tương đương tiền",
            "sortOrder": 100,
            "payload": {
                "reportKey": report_key,
                "rowCode": "V.1",
                "rowLabel": "Tiền và tương đương tiền",
                "rowType": "DATA",
                "disclosureItemKey": f"{report_key}:V.1",
            },
        },
        {
            "recordType": "DISCLOSURE_TEMPLATE",
            "entityKey": report_key,
            "sortOrder": 1,
            "sourceSheet": "B09",
            "payload": {"templateKey": report_key, "sourceSheet": "B09"},
        },
        {
            "recordType": "DISCLOSURE_SECTION",
            "entityKey": f"{report_key}:V",
            "sortOrder": 1,
            "payload": {
                "templateKey": report_key,
                "startRow": 1,
                "endRow": 12,
                "cells": [
                    {"address": "A1", "row": 1, "column": 1, "role": "TITLE", "text": "BẢN THUYẾT MINH BÁO CÁO TÀI CHÍNH"},
                    {"address": "D5", "row": 5, "column": 4, "role": "HEADER", "text": "Năm nay"},
                    {"address": "E5", "row": 5, "column": 5, "role": "HEADER", "text": "Năm trước"},
                    {"address": "A10", "row": 10, "column": 1, "role": "LABEL", "text": "Cộng"},
                    {"address": "D10", "row": 10, "column": 4, "role": "REFERENCE_VALUE"},
                    {"address": "E10", "row": 10, "column": 5, "role": "REFERENCE_VALUE"},
                ],
            },
        },
        {
            "recordType": "DISCLOSURE_ITEM",
            "entityKey": f"{report_key}:V.1",
            "sortOrder": 1,
            "payload": {
                "templateKey": report_key,
                "itemKey": f"{report_key}:V.1",
                "noteCode": "V.1",
                "startRow": 6,
                "endRow": 10,
            },
        },
    ]
    monkeypatch.setattr(worker, "_load_config_records", lambda version_id: records)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "B09"
    sheet["A1"] = "BẢN THUYẾT MINH BÁO CÁO TÀI CHÍNH"
    sheet["D5"] = "Năm nay"
    sheet["E5"] = "Năm trước"
    sheet["A10"] = "Cộng"
    sheet["D10"] = 125_000
    sheet["E10"] = 90_000
    stream = io.BytesIO()
    workbook.save(stream)
    parsed = load_workbook(io.BytesIO(stream.getvalue()), read_only=True, data_only=True)
    try:
        preview = worker._accounting_report_file_preview(
            parsed,
            worker.AccountingFilePreviewRequest(
                fileName="B09 2025.xlsx",
                fileDataUrl="unused",
                companyKey="TVG",
                companyLabel="TVG",
                datasetType="REPORT_RESULT",
                versionId="TEST",
                reportKey=report_key,
            ),
            "B09 2025.xlsx",
            date(2025, 1, 1),
            date(2025, 12, 31),
        )
    finally:
        parsed.close()

    assert preview["reportKey"] == report_key
    assert preview["rowCount"] == 1
    assert preview["reportOutput"]["controlStatus"] == "READY"
    row = preview["reportOutput"]["rows"][0]
    assert row["cells"][2]["value"] == 125_000
    assert row["cells"][3]["value"] == 90_000
    assert row["cells"][2]["evidence"]["components"][0]["tracePath"] == "File B09 → ô D10 → V.1"


def test_b09_import_does_not_turn_missing_template_cells_into_zero(monkeypatch) -> None:
    report_key = worker.DISCLOSURE_REPORT_KEY
    records = [
        {"recordType": "REPORT", "entityKey": report_key, "sortOrder": 1, "payload": {"reportKey": report_key, "reportName": "B09"}},
        {"recordType": "REPORT_COLUMN", "entityKey": "LABEL", "sortOrder": 1, "payload": {"reportKey": report_key, "columnKey": "LABEL", "columnType": "LABEL"}},
        {"recordType": "REPORT_COLUMN", "entityKey": "CURRENT", "sortOrder": 2, "payload": {"reportKey": report_key, "columnKey": "CURRENT", "columnType": "CURRENT_PERIOD"}},
        {"recordType": "REPORT_ROW", "entityKey": "V.1", "sortOrder": 3, "payload": {"reportKey": report_key, "rowCode": "V.1", "rowLabel": "Tiền", "disclosureItemKey": f"{report_key}:V.1"}},
        {"recordType": "DISCLOSURE_TEMPLATE", "entityKey": report_key, "sortOrder": 1, "payload": {"templateKey": report_key, "sourceSheet": "B09"}},
        {"recordType": "DISCLOSURE_SECTION", "entityKey": "V", "sortOrder": 1, "payload": {"templateKey": report_key, "startRow": 1, "endRow": 10, "cells": [
            {"address": "A1", "row": 1, "column": 1, "role": "TITLE", "text": "BẢN THUYẾT MINH BÁO CÁO TÀI CHÍNH"},
            {"address": "D5", "row": 5, "column": 4, "role": "HEADER", "text": "Năm nay"},
            {"address": "D10", "row": 10, "column": 4, "role": "REFERENCE_VALUE"},
        ]}},
        {"recordType": "DISCLOSURE_ITEM", "entityKey": f"{report_key}:V.1", "sortOrder": 1, "payload": {"templateKey": report_key, "itemKey": f"{report_key}:V.1", "noteCode": "V.1", "startRow": 6, "endRow": 10}},
    ]
    monkeypatch.setattr(worker, "_load_config_records", lambda version_id: records)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "B09"
    sheet["A1"] = "BẢN THUYẾT MINH BÁO CÁO TÀI CHÍNH"
    stream = io.BytesIO()
    workbook.save(stream)
    parsed = load_workbook(io.BytesIO(stream.getvalue()), read_only=True, data_only=True)
    try:
        preview = worker._accounting_report_file_preview(
            parsed,
            worker.AccountingFilePreviewRequest(
                fileName="B09.xlsx",
                fileDataUrl="unused",
                companyKey="TVG",
                datasetType="REPORT_RESULT",
                versionId="TEST",
                reportKey=report_key,
            ),
            "B09.xlsx",
            date(2025, 1, 1),
            date(2025, 12, 31),
        )
    finally:
        parsed.close()

    cell = preview["reportOutput"]["rows"][0]["cells"][1]
    assert cell["value"] is None
    assert cell["formulaStatus"] == "MISSING"
    assert preview["reportOutput"]["controlStatus"] == "NEEDS_REVIEW"
