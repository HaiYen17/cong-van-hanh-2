from __future__ import annotations

from datetime import date
from io import BytesIO

from openpyxl import load_workbook

from disclosure_engine import (
    DISCLOSURE_REPORT_KEY,
    build_disclosure_output,
    disclosure_output_to_xlsx,
)


def record(record_type: str, entity_key: str, payload: dict, sort_order: int = 0) -> dict:
    return {
        "recordId": f"{record_type}:{entity_key}:{sort_order}",
        "recordType": record_type,
        "entityKey": entity_key,
        "sortOrder": sort_order,
        "payload": payload,
    }


def disclosure_records(
    *,
    item_key: str,
    automation_level: str = "AUTOMATIC",
    links: list[dict],
) -> list[dict]:
    return [
        record(
            "DISCLOSURE_TEMPLATE",
            DISCLOSURE_REPORT_KEY,
            {
                "templateKey": DISCLOSURE_REPORT_KEY,
                "templateName": "B09-DN",
            },
        ),
        record(
            "DISCLOSURE_SECTION",
            "SECTION_VI",
            {
                "templateKey": DISCLOSURE_REPORT_KEY,
                "sectionKey": "SECTION_VI",
                "sectionCode": "VI",
                "sectionTitle": "Thông tin bổ sung",
                "startRow": 1,
            },
        ),
        record(
            "DISCLOSURE_ITEM",
            item_key,
            {
                "templateKey": DISCLOSURE_REPORT_KEY,
                "sectionKey": "SECTION_VI",
                "itemKey": item_key,
                "noteCode": "VI.6",
                "noteTitle": "Chi phí tài chính",
                "automationLevel": automation_level,
                "startRow": 2,
            },
        ),
        *[
            record(
                "DISCLOSURE_LINK",
                f"{item_key}:{index}",
                {"disclosureItemKey": item_key, **link},
                index,
            )
            for index, link in enumerate(links, start=1)
        ],
    ]


def evidence(value: float, *, complete: bool = True, account: str = "635") -> dict:
    components = (
        [
            {
                "label": f"Tài khoản {account}",
                "accountNumber": account,
                "signedValue": value,
                "sourceName": "Bảng cân đối tài khoản",
            }
        ]
        if complete
        else []
    )
    return {
        "value": value,
        "components": components,
        "isComplete": complete,
        "difference": 0 if complete else value,
    }


def report_output(
    report_key: str,
    rows: list[tuple[str, str, list[float]]],
    *,
    column_keys: list[str] | None = None,
    column_types: list[str] | None = None,
    complete: bool = True,
    control_status: str = "READY",
) -> dict:
    keys = column_keys or ["CURRENT", "PREVIOUS"]
    types = column_types or ["CURRENT_PERIOD", "PREVIOUS_PERIOD"]
    columns = [
        {"columnKey": "LABEL", "columnLabel": "Chỉ tiêu", "columnType": "LABEL"},
        *[
            {
                "columnKey": key,
                "columnLabel": key.replace("_", " ").title(),
                "columnType": types[index],
            }
            for index, key in enumerate(keys)
        ],
    ]
    output_rows = []
    for row_code, row_label, values in rows:
        cells = [{"value": row_label}]
        cells.extend(
            {
                "value": value,
                "evidence": evidence(value, complete=complete, account=f"635{index + 1}"),
            }
            for index, value in enumerate(values)
        )
        output_rows.append({"rowCode": row_code, "rowLabel": row_label, "cells": cells})
    return {
        "report": {"reportKey": report_key, "reportName": report_key},
        "columns": columns,
        "rows": output_rows,
        "controlStatus": control_status,
    }


def build(records: list[dict], reports: dict[str, dict]) -> dict:
    return build_disclosure_output(
        records=records,
        period_from=date(2026, 1, 1),
        period_to=date(2026, 1, 31),
        selection={"snapshotIdsByOffset": {"0": ["snapshot-current"]}},
        resolve_report=lambda report_key: reports[report_key],
    )


def test_subset_detail_is_shown_but_never_added_twice() -> None:
    item_key = "NOTE_VI_6"
    records = disclosure_records(
        item_key=item_key,
        links=[
            {
                "reportKey": "INCOME_STATEMENT",
                "reportRowCode": "23",
                "linkRole": "PRIMARY_NOTE_REFERENCE",
                "aggregationRole": "ADDITIVE",
            },
            {
                "reportKey": "INCOME_STATEMENT",
                "reportRowCode": "24",
                "linkRole": "PRIMARY_NOTE_REFERENCE",
            },
        ],
    )
    output = build(
        records,
        {
            "INCOME_STATEMENT": report_output(
                "INCOME_STATEMENT",
                [
                    ("23", "Chi phí tài chính", [100, 80]),
                    ("24", "Trong đó: Chi phí lãi vay", [40, 30]),
                ],
            )
        },
    )

    item = output["items"][0]
    assert [total["value"] for total in item["totals"]] == [100, 80]
    assert item["links"][1]["aggregationRole"] == "SUBSET_DETAIL"
    assert item["status"] == "READY"
    assert all(control["status"] == "PASS" for control in item["controlChecks"])


def test_debit_and_credit_columns_with_same_semantic_type_stay_separate() -> None:
    item_key = "NOTE_BALANCE"
    records = disclosure_records(
        item_key=item_key,
        links=[
            {
                "reportKey": "BALANCE_SHEET",
                "reportRowCode": "141",
                "linkRole": "PRIMARY_NOTE_REFERENCE",
                "aggregationRole": "ADDITIVE",
            }
        ],
    )
    output = build(
        records,
        {
            "BALANCE_SHEET": report_output(
                "BALANCE_SHEET",
                [("141", "Hàng tồn kho", [120, 20])],
                column_keys=["CURRENT_DEBIT", "CURRENT_CREDIT"],
                column_types=["CURRENT_PERIOD", "CURRENT_PERIOD"],
            )
        },
    )

    totals = output["items"][0]["totals"]
    assert [(total["columnKey"], total["value"]) for total in totals] == [
        ("CURRENT_DEBIT", 120),
        ("CURRENT_CREDIT", 20),
    ]


def test_b09_shared_formula_is_the_value_source_and_links_only_reconcile() -> None:
    item_key = "NOTE_VI_6"
    records = disclosure_records(
        item_key=item_key,
        links=[
            {
                "reportKey": "INCOME_STATEMENT",
                "reportRowCode": "23",
                "linkRole": "PRIMARY_NOTE_REFERENCE",
                "aggregationRole": "ADDITIVE",
            }
        ],
    )
    records.append(
        record(
            "REPORT",
            DISCLOSURE_REPORT_KEY,
            {
                "reportKey": DISCLOSURE_REPORT_KEY,
                "reportName": "B09-DN",
            },
        )
    )
    formula = report_output(
        DISCLOSURE_REPORT_KEY,
        [("VI.6", "Chi phí tài chính", [777, 666])],
    )
    for row in formula["rows"]:
        for cell in row["cells"][1:]:
            cell["kind"] = "VALUE"
            cell["formulaStatus"] = "READY"
    output = build(
        records,
        {
            "INCOME_STATEMENT": report_output(
                "INCOME_STATEMENT",
                [("23", "Chi phí tài chính", [100, 80])],
            ),
            DISCLOSURE_REPORT_KEY: formula,
        },
    )

    item = output["items"][0]
    assert [total["value"] for total in item["totals"]] == [777, 666]
    assert [value["value"] for value in item["links"][0]["values"]] == [100, 80]


def test_reference_reports_are_compared_by_period_semantics() -> None:
    item_key = "NOTE_CASH"
    records = disclosure_records(
        item_key=item_key,
        links=[
            {
                "reportKey": "CASH_FLOW_DIRECT",
                "reportRowCode": "50",
                "linkRole": "RECONCILIATION_ONLY",
                "aggregationRole": "REFERENCE_ONLY",
            },
            {
                "reportKey": "CASH_FLOW_INDIRECT",
                "reportRowCode": "50",
                "linkRole": "RECONCILIATION_ONLY",
                "aggregationRole": "REFERENCE_ONLY",
            },
        ],
    )
    reports = {
        "CASH_FLOW_DIRECT": report_output(
            "CASH_FLOW_DIRECT",
            [("50", "Lưu chuyển tiền thuần", [250, 200])],
            column_keys=["DIRECT_CURRENT", "DIRECT_PREVIOUS"],
        ),
        "CASH_FLOW_INDIRECT": report_output(
            "CASH_FLOW_INDIRECT",
            [("50", "Lưu chuyển tiền thuần", [250, 200])],
            column_keys=["INDIRECT_CURRENT", "INDIRECT_PREVIOUS"],
        ),
    }

    output = build(records, reports)
    item = output["items"][0]
    assert len(item["controlChecks"]) == 2
    assert all(control["status"] == "PASS" for control in item["controlChecks"])
    assert item["status"] == "READY"

    reports["CASH_FLOW_INDIRECT"] = report_output(
        "CASH_FLOW_INDIRECT",
        [("50", "Lưu chuyển tiền thuần", [249, 200])],
        column_keys=["INDIRECT_CURRENT", "INDIRECT_PREVIOUS"],
    )
    mismatch = build(records, reports)
    failed = mismatch["items"][0]
    assert failed["status"] == "NEEDS_REVIEW"
    assert failed["controlChecks"][0]["status"] == "FAIL"
    assert failed["controlChecks"][0]["difference"] == 1


def test_incomplete_evidence_fails_without_creating_a_balancing_plug() -> None:
    item_key = "NOTE_UNPROVEN"
    records = disclosure_records(
        item_key=item_key,
        links=[
            {
                "reportKey": "INCOME_STATEMENT",
                "reportRowCode": "01",
                "linkRole": "PRIMARY_NOTE_REFERENCE",
                "aggregationRole": "ADDITIVE",
            }
        ],
    )
    output = build(
        records,
        {
            "INCOME_STATEMENT": report_output(
                "INCOME_STATEMENT",
                [("01", "Doanh thu", [500, 400])],
                complete=False,
            )
        },
    )

    item = output["items"][0]
    assert item["status"] == "NEEDS_REVIEW"
    assert all(control["status"] == "FAIL" for control in item["controlChecks"])
    assert item["totals"][0]["evidence"]["components"] == []
    assert "không tự tạo số bù" in item["totals"][0]["evidence"]["warning"]


def test_disclosure_xlsx_keeps_items_without_automatic_values() -> None:
    item_key = "NOTE_MANUAL"
    output = build(
        disclosure_records(item_key=item_key, automation_level="MANUAL", links=[]),
        {},
    )

    workbook = load_workbook(BytesIO(disclosure_output_to_xlsx(output)), data_only=True)
    sheet = workbook["B09-DN"]
    values = [cell.value for row in sheet.iter_rows() for cell in row if cell.value is not None]

    assert "VI.6" in values
    assert "Chi phí tài chính" in values
    assert "Chưa có số tự động" in values
    assert "Kế toán nhập" in values
