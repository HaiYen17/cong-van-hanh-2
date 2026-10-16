from __future__ import annotations

from datetime import date
from io import BytesIO

import pytest
from openpyxl import load_workbook

from report_engine import (
    ReportEngineError,
    _account_matches,
    _filters_match,
    _rows_for_set,
    _source_rows_for_term,
    _source_matches,
    build_output,
    evaluate_expression,
    note_required_offsets,
    normalize_key,
    output_to_xlsx,
    required_offsets,
    shift_year,
)


def record(
    record_id: str,
    record_type: str,
    entity_key: str,
    payload: dict,
    sort_order: int = 0,
    display_name: str = "",
) -> dict:
    return {
        "recordId": record_id,
        "recordType": record_type,
        "entityKey": entity_key,
        "displayCode": payload.get("rowCode", entity_key),
        "displayName": display_name or entity_key,
        "sortOrder": sort_order,
        "payload": payload,
    }


def approved_adjustment_cell(report_key: str, row_code: str, column_key: str = "CURRENT") -> dict:
    return {
        "operator": "MANUAL_ADJUSTMENT",
        "sourceKey": "BCTC_DIEU_CHINH",
        "datasetType": "MANUAL_ADJUSTMENT",
        "filters": {"report_key": report_key, "row_code": row_code},
        "_columnKey": column_key,
    }


def approved_adjustment_row(
    report_key: str,
    row_code: str,
    amount: float,
    column_key: str = "CURRENT",
) -> dict:
    return {
        "source_key": "BCTC_DIEU_CHINH",
        "base_source_key": "BCTC_DIEU_CHINH",
        "dataset_type": "MANUAL_ADJUSTMENT",
        "canonical_data": {
            "report_key": report_key,
            "row_code": row_code,
            "column_key": column_key,
            "amount": amount,
            "approval_status": "APPROVED",
            "reason": "Dữ liệu kiểm thử có lịch sử",
            "source_reference": "TEST-FIXTURE",
        },
    }


def test_normalize_key_and_shift_year_keep_legacy_behavior() -> None:
    assert normalize_key("Báo cáo: Tiền và tương đương tiền") == "BAO_CAO:_TIEN_VA_TUONG_DUONG_TIEN"
    assert shift_year(date(2024, 2, 29), -1) == date(2023, 2, 28)


def test_structured_formula_conditions_are_generic_and_ignore_person_names() -> None:
    filters = {
        "logic": "ANY",
        "conditions": [
            {"field": "cash_flow_code", "operator": "EQUALS_ANY", "value": ["14"]},
            {
                "field": "search_text",
                "operator": "CONTAINS_ANY",
                "value": ["lãi vay", "interest"],
            },
        ],
    }
    interest_payment = {
        "cash_flow_code": "",
        "description": "Thanh toán lãi vay tháng 12/2025",
        "account_object_name": "NGÂN HÀNG A",
    }
    unrelated_person = {
        "cash_flow_code": "",
        "description": "Khách hàng thanh toán tiền mua hàng",
        "account_object_name": "Anh Bình - cửa hàng 562",
    }
    coded_payment = {"cash_flow_code": "14", "description": "UNC ngân hàng"}

    assert _filters_match(interest_payment, filters, []) is True
    assert _filters_match(coded_payment, filters, []) is True
    assert _filters_match(unrelated_person, filters, []) is False


def test_structured_formula_conditions_support_all_and_empty_checks() -> None:
    filters = {
        "logic": "ALL",
        "conditions": [
            {"field": "cash_flow_code", "operator": "IS_EMPTY", "value": []},
            {"field": "expense_item_code", "operator": "STARTS_WITH_ANY", "value": ["CP"]},
        ],
    }
    assert _filters_match({"cash_flow_code": "", "expense_item_code": "CP01"}, filters, []) is True
    assert _filters_match({"cash_flow_code": "14", "expense_item_code": "CP01"}, filters, []) is False


def test_journal_term_uses_detail_only_for_text_conditions_and_monthly_for_other_terms() -> None:
    rows = [
        {
            "source_key": "GL",
            "dataset_type": "GENERAL_LEDGER",
            "canonical_data": {
                "journal_granularity": "MONTHLY",
                "debit_account": "335",
                "credit_account": "1121",
                "amount": 30,
            },
        },
        {
            "source_key": "GL",
            "dataset_type": "GENERAL_LEDGER",
            "canonical_data": {
                "journal_granularity": "DETAIL",
                "debit_account": "335",
                "credit_account": "1121",
                "amount": 30,
                "description": "Thanh toán lãi vay",
            },
        },
    ]
    base_term = {"sourceKey": "GL", "datasetType": "GENERAL_LEDGER"}
    detail_term = {
        **base_term,
        "filters": {
            "conditions": [
                {
                    "field": "search_text",
                    "operator": "CONTAINS_ANY",
                    "value": ["lãi vay"],
                }
            ]
        },
    }

    assert [row["canonical_data"]["journal_granularity"] for row in _source_rows_for_term(rows, base_term, [])] == ["MONTHLY"]
    assert [row["canonical_data"]["journal_granularity"] for row in _source_rows_for_term(rows, detail_term, [])] == ["DETAIL"]


def test_cash_flow_line_11_excludes_interest_payable_without_period_hardcode() -> None:
    report_key = "CF_INDIRECT"
    records = [
        record("report", "REPORT", report_key, {"reportName": "Lưu chuyển tiền tệ gián tiếp"}),
        record("label", "REPORT_COLUMN", "CF_LABEL", {"reportKey": report_key, "columnKey": "LABEL", "columnType": "LABEL", "columnLabel": "Chỉ tiêu"}, 1),
        record("current", "REPORT_COLUMN", "CF_CURRENT", {"reportKey": report_key, "columnKey": "CURRENT", "columnType": "CURRENT_PERIOD", "columnLabel": "Kỳ này", "valueBasis": "PERIOD"}, 2),
        record("row-06", "REPORT_ROW", "CF_06", {"reportKey": report_key, "rowCode": "06", "rowLabel": "Chi phí đi vay", "rowType": "DATA", "cells": {"CURRENT": {"operator": "MOVEMENT", "sourceKey": "GL", "datasetType": "GENERAL_LEDGER", "accountSetKey": "INTEREST_EXPENSE", "side": "DEBIT", "periodBasis": "PERIOD", "useColumnBasis": False}}}, 10),
        record("row-14", "REPORT_ROW", "CF_14", {"reportKey": report_key, "rowCode": "14", "rowLabel": "Chi phí đi vay đã trả", "rowType": "DATA", "cells": {"CURRENT": {"operator": "CASH_MOVEMENT", "sourceKey": "GL", "datasetType": "GENERAL_LEDGER", "accountSetKey": "CASH", "counterAccountSetKey": "INTEREST_PAYMENT_COUNTER", "side": "CREDIT", "sign": -1, "periodBasis": "PERIOD", "useColumnBasis": False, "filters": {"logic": "ANY", "conditions": [{"field": "cash_flow_code", "operator": "EQUALS_ANY", "value": ["14"]}, {"field": "search_text", "operator": "CONTAINS_ANY", "value": ["lãi vay", "interest"]}]}}}}, 20),
        record("row-11", "REPORT_ROW", "CF_11", {"reportKey": report_key, "rowCode": "11", "rowLabel": "Tăng giảm các khoản phải trả, không kể lãi vay phải trả", "rowType": "DATA", "cells": {"CURRENT": {"terms": [{"operator": "DELTA_BALANCE", "sourceKey": "TB", "datasetType": "TRIAL_BALANCE", "accountSetKey": "PAYABLES", "side": "CREDIT", "sign": 1}, {"operator": "LINE_REFERENCE", "targetRowCode": "06", "sign": -1}, {"operator": "LINE_REFERENCE", "targetRowCode": "14", "sign": -1}]}}}, 30),
        record("member-interest", "ACCOUNT_SET_MEMBER", "INTEREST_EXPENSE_635", {"accountSetKey": "INTEREST_EXPENSE", "matchType": "PREFIX", "matchValue": "635", "role": "INCLUDE"}),
        record("member-cash", "ACCOUNT_SET_MEMBER", "CASH_112", {"accountSetKey": "CASH", "matchType": "PREFIX", "matchValue": "112", "role": "INCLUDE"}),
        record("member-counter", "ACCOUNT_SET_MEMBER", "INTEREST_PAYMENT_COUNTER_MIXED", {"accountSetKey": "INTEREST_PAYMENT_COUNTER", "matchType": "MIXED", "matchValue": "331*,335*,635*", "role": "INCLUDE"}),
        record("member-payable", "ACCOUNT_SET_MEMBER", "PAYABLES_MIXED", {"accountSetKey": "PAYABLES", "matchType": "MIXED", "matchValue": "331*,335*", "role": "INCLUDE"}),
    ]
    rows = [
        {"source_key": "TB", "dataset_type": "TRIAL_BALANCE", "company_key": "TVG", "account_number": "335", "is_leaf": True, "opening_credit": 500, "closing_credit": 1_500},
        {"source_key": "GL", "dataset_type": "GENERAL_LEDGER", "company_key": "TVG", "debit_account": "6351", "credit_account": "335", "amount": 120, "description": "Trích trước chi phí lãi vay"},
        {"source_key": "GL", "dataset_type": "GENERAL_LEDGER", "company_key": "TVG", "debit_account": "335", "credit_account": "1121", "amount": 30, "description": "Thanh toán lãi vay cho tổ chức A", "cash_flow_code": ""},
        {"source_key": "GL", "dataset_type": "GENERAL_LEDGER", "company_key": "TVG", "debit_account": "3311", "credit_account": "1121", "amount": 20, "description": "UNC ngân hàng", "cash_flow_code": "14"},
        {"source_key": "GL", "dataset_type": "GENERAL_LEDGER", "company_key": "TVG", "debit_account": "3311", "credit_account": "1121", "amount": 999, "description": "Thanh toán tiền hàng cho Anh Bình", "cash_flow_code": ""},
    ]

    output = build_output(
        records=records,
        report_key=report_key,
        period_from=date(2026, 1, 1),
        period_to=date(2026, 1, 31),
        rows_by_offset={0: rows},
        selection={"sourceKeys": ["TB", "GL"], "companyKeys": ["TVG"], "dataSpaceKeys": ["SPACE"]},
    )
    values = {row["rowCode"]: row["cells"][1]["value"] for row in output["rows"]}
    assert values["06"] == 120
    assert values["14"] == -50
    assert values["11"] == 930


def test_expression_parser_supports_precedence_and_rejects_unknown_characters() -> None:
    values = {"A": 10.0, "B": 2.0}
    assert evaluate_expression("=A + B * 3", values.get) == 16.0
    assert evaluate_expression("(A-B)/2", values.get) == 4.0
    with pytest.raises(ReportEngineError, match="ký tự"):
        evaluate_expression("__import__('os')", values.get)


def test_expression_parser_supports_legal_report_codes_with_suffixes() -> None:
    values = {"411a": 5_000, "411b": 250, "420a": 90, "420b": 10, "270_KHAC": 7}
    assert evaluate_expression("411a+411b", values.get) == 5_250
    assert evaluate_expression("420a+420b", values.get) == 100
    assert evaluate_expression("270_KHAC+3", values.get) == 10


def test_account_set_member_supports_mixed_exact_and_prefix_accounts() -> None:
    member = {
        "matchType": "MIXED",
        "matchValue": "151*, 152*, 1531, 1541, 1542; 1543\n155*, TK156*, 157*, 158*",
    }
    leaf = {"is_leaf": True}
    assert _account_matches("1512", member, leaf)
    assert _account_matches("1528", member, leaf)
    assert _account_matches("1531", member, leaf)
    assert _account_matches("1543", member, leaf)
    assert _account_matches("1569", member, leaf)
    assert not _account_matches("1532", member, leaf)
    assert not _account_matches("1544", member, leaf)
    assert not _account_matches("151", member, {"is_leaf": False})


def test_short_term_payable_group_includes_new_subaccounts_but_excludes_long_term_branch() -> None:
    members = [
        record(
            "include-331",
            "ACCOUNT_SET_MEMBER",
            "SHORT_PAYABLE_INCLUDE",
            {
                "accountSetKey": "SHORT_PAYABLE",
                "matchType": "PREFIX",
                "matchValue": "331",
                "role": "INCLUDE",
                "companyKey": "TVG",
            },
        ),
        record(
            "exclude-3312",
            "ACCOUNT_SET_MEMBER",
            "SHORT_PAYABLE_EXCLUDE",
            {
                "accountSetKey": "SHORT_PAYABLE",
                "matchType": "PREFIX",
                "matchValue": "3312",
                "role": "EXCLUDE",
                "companyKey": "TVG",
            },
        ),
    ]
    rows = [
        {"company_key": "TVG", "account_number": "3311", "is_leaf": True},
        {"company_key": "TVG", "account_number": "3312", "is_leaf": True},
        {"company_key": "TVG", "account_number": "3313", "is_leaf": True},
    ]

    selected = _rows_for_set(rows, "SHORT_PAYABLE", members, date(2026, 1, 31))

    assert [row["account_number"] for row in selected] == ["3311", "3313"]


def test_source_matching_keeps_legacy_report_templates_compatible_with_v3_profiles() -> None:
    journal_row = {
        "source_key": "AMIS_GENERAL_LEDGER__A538C6BE2AE67652",
        "base_source_key": "AMIS_GENERAL_LEDGER",
    }
    trial_row = {
        "source_key": "AMIS_TRIAL_BALANCE__47B24432AFBF37AD",
        "base_source_key": "AMIS_TRIAL_BALANCE",
    }
    file_journal_row = {
        "source_key": "FILE_GENERAL_LEDGER__3D90476E5ACF6215",
        "base_source_key": "FILE_GENERAL_LEDGER",
    }
    file_trial_row = {
        "source_key": "FILE_TRIAL_BALANCE__96A850FC09E48B69",
        "base_source_key": "FILE_TRIAL_BALANCE",
    }

    assert _source_matches(journal_row, "AMIS_SO_CHI_TIET")
    assert _source_matches(journal_row, "AMIS_GENERAL_LEDGER")
    assert _source_matches(trial_row, "AMIS_CAN_DOI_THOI_GIAN_THUC")
    assert _source_matches(trial_row, "AMIS_TRIAL_BALANCE")
    assert _source_matches(file_journal_row, "AMIS_SO_CHI_TIET")
    assert _source_matches(file_journal_row, "AMIS_GENERAL_LEDGER")
    assert _source_matches(file_trial_row, "AMIS_CAN_DOI_THOI_GIAN_THUC")
    assert _source_matches(file_trial_row, "AMIS_TRIAL_BALANCE")
    assert not _source_matches(journal_row, "AMIS_TRIAL_BALANCE")
    assert not _source_matches(file_journal_row, "AMIS_TRIAL_BALANCE")
    assert not _source_matches(journal_row, "AMIS_GENERAL_LEDGER__OTHER")


def test_report_engine_applies_include_exclude_company_and_previous_period() -> None:
    records = [
        record(
            "report-1",
            "REPORT",
            "B01",
            {"reportName": "Bảng cân đối", "displayTitle": "BẢNG CÂN ĐỐI"},
        ),
        record(
            "column-label",
            "REPORT_COLUMN",
            "B01_LABEL",
            {"reportKey": "B01", "columnKey": "LABEL", "columnType": "LABEL", "columnLabel": "Chỉ tiêu"},
            1,
        ),
        record(
            "column-current",
            "REPORT_COLUMN",
            "B01_CURRENT",
            {
                "reportKey": "B01",
                "columnKey": "CURRENT",
                "columnType": "CURRENT_PERIOD",
                "columnLabel": "Năm nay",
            },
            2,
        ),
        record(
            "column-previous",
            "REPORT_COLUMN",
            "B01_PREVIOUS",
            {
                "reportKey": "B01",
                "columnKey": "PREVIOUS",
                "columnType": "PREVIOUS_PERIOD",
                "columnLabel": "Năm trước",
            },
            3,
        ),
        record(
            "row-1",
            "REPORT_ROW",
            "B01_ROW_110",
            {
                "reportKey": "B01",
                "rowCode": "110",
                "rowLabel": "Tiền",
                "rowType": "DATA",
                "cells": {
                    "CURRENT": {"accountSetKey": "CASH", "periodBasis": "ENDING", "side": "NET"},
                    "PREVIOUS": {"accountSetKey": "CASH", "periodBasis": "ENDING", "side": "NET"},
                },
            },
            1,
            "Tiền",
        ),
        record(
            "member-include",
            "ACCOUNT_SET_MEMBER",
            "CASH_111",
            {
                "accountSetKey": "CASH",
                "matchType": "PREFIX",
                "matchValue": "111",
                "role": "INCLUDE",
                "companyKey": "NTV",
                "status": "ACTIVE",
            },
        ),
        record(
            "member-exclude",
            "ACCOUNT_SET_MEMBER",
            "CASH_1118",
            {
                "accountSetKey": "CASH",
                "matchType": "PREFIX",
                "matchValue": "1118",
                "role": "EXCLUDE",
                "companyKey": "NTV",
                "status": "ACTIVE",
            },
        ),
    ]
    assert required_offsets(records, "B01") == [0, -1]
    rows_by_offset = {
        0: [
            {
                "company_key": "NTV",
                "account_number": "1111",
                "is_leaf": True,
                "closing_debit": 120,
                "closing_credit": 20,
            },
            {
                "company_key": "NTV",
                "account_number": "1118",
                "is_leaf": True,
                "closing_debit": 999,
                "closing_credit": 0,
            },
            {
                "company_key": "OTHER",
                "account_number": "1111",
                "is_leaf": True,
                "closing_debit": 700,
                "closing_credit": 0,
            },
        ],
        -1: [
            {
                "company_key": "NTV",
                "account_number": "1111",
                "is_leaf": True,
                "closing_debit": 80,
                "closing_credit": 5,
            }
        ],
    }
    output = build_output(
        records=records,
        report_key="B01",
        period_from=date(2026, 1, 1),
        period_to=date(2026, 12, 31),
        rows_by_offset=rows_by_offset,
        selection={"sourceKeys": ["SRC"], "companyKeys": ["NTV"], "dataSpaceKeys": ["SPACE"]},
    )
    cells = output["rows"][0]["cells"]
    assert cells[1]["value"] == 100
    assert cells[2]["value"] == 75
    assert output["warnings"] == []


def test_report_engine_can_return_account_level_evidence_that_reconciles_to_cell() -> None:
    records = [
        record("report", "REPORT", "B01", {"reportName": "Bảng cân đối"}),
        record(
            "label",
            "REPORT_COLUMN",
            "B01_LABEL",
            {"reportKey": "B01", "columnKey": "LABEL", "columnType": "LABEL", "columnLabel": "Chỉ tiêu"},
            1,
        ),
        record(
            "current",
            "REPORT_COLUMN",
            "B01_CURRENT",
            {
                "reportKey": "B01",
                "columnKey": "CURRENT",
                "columnType": "CURRENT_PERIOD",
                "columnLabel": "Cuối kỳ",
                "valueBasis": "ENDING",
            },
            2,
        ),
        record(
            "row",
            "REPORT_ROW",
            "B01_141",
            {
                "reportKey": "B01",
                "rowCode": "141",
                "rowLabel": "Hàng tồn kho",
                "rowType": "DATA",
                "cells": {
                    "CURRENT": {
                        "accountSetKey": "INVENTORY",
                        "operator": "BALANCE",
                        "side": "DEBIT",
                    }
                },
            },
            3,
        ),
        record(
            "member",
            "ACCOUNT_SET_MEMBER",
            "INVENTORY_MIXED",
            {
                "accountSetKey": "INVENTORY",
                "matchType": "MIXED",
                "matchValue": "151*,152*,1531,1541,1542,1543,155*,156*,157*,158*",
                "role": "INCLUDE",
            },
            4,
        ),
    ]
    rows = [
        {
            "source_key": "TB_NTV",
            "dataset_type": "TRIAL_BALANCE",
            "company_key": "NTV",
            "data_space_name": "KETOAN_2026",
            "period_from": "2026-01-01",
            "period_to": "2026-01-31",
            "account_number": "1511",
            "account_name": "Hàng mua đang đi đường",
            "is_leaf": True,
            "closing_debit": 1_200,
        },
        {
            "source_key": "TB_NTV",
            "dataset_type": "TRIAL_BALANCE",
            "company_key": "NTV",
            "data_space_name": "KETOAN_2026",
            "period_from": "2026-01-01",
            "period_to": "2026-01-31",
            "account_number": "1542",
            "account_name": "Chi phí sản xuất dở dang",
            "is_leaf": True,
            "closing_debit": 800,
        },
    ]
    output = build_output(
        records=records,
        report_key="B01",
        period_from=date(2026, 1, 1),
        period_to=date(2026, 1, 31),
        rows_by_offset={0: rows},
        selection={"sourceKeys": ["TB_NTV"], "companyKeys": ["NTV"], "dataSpaceKeys": ["KETOAN_2026"]},
        include_evidence=True,
    )
    cell = output["rows"][0]["cells"][1]
    assert cell["value"] == 2_000
    components = cell["evidence"]["components"]
    assert sum(component["signedValue"] for component in components) == cell["value"]
    assert [component["accountNumber"] for component in components] == ["1511", "1542"]
    assert components[0]["label"] == "Nợ cuối kỳ TK 1511 - Hàng mua đang đi đường"


def test_evidence_expands_line_formula_to_source_accounts_without_balancing_plug() -> None:
    records = [
        record("report", "REPORT", "B02", {"reportName": "Kết quả kinh doanh"}),
        record(
            "label",
            "REPORT_COLUMN",
            "B02_LABEL",
            {"reportKey": "B02", "columnKey": "LABEL", "columnType": "LABEL", "columnLabel": "Chỉ tiêu"},
            1,
        ),
        record(
            "current",
            "REPORT_COLUMN",
            "B02_CURRENT",
            {
                "reportKey": "B02",
                "columnKey": "CURRENT",
                "columnType": "CURRENT_PERIOD",
                "columnLabel": "Năm nay",
                "valueBasis": "PERIOD",
            },
            2,
        ),
        record(
            "row-01",
            "REPORT_ROW",
            "B02_01",
            {
                "reportKey": "B02",
                "rowCode": "01",
                "rowLabel": "Doanh thu bán hàng",
                "rowType": "DATA",
                "cells": {"CURRENT": {"accountSetKey": "REVENUE", "operator": "MOVEMENT", "side": "CREDIT"}},
            },
            10,
        ),
        record(
            "row-02",
            "REPORT_ROW",
            "B02_02",
            {
                "reportKey": "B02",
                "rowCode": "02",
                "rowLabel": "Các khoản giảm trừ doanh thu",
                "rowType": "DATA",
                "cells": {"CURRENT": {"accountSetKey": "DEDUCTION", "operator": "MOVEMENT", "side": "DEBIT"}},
            },
            20,
        ),
        record(
            "row-10",
            "REPORT_ROW",
            "B02_10",
            {
                "reportKey": "B02",
                "rowCode": "10",
                "rowLabel": "Doanh thu thuần",
                "rowType": "FORMULA",
                "cells": {"CURRENT": {"operator": "LINE_EXPRESSION", "expression": "01-02"}},
            },
            30,
        ),
        record(
            "revenue-member",
            "ACCOUNT_SET_MEMBER",
            "REVENUE_511",
            {"accountSetKey": "REVENUE", "matchType": "PREFIX", "matchValue": "511", "role": "INCLUDE"},
            40,
        ),
        record(
            "deduction-member",
            "ACCOUNT_SET_MEMBER",
            "DEDUCTION_521",
            {"accountSetKey": "DEDUCTION", "matchType": "PREFIX", "matchValue": "521", "role": "INCLUDE"},
            50,
        ),
    ]
    rows = [
        {
            "source_key": "TB_NTV",
            "dataset_type": "TRIAL_BALANCE",
            "company_key": "NTV",
            "data_space_name": "KETOAN_2026",
            "period_from": "2026-01-01",
            "period_to": "2026-01-31",
            "account_number": "5111",
            "account_name": "Doanh thu bán hàng",
            "is_leaf": True,
            "period_credit": 1_000,
        },
        {
            "source_key": "TB_NTV",
            "dataset_type": "TRIAL_BALANCE",
            "company_key": "NTV",
            "data_space_name": "KETOAN_2026",
            "period_from": "2026-01-01",
            "period_to": "2026-01-31",
            "account_number": "5211",
            "account_name": "Chiết khấu thương mại",
            "is_leaf": True,
            "period_debit": 200,
        },
    ]
    output = build_output(
        records=records,
        report_key="B02",
        period_from=date(2026, 1, 1),
        period_to=date(2026, 1, 31),
        rows_by_offset={0: rows},
        selection={"sourceKeys": ["TB_NTV"], "companyKeys": ["NTV"], "dataSpaceKeys": ["KETOAN_2026"]},
        include_evidence=True,
    )
    evidence = output["rows"][2]["cells"][1]["evidence"]
    assert evidence["value"] == 800
    assert evidence["isComplete"] is True
    assert evidence["difference"] == 0
    assert [component["accountNumber"] for component in evidence["components"]] == ["5111", "5211"]
    assert sum(component["signedValue"] for component in evidence["components"]) == 800
    assert "Dòng 10 - Doanh thu thuần → Dòng 01 - Doanh thu bán hàng" in evidence["components"][0]["tracePath"]
    assert all("Phần còn lại" not in component["label"] for component in evidence["components"])


def test_evidence_keeps_every_zero_formula_branch_and_mixed_zero_branch() -> None:
    records = [
        record("report", "REPORT", "B04", {"reportName": "Kết quả kinh doanh"}),
        record(
            "label",
            "REPORT_COLUMN",
            "B04_LABEL",
            {"reportKey": "B04", "columnKey": "LABEL", "columnType": "LABEL", "columnLabel": "Chỉ tiêu"},
            1,
        ),
        record(
            "current",
            "REPORT_COLUMN",
            "B04_CURRENT",
            {
                "reportKey": "B04",
                "columnKey": "CURRENT",
                "columnType": "CURRENT_PERIOD",
                "columnLabel": "Năm nay",
                "valueBasis": "PERIOD",
            },
            2,
        ),
        record(
            "row-01",
            "REPORT_ROW",
            "B04_01",
            {
                "reportKey": "B04",
                "rowCode": "01",
                "rowLabel": "Doanh thu bán hàng",
                "rowType": "DATA",
                "cells": {"CURRENT": {"accountSetKey": "REVENUE", "operator": "MOVEMENT", "side": "CREDIT"}},
            },
            10,
        ),
        record(
            "row-02",
            "REPORT_ROW",
            "B04_02",
            {
                "reportKey": "B04",
                "rowCode": "02",
                "rowLabel": "Các khoản giảm trừ doanh thu",
                "rowType": "DATA",
                "cells": {"CURRENT": {"accountSetKey": "DEDUCTION", "operator": "MOVEMENT", "side": "DEBIT"}},
            },
            20,
        ),
        record(
            "row-10",
            "REPORT_ROW",
            "B04_10",
            {
                "reportKey": "B04",
                "rowCode": "10",
                "rowLabel": "Doanh thu thuần",
                "rowType": "FORMULA",
                "cells": {"CURRENT": {"operator": "LINE_EXPRESSION", "expression": "01-02"}},
            },
            30,
        ),
        record(
            "revenue-member",
            "ACCOUNT_SET_MEMBER",
            "REVENUE_511",
            {"accountSetKey": "REVENUE", "matchType": "PREFIX", "matchValue": "511", "role": "INCLUDE"},
            40,
        ),
        record(
            "deduction-member",
            "ACCOUNT_SET_MEMBER",
            "DEDUCTION_521",
            {"accountSetKey": "DEDUCTION", "matchType": "PREFIX", "matchValue": "521", "role": "INCLUDE"},
            50,
        ),
    ]
    selection = {
        "sourceKeys": ["TB_NTV"],
        "companyKeys": ["NTV"],
        "dataSpaceKeys": ["KETOAN_2026"],
    }

    empty_output = build_output(
        records=records,
        report_key="B04",
        period_from=date(2026, 1, 1),
        period_to=date(2026, 1, 31),
        rows_by_offset={0: []},
        selection=selection,
        include_evidence=True,
    )
    empty_evidence = empty_output["rows"][2]["cells"][1]["evidence"]
    empty_components = empty_evidence["components"]
    assert empty_evidence["value"] == 0
    assert empty_evidence["isComplete"] is True
    assert empty_evidence["difference"] == 0
    assert len(empty_components) == 2
    assert [component["formulaSign"] for component in empty_components] == [1, -1]
    assert all(component["signedValue"] == 0 for component in empty_components)
    assert "TK 511*" in empty_components[0]["label"]
    assert "TK 521*" in empty_components[1]["label"]
    assert all("Công thức vẫn áp dụng" in component["tracePath"] for component in empty_components)

    mixed_output = build_output(
        records=records,
        report_key="B04",
        period_from=date(2026, 1, 1),
        period_to=date(2026, 1, 31),
        rows_by_offset={
            0: [
                {
                    "source_key": "TB_NTV",
                    "dataset_type": "TRIAL_BALANCE",
                    "company_key": "NTV",
                    "data_space_name": "KETOAN_2026",
                    "period_from": "2026-01-01",
                    "period_to": "2026-01-31",
                    "account_number": "5111",
                    "account_name": "Doanh thu bán hàng",
                    "is_leaf": True,
                    "period_credit": 1_000,
                }
            ]
        },
        selection=selection,
        include_evidence=True,
    )
    mixed_evidence = mixed_output["rows"][2]["cells"][1]["evidence"]
    mixed_components = mixed_evidence["components"]
    assert mixed_evidence["value"] == 1_000
    assert mixed_evidence["isComplete"] is True
    assert len(mixed_components) == 2
    assert mixed_components[0]["accountNumber"] == "5111"
    assert mixed_components[1]["signedValue"] == 0
    assert mixed_components[1]["formulaSign"] == -1
    assert "TK 521*" in mixed_components[1]["label"]


def test_incomplete_evidence_reports_difference_instead_of_creating_a_plug() -> None:
    records = [
        record("report", "REPORT", "B03", {"reportName": "Báo cáo kiểm tra"}),
        record("label", "REPORT_COLUMN", "B03_LABEL", {"reportKey": "B03", "columnKey": "LABEL", "columnType": "LABEL"}, 1),
        record("current", "REPORT_COLUMN", "B03_CURRENT", {"reportKey": "B03", "columnKey": "CURRENT", "columnType": "CURRENT_PERIOD"}, 2),
        record("row-a", "REPORT_ROW", "B03_A", {"reportKey": "B03", "rowCode": "A", "rowLabel": "A", "rowType": "DATA", "cells": {"CURRENT": approved_adjustment_cell("B03", "A")}}, 10),
        record("row-b", "REPORT_ROW", "B03_B", {"reportKey": "B03", "rowCode": "B", "rowLabel": "B", "rowType": "DATA", "cells": {"CURRENT": approved_adjustment_cell("B03", "B")}}, 20),
        record("row-total", "REPORT_ROW", "B03_TOTAL", {"reportKey": "B03", "rowCode": "TOTAL", "rowLabel": "Tổng", "rowType": "FORMULA", "cells": {"CURRENT": {"operator": "LINE_EXPRESSION", "expression": "A*B"}}}, 30),
    ]
    output = build_output(
        records=records,
        report_key="B03",
        period_from=date(2026, 1, 1),
        period_to=date(2026, 1, 31),
        rows_by_offset={0: [
            approved_adjustment_row("B03", "A", 10),
            approved_adjustment_row("B03", "B", 20),
        ]},
        selection={"sourceKeys": [], "companyKeys": [], "dataSpaceKeys": []},
        include_evidence=True,
    )
    evidence = output["rows"][2]["cells"][1]["evidence"]
    assert evidence["value"] == 200
    assert evidence["isComplete"] is False
    assert evidence["difference"] == 200
    assert "không tự tạo số bù" in evidence["warning"]
    assert all("Phần còn lại" not in component["label"] for component in evidence["components"])


def test_opening_column_uses_current_snapshot_instead_of_shifting_one_year() -> None:
    records = [
        record("report", "REPORT", "BS", {"reportName": "Tình hình tài chính"}),
        record(
            "label",
            "REPORT_COLUMN",
            "BS_LABEL",
            {"reportKey": "BS", "columnKey": "LABEL", "columnType": "LABEL", "columnLabel": "Chỉ tiêu"},
            1,
        ),
        record(
            "opening",
            "REPORT_COLUMN",
            "BS_OPENING",
            {
                "reportKey": "BS",
                "columnKey": "OPENING",
                "columnType": "PREVIOUS_PERIOD",
                "columnLabel": "Số đầu năm",
                "periodOffset": -1,
                "valueBasis": "OPENING",
            },
            2,
        ),
        record(
            "cash-row",
            "REPORT_ROW",
            "BS_110",
            {
                "reportKey": "BS",
                "rowCode": "110",
                "rowLabel": "Tiền",
                "rowType": "DATA",
                "cells": {
                    "OPENING": {
                        "accountSetKey": "CASH",
                        "metricKey": "CLOSING_DEBIT",
                        "periodBasis": "ENDING",
                        "side": "DEBIT",
                    }
                },
            },
        ),
        record(
            "cash-member",
            "ACCOUNT_SET_MEMBER",
            "CASH_111",
            {"accountSetKey": "CASH", "matchType": "PREFIX", "matchValue": "111", "role": "INCLUDE"},
        ),
    ]
    assert required_offsets(records, "BS") == [0]
    output = build_output(
        records=records,
        report_key="BS",
        period_from=date(2026, 1, 1),
        period_to=date(2026, 1, 31),
        rows_by_offset={
            0: [
                {
                    "account_number": "1111",
                    "is_leaf": True,
                    "opening_debit": 100,
                    "closing_debit": 130,
                }
            ]
        },
        selection={"sourceKeys": ["TB"], "companyKeys": [], "dataSpaceKeys": []},
    )
    assert output["rows"][0]["cells"][1]["value"] == 100


def test_balance_sides_and_delta_basis_are_accounting_aware() -> None:
    records = [
        record("report", "REPORT", "BS", {"reportName": "BS"}),
        record(
            "label",
            "REPORT_COLUMN",
            "BS_LABEL",
            {"reportKey": "BS", "columnKey": "LABEL", "columnType": "LABEL", "columnLabel": "Chỉ tiêu"},
            1,
        ),
        record(
            "current",
            "REPORT_COLUMN",
            "BS_CURRENT",
            {
                "reportKey": "BS",
                "columnKey": "CURRENT",
                "columnType": "CURRENT_PERIOD",
                "columnLabel": "Kỳ này",
                "valueBasis": "YTD",
            },
            2,
        ),
        record(
            "asset",
            "REPORT_ROW",
            "BS_100",
            {
                "reportKey": "BS",
                "rowCode": "100",
                "rowLabel": "Tài sản",
                "rowType": "DATA",
                "cells": {"CURRENT": {"accountSetKey": "ASSET", "side": "DEBIT_BALANCE", "useColumnBasis": False}},
            },
            1,
        ),
        record(
            "liability",
            "REPORT_ROW",
            "BS_300",
            {
                "reportKey": "BS",
                "rowCode": "300",
                "rowLabel": "Nợ phải trả",
                "rowType": "DATA",
                "cells": {"CURRENT": {"accountSetKey": "LIABILITY", "side": "CREDIT_BALANCE", "useColumnBasis": False}},
            },
            2,
        ),
        record(
            "delta",
            "REPORT_ROW",
            "BS_DELTA",
            {
                "reportKey": "BS",
                "rowCode": "D",
                "rowLabel": "Biến động tài sản",
                "rowType": "DATA",
                "cells": {
                    "CURRENT": {
                        "operator": "DELTA_BALANCE",
                        "accountSetKey": "ASSET",
                        "periodBasis": "DELTA",
                        "side": "DEBIT_BALANCE",
                    }
                },
            },
            3,
        ),
        record("asset-member", "ACCOUNT_SET_MEMBER", "ASSET", {"accountSetKey": "ASSET", "matchType": "PREFIX", "matchValue": "1", "role": "INCLUDE"}),
        record("liability-member", "ACCOUNT_SET_MEMBER", "LIABILITY", {"accountSetKey": "LIABILITY", "matchType": "PREFIX", "matchValue": "3", "role": "INCLUDE"}),
    ]
    output = build_output(
        records=records,
        report_key="BS",
        period_from=date(2026, 1, 1),
        period_to=date(2026, 1, 31),
        rows_by_offset={
            0: [
                {"account_number": "111", "is_leaf": True, "opening_debit": 80, "closing_debit": 120, "closing_credit": 20, "ytd_debit": 999},
                {"account_number": "331", "is_leaf": True, "opening_credit": 60, "closing_debit": 10, "closing_credit": 110, "ytd_credit": 999},
            ]
        },
        selection={"sourceKeys": ["TB"], "companyKeys": [], "dataSpaceKeys": []},
    )
    values = {row["rowCode"]: row["cells"][1]["value"] for row in output["rows"]}
    assert values == {"100": 100, "300": 100, "D": 20}


def test_formula_terms_combine_cross_report_and_audited_adjustment() -> None:
    records = [
        record("report", "REPORT", "BS", {"reportName": "BS"}),
        record("label", "REPORT_COLUMN", "LABEL", {"reportKey": "BS", "columnKey": "LABEL", "columnType": "LABEL", "columnLabel": "Chỉ tiêu"}, 1),
        record("current", "REPORT_COLUMN", "CURRENT", {"reportKey": "BS", "columnKey": "CURRENT", "columnType": "CURRENT_PERIOD", "columnLabel": "Kỳ này"}, 2),
        record(
            "profit",
            "REPORT_ROW",
            "BS_420B",
            {
                "reportKey": "BS",
                "rowCode": "420b",
                "rowLabel": "LNST kỳ này",
                "rowType": "CROSS_REPORT",
                "cells": {
                    "CURRENT": {
                        "terms": [
                            {"operator": "CROSS_REPORT", "targetReportKey": "PL", "targetRowCode": "60", "sign": 1},
                            {"operator": "MANUAL_ADJUSTMENT", "sourceKey": "BCTC_DIEU_CHINH", "filters": {"report_key": "BS", "row_code": "420b"}, "sign": 1},
                        ]
                    }
                },
            },
        ),
    ]

    output = build_output(
        records=records,
        report_key="BS",
        period_from=date(2026, 1, 1),
        period_to=date(2026, 1, 31),
        rows_by_offset={
            0: [
                {
                    "source_key": "BCTC_DIEU_CHINH",
                    "dataset_type": "MANUAL_ADJUSTMENT",
                    "canonical_data": {"report_key": "BS", "row_code": "420b", "amount": -15},
                }
            ]
        },
        selection={"sourceKeys": ["TB", "BCTC_DIEU_CHINH"], "companyKeys": [], "dataSpaceKeys": []},
        cross_report_resolver=lambda *_: 100,
    )
    assert output["rows"][0]["cells"][1]["value"] == 85


def test_report_engine_rejects_manual_value_embedded_in_formula() -> None:
    records = [
        record("report", "REPORT", "CF", {"reportName": "Lưu chuyển tiền tệ"}),
        record("label", "REPORT_COLUMN", "CF_LABEL", {"reportKey": "CF", "columnKey": "LABEL", "columnType": "LABEL"}, 1),
        record("current", "REPORT_COLUMN", "CF_CURRENT", {"reportKey": "CF", "columnKey": "CURRENT", "columnType": "CURRENT_PERIOD"}, 2),
        record(
            "row-27",
            "REPORT_ROW",
            "CF_27",
            {
                "reportKey": "CF",
                "rowCode": "27",
                "rowType": "DATA",
                "cells": {"CURRENT": {"operator": "MANUAL_ADJUSTMENT", "manualValue": 742_683}},
            },
            3,
        ),
    ]
    output = build_output(
        records=records,
        report_key="CF",
        period_from=date(2026, 1, 1),
        period_to=date(2026, 1, 31),
        rows_by_offset={0: []},
        selection={"sourceKeys": [], "companyKeys": [], "dataSpaceKeys": []},
        include_evidence=True,
    )
    cell = output["rows"][0]["cells"][1]
    assert cell["value"] == 0
    assert output["controlStatus"] == "NEEDS_REVIEW"
    assert any("manualValue bị từ chối" in message for message in output["blockingIssues"])
    assert cell["evidence"]["components"][0]["signedValue"] == 0


def test_unconfigured_numeric_cells_stay_blank_while_configured_zero_stays_zero() -> None:
    records = [
        record("report", "REPORT", "PL", {"reportName": "Kết quả kinh doanh"}),
        record(
            "label",
            "REPORT_COLUMN",
            "PL_LABEL",
            {
                "reportKey": "PL",
                "columnKey": "LABEL",
                "columnType": "LABEL",
                "columnLabel": "Chỉ tiêu",
            },
            1,
        ),
        record(
            "current",
            "REPORT_COLUMN",
            "PL_CURRENT",
            {
                "reportKey": "PL",
                "columnKey": "CURRENT",
                "columnType": "CURRENT_PERIOD",
                "columnLabel": "Năm nay",
            },
            2,
        ),
        record(
            "row-70",
            "REPORT_ROW",
            "PL_70",
            {
                "reportKey": "PL",
                "rowCode": "70",
                "rowLabel": "Lãi cơ bản trên cổ phiếu",
                "rowType": "DATA",
            },
            10,
        ),
        record(
            "row-71",
            "REPORT_ROW",
            "PL_71",
            {
                "reportKey": "PL",
                "rowCode": "71",
                "rowLabel": "Lãi suy giảm trên cổ phiếu",
                "rowType": "DATA",
                "ruleKey": "LEGACY_EPS_RULE",
                "cells": {
                    "CURRENT": {
                        "contentKind": "ACCOUNT_SET",
                        "terms": [],
                        "formulaRequired": True,
                    }
                },
            },
            20,
        ),
        record(
            "row-01",
            "REPORT_ROW",
            "PL_01",
            {
                "reportKey": "PL",
                "rowCode": "01",
                "rowLabel": "Doanh thu",
                "rowType": "DATA",
                "cells": {
                    "CURRENT": {
                        "operator": "MOVEMENT",
                        "accountSetKey": "REVENUE",
                        "side": "CREDIT",
                    }
                },
            },
            30,
        ),
        record(
            "legacy-rule",
            "RULE",
            "LEGACY_EPS_RULE",
            {
                "ruleKey": "LEGACY_EPS_RULE",
                "operator": "MANUAL_ADJUSTMENT",
                "sourceKey": "BCTC_DIEU_CHINH",
                "datasetType": "MANUAL_ADJUSTMENT",
                "filters": {"report_key": "PL", "row_code": "71"},
            },
        ),
        record(
            "revenue-member",
            "ACCOUNT_SET_MEMBER",
            "REVENUE_511",
            {
                "accountSetKey": "REVENUE",
                "matchType": "PREFIX",
                "matchValue": "511",
                "role": "INCLUDE",
            },
        ),
    ]

    output = build_output(
        records=records,
        report_key="PL",
        period_from=date(2026, 1, 1),
        period_to=date(2026, 1, 31),
        rows_by_offset={0: []},
        selection={"sourceKeys": [], "companyKeys": [], "dataSpaceKeys": []},
        include_evidence=True,
    )
    cells = {row["rowCode"]: row["cells"][1] for row in output["rows"]}

    assert cells["70"]["value"] is None
    assert cells["70"]["formulaStatus"] == "MISSING"
    assert cells["70"]["evidence"]["value"] is None
    assert cells["70"]["evidence"]["isComplete"] is False
    assert cells["71"]["value"] is None
    assert cells["71"]["formulaStatus"] == "MISSING"
    assert cells["01"]["value"] == 0
    assert cells["01"]["formulaStatus"] == "CONFIGURED"
    assert any("Dòng 70 chưa có công thức" in message for message in output["warnings"])
    assert any("Dòng 71 còn ô chưa thiết lập công thức" in message for message in output["warnings"])


def test_approved_adjustment_replaces_unclassified_cash_flow_without_guessing() -> None:
    records = [
        record("report", "REPORT", "CF", {"reportName": "Lưu chuyển tiền tệ"}),
        record(
            "label",
            "REPORT_COLUMN",
            "CF_LABEL",
            {"reportKey": "CF", "columnKey": "LABEL", "columnType": "LABEL", "columnLabel": "Chỉ tiêu"},
            1,
        ),
        record(
            "current",
            "REPORT_COLUMN",
            "CF_CURRENT",
            {"reportKey": "CF", "columnKey": "KY_NAY", "columnType": "CURRENT_PERIOD", "columnLabel": "Kỳ này"},
            2,
        ),
        record(
            "row-14",
            "REPORT_ROW",
            "CF_14",
            {
                "reportKey": "CF",
                "rowCode": "14",
                "rowLabel": "Chi phí đi vay đã trả",
                "rowType": "DATA",
                "cells": {
                    "KY_NAY": {
                        "terms": [
                            {
                                "operator": "CASH_MOVEMENT",
                                "sourceKey": "AMIS_GENERAL_LEDGER",
                                "datasetType": "GENERAL_LEDGER",
                                "accountSetKey": "CASH",
                                "counterAccountSetKey": "INTEREST",
                                "side": "CREDIT",
                                "sign": -1,
                                "conditionKey": "CF_14",
                                "conditionRequirement": "REQUIRED",
                                "classificationPolicy": "REQUIRE_CASH_FLOW_CODE_OR_APPROVED_ADJUSTMENT",
                            },
                            {
                                "operator": "MANUAL_ADJUSTMENT",
                                "sourceKey": "BCTC_DIEU_CHINH",
                                "datasetType": "MANUAL_ADJUSTMENT",
                                "filters": {"report_key": "CF", "row_code": "14"},
                                "sign": 1,
                            },
                        ]
                    }
                },
            },
            3,
        ),
        record(
            "cash-member",
            "ACCOUNT_SET_MEMBER",
            "CASH_111",
            {"accountSetKey": "CASH", "matchType": "PREFIX", "matchValue": "111", "role": "INCLUDE"},
        ),
        record(
            "interest-member",
            "ACCOUNT_SET_MEMBER",
            "INTEREST_635",
            {"accountSetKey": "INTEREST", "matchType": "PREFIX", "matchValue": "635", "role": "INCLUDE"},
        ),
    ]
    rows = [
        {
            "source_key": "AMIS_GENERAL_LEDGER",
            "dataset_type": "GENERAL_LEDGER",
            "company_key": "TVG",
            "canonical_data": {
                "debit_account": "6351",
                "credit_account": "1111",
                "amount": 50,
                "entry_count": 1,
                "cash_flow_code": "",
            },
        },
        {
            "source_key": "BCTC_DIEU_CHINH",
            "dataset_type": "MANUAL_ADJUSTMENT",
            "company_key": "TVG",
            "canonical_data": {
                "report_key": "CF",
                "row_code": "14",
                "column_key": "KY_NAY",
                "amount": -50,
                "approval_status": "APPROVED",
                "reason": "Đối chiếu phiếu trả lãi vay",
                "source_reference": "UNC-001",
            },
        },
    ]

    output = build_output(
        records=records,
        report_key="CF",
        period_from=date(2026, 1, 1),
        period_to=date(2026, 1, 31),
        rows_by_offset={0: rows},
        selection={"sourceKeys": ["AMIS_GENERAL_LEDGER"], "companyKeys": ["TVG"], "dataSpaceKeys": ["SPACE"]},
        include_evidence=True,
    )

    assert output["rows"][0]["cells"][1]["value"] == -50
    assert output["blockingIssues"] == []
    assert any("điều chỉnh đã được duyệt" in warning for warning in output["warnings"])
    evidence = output["rows"][0]["cells"][1]["evidence"]["components"]
    assert any("Đối chiếu phiếu trả lãi vay" in component["label"] for component in evidence)


def test_detailed_movement_honors_source_counter_account_and_excludes_closing() -> None:
    records = [
        record("report", "REPORT", "PL", {"reportName": "Kết quả kinh doanh"}),
        record(
            "label",
            "REPORT_COLUMN",
            "PL_LABEL",
            {"reportKey": "PL", "columnKey": "LABEL", "columnType": "LABEL", "columnLabel": "Chỉ tiêu"},
            1,
        ),
        record(
            "current",
            "REPORT_COLUMN",
            "PL_CURRENT",
            {"reportKey": "PL", "columnKey": "CURRENT", "columnType": "CURRENT_PERIOD", "columnLabel": "Kỳ này"},
            2,
        ),
        record(
            "cogs",
            "REPORT_ROW",
            "PL_11",
            {
                "reportKey": "PL",
                "rowCode": "11",
                "rowLabel": "Giá vốn hàng bán",
                "rowType": "DATA",
                "cells": {
                    "CURRENT": {
                        "operator": "MOVEMENT",
                        "sourceKey": "DETAIL",
                        "accountSetKey": "COGS",
                        "counterAccountSetKey": "INVENTORY",
                        "side": "DEBIT",
                    }
                },
            },
        ),
        record(
            "cogs-member",
            "ACCOUNT_SET_MEMBER",
            "COGS_632",
            {"accountSetKey": "COGS", "matchType": "PREFIX", "matchValue": "632", "role": "INCLUDE"},
        ),
        record(
            "inventory-member",
            "ACCOUNT_SET_MEMBER",
            "INVENTORY_15",
            {"accountSetKey": "INVENTORY", "matchType": "PREFIX", "matchValue": "15", "role": "INCLUDE"},
        ),
    ]
    rows = [
        {
            "source_key": "DETAIL__COMPANY",
            "base_source_key": "DETAIL",
            "dataset_type": "GENERAL_LEDGER",
            "canonical_data": {"debit_account": "6321", "credit_account": "1561", "amount": 100},
        },
        {
            "source_key": "DETAIL__COMPANY",
            "base_source_key": "DETAIL",
            "dataset_type": "GENERAL_LEDGER",
            "canonical_data": {"debit_account": "6321", "credit_account": "217", "amount": 50},
        },
        {
            "source_key": "DETAIL__COMPANY",
            "base_source_key": "DETAIL",
            "dataset_type": "GENERAL_LEDGER",
            "canonical_data": {"debit_account": "911", "credit_account": "6321", "amount": 100},
        },
        {
            "source_key": "OTHER",
            "base_source_key": "OTHER",
            "dataset_type": "GENERAL_LEDGER",
            "canonical_data": {"debit_account": "6321", "credit_account": "1561", "amount": 999},
        },
    ]
    output = build_output(
        records=records,
        report_key="PL",
        period_from=date(2026, 1, 1),
        period_to=date(2026, 1, 31),
        rows_by_offset={0: rows},
        selection={"sourceKeys": ["DETAIL__COMPANY"], "companyKeys": [], "dataSpaceKeys": []},
    )
    assert output["rows"][0]["cells"][1]["value"] == 100
    assert output["warnings"] == []


def test_detailed_movement_preserves_negative_adjustment_sign() -> None:
    records = [
        record("report", "REPORT", "PL", {"reportName": "Kết quả kinh doanh"}),
        record(
            "label",
            "REPORT_COLUMN",
            "PL_LABEL",
            {"reportKey": "PL", "columnKey": "LABEL", "columnType": "LABEL", "columnLabel": "Chỉ tiêu"},
            1,
        ),
        record(
            "current",
            "REPORT_COLUMN",
            "PL_CURRENT",
            {"reportKey": "PL", "columnKey": "CURRENT", "columnType": "CURRENT_PERIOD", "columnLabel": "Kỳ này"},
            2,
        ),
        record(
            "revenue",
            "REPORT_ROW",
            "PL_01",
            {
                "reportKey": "PL",
                "rowCode": "01",
                "rowLabel": "Doanh thu",
                "rowType": "DATA",
                "cells": {
                    "CURRENT": {
                        "operator": "MOVEMENT_DETAIL",
                        "sourceKey": "DETAIL",
                        "datasetType": "GENERAL_LEDGER",
                        "accountSetKey": "REVENUE",
                        "side": "CREDIT",
                    }
                },
            },
        ),
        record(
            "revenue-member",
            "ACCOUNT_SET_MEMBER",
            "REVENUE_511",
            {"accountSetKey": "REVENUE", "matchType": "PREFIX", "matchValue": "511", "role": "INCLUDE"},
        ),
    ]
    output = build_output(
        records=records,
        report_key="PL",
        period_from=date(2026, 1, 1),
        period_to=date(2026, 1, 31),
        rows_by_offset={
            0: [
                {
                    "source_key": "DETAIL__COMPANY",
                    "base_source_key": "DETAIL",
                    "dataset_type": "GENERAL_LEDGER",
                    "canonical_data": {"debit_account": "131", "credit_account": "5111", "amount": 100},
                },
                {
                    "source_key": "DETAIL__COMPANY",
                    "base_source_key": "DETAIL",
                    "dataset_type": "GENERAL_LEDGER",
                    "canonical_data": {"debit_account": "131", "credit_account": "5111", "amount": -20},
                },
            ]
        },
        selection={"sourceKeys": ["DETAIL__COMPANY"], "companyKeys": [], "dataSpaceKeys": []},
    )
    assert output["rows"][0]["cells"][1]["value"] == 80


def test_cash_movement_uses_cash_flow_code_and_keeps_legacy_fallback() -> None:
    records = [
        record("report", "REPORT", "CF", {"reportName": "Lưu chuyển tiền tệ"}),
        record(
            "label",
            "REPORT_COLUMN",
            "CF_LABEL",
            {"reportKey": "CF", "columnKey": "LABEL", "columnType": "LABEL"},
            1,
        ),
        record(
            "current",
            "REPORT_COLUMN",
            "CF_CURRENT",
            {"reportKey": "CF", "columnKey": "CURRENT", "columnType": "CURRENT_PERIOD"},
            2,
        ),
        record(
            "cash-receipt",
            "REPORT_ROW",
            "CF_01",
            {
                "reportKey": "CF",
                "rowCode": "01",
                "rowLabel": "Tiền thu bán hàng",
                "rowType": "DATA",
                "cells": {
                    "CURRENT": {
                        "operator": "CASH_MOVEMENT",
                        "sourceKey": "JOURNAL",
                        "accountSetKey": "CASH",
                        "counterAccountSetKey": "SALES",
                        "conditionKey": "DIEU_KIEN_DONG_TIEN_01",
                        "side": "DEBIT",
                    }
                },
            },
        ),
        record(
            "cash-member",
            "ACCOUNT_SET_MEMBER",
            "CASH_112",
            {"accountSetKey": "CASH", "matchType": "PREFIX", "matchValue": "112", "role": "INCLUDE"},
        ),
        record(
            "sales-member",
            "ACCOUNT_SET_MEMBER",
            "SALES_511",
            {"accountSetKey": "SALES", "matchType": "PREFIX", "matchValue": "511", "role": "INCLUDE"},
        ),
    ]
    output = build_output(
        records=records,
        report_key="CF",
        period_from=date(2026, 1, 1),
        period_to=date(2026, 1, 31),
        rows_by_offset={
            0: [
                {"base_source_key": "JOURNAL", "canonical_data": {"debit_account": "1121", "credit_account": "5111", "amount": 100, "cash_flow_code": "01"}},
                {"base_source_key": "JOURNAL", "canonical_data": {"debit_account": "1121", "credit_account": "5111", "amount": 50, "cash_flow_code": "16"}},
                {"base_source_key": "JOURNAL", "canonical_data": {"debit_account": "1121", "credit_account": "5111", "amount": 25}},
            ]
        },
        selection={"sourceKeys": ["JOURNAL"], "companyKeys": [], "dataSpaceKeys": []},
    )
    assert output["rows"][0]["cells"][1]["value"] == 125


def test_standard_cash_flow_corrections_do_not_guess_missing_classification() -> None:
    report_key = "LUU_CHUYEN_TIEN_TE_GIAN_TIEP"
    records = [
        record("report", "REPORT", report_key, {"reportName": "Lưu chuyển tiền tệ"}),
        record(
            "label",
            "REPORT_COLUMN",
            "CF_LABEL",
            {"reportKey": report_key, "columnKey": "LABEL", "columnType": "LABEL"},
            1,
        ),
        record(
            "current",
            "REPORT_COLUMN",
            "CF_CURRENT",
            {"reportKey": report_key, "columnKey": "CURRENT", "columnType": "CURRENT_PERIOD"},
            2,
        ),
        record(
            "row-05",
            "REPORT_ROW",
            "CF_05",
            {
                "reportKey": report_key,
                "rowCode": "05",
                "rowLabel": "Lãi, lỗ từ hoạt động đầu tư",
                "rowType": "DATA",
                "cells": {
                    "CURRENT": {
                        "operator": "PROFIT_LOSS_ACTIVITY",
                        "sourceKey": "JOURNAL",
                        "datasetType": "GENERAL_LEDGER",
                        "accountSetKey": "INVESTMENT_RESULT",
                        "side": "CREDIT",
                        "sign": -1,
                        "useColumnBasis": False,
                    }
                },
            },
            5,
        ),
        record(
            "row-22",
            "REPORT_ROW",
            "CF_22",
            {
                "reportKey": report_key,
                "rowCode": "22",
                "rowLabel": "Tiền thu thanh lý, nhượng bán tài sản",
                "rowType": "DATA",
                "cells": {
                    "CURRENT": {
                        "operator": "CASH_MOVEMENT",
                        "sourceKey": "JOURNAL",
                        "accountSetKey": "CASH",
                        "counterAccountSetKey": "DISPOSAL",
                        "conditionKey": "DIEU_KIEN_DONG_TIEN_22",
                        "conditionRequirement": "REQUIRED",
                        "side": "DEBIT",
                        "sign": 1,
                    }
                },
            },
            22,
        ),
        record(
            "row-36",
            "REPORT_ROW",
            "CF_36",
            {
                "reportKey": report_key,
                "rowCode": "36",
                "rowLabel": "Cổ tức, lợi nhuận đã trả",
                "rowType": "DATA",
                "cells": {
                    "CURRENT": {
                        "operator": "CASH_MOVEMENT",
                        "sourceKey": "JOURNAL",
                        "accountSetKey": "CASH",
                        "counterAccountSetKey": "RETAINED_EARNINGS",
                        "conditionKey": "DIEU_KIEN_DONG_TIEN_36",
                        "conditionRequirement": "REQUIRED",
                        "side": "CREDIT",
                        "sign": -1,
                    }
                },
            },
            36,
        ),
        record(
            "investment-result",
            "ACCOUNT_SET_MEMBER",
            "INVESTMENT_RESULT_5152",
            {"accountSetKey": "INVESTMENT_RESULT", "matchType": "PREFIX", "matchValue": "5152", "role": "INCLUDE"},
        ),
        record(
            "cash",
            "ACCOUNT_SET_MEMBER",
            "CASH_112",
            {"accountSetKey": "CASH", "matchType": "PREFIX", "matchValue": "112", "role": "INCLUDE"},
        ),
        record(
            "disposal",
            "ACCOUNT_SET_MEMBER",
            "DISPOSAL_711",
            {"accountSetKey": "DISPOSAL", "matchType": "PREFIX", "matchValue": "711", "role": "INCLUDE"},
        ),
        record(
            "retained",
            "ACCOUNT_SET_MEMBER",
            "RETAINED_421",
            {"accountSetKey": "RETAINED_EARNINGS", "matchType": "PREFIX", "matchValue": "421", "role": "INCLUDE"},
        ),
    ]
    journal_rows = [
        {"debit_account": "5152", "credit_account": "911", "amount": 80, "entry_type": "CLOSING"},
        {"debit_account": "5151", "credit_account": "911", "amount": 20, "entry_type": "CLOSING"},
        {"debit_account": "1121", "credit_account": "7111", "amount": 100, "cash_flow_code": "22"},
        {"debit_account": "7111", "credit_account": "1121", "amount": 30, "cash_flow_code": "22"},
        {"debit_account": "1121", "credit_account": "7111", "amount": 25},
        {"debit_account": "4212", "credit_account": "1121", "amount": 40, "cash_flow_code": "36"},
        {"debit_account": "4212", "credit_account": "1121", "amount": 10},
        {"debit_account": "4212", "credit_account": "1121", "amount": 15, "cash_flow_code": "34"},
    ]
    output = build_output(
        records=records,
        report_key=report_key,
        period_from=date(2026, 1, 1),
        period_to=date(2026, 1, 31),
        rows_by_offset={
            0: [
                {
                    "source_key": "JOURNAL__NTV",
                    "base_source_key": "JOURNAL",
                    "dataset_type": "GENERAL_LEDGER",
                    "canonical_data": row,
                }
                for row in journal_rows
            ]
        },
        selection={"sourceKeys": ["JOURNAL__NTV"], "companyKeys": ["NTV"], "dataSpaceKeys": []},
        include_evidence=True,
    )
    values = {row["rowCode"]: row["cells"][1]["value"] for row in output["rows"]}
    assert values == {"05": -80, "22": 100, "36": -40}
    assert output["controlStatus"] == "NEEDS_REVIEW"
    assert len(output["blockingIssues"]) == 2
    assert all("không tự gán" in message.lower() for message in output["blockingIssues"])
    assert all(
        check["status"] == "FAIL"
        for check in output["controlChecks"]
        if check["checkKey"].startswith("SOURCE_CLASSIFICATION_COMPLETE_")
    )


@pytest.mark.parametrize(
    ("side", "activity_account", "movement_rows", "closing_row", "expected"),
    [
        (
            "CREDIT",
            "5111",
            [
                {"debit_account": "131", "credit_account": "5111", "amount": 120},
                {"debit_account": "5111", "credit_account": "131", "amount": 20},
            ],
            {"debit_account": "5111", "credit_account": "911", "amount": 100},
            100,
        ),
        (
            "DEBIT",
            "6321",
            [
                {"debit_account": "6321", "credit_account": "1561", "amount": 110},
                {"debit_account": "1561", "credit_account": "6321", "amount": 10},
            ],
            {"debit_account": "911", "credit_account": "6321", "amount": 100},
            100,
        ),
    ],
)
def test_profit_loss_activity_prefers_period_closing_transfer(
    side: str,
    activity_account: str,
    movement_rows: list[dict],
    closing_row: dict,
    expected: float,
) -> None:
    records = [
        record("report", "REPORT", "PL", {"reportName": "Kết quả kinh doanh"}),
        record(
            "label",
            "REPORT_COLUMN",
            "PL_LABEL",
            {"reportKey": "PL", "columnKey": "LABEL", "columnType": "LABEL", "columnLabel": "Chỉ tiêu"},
            1,
        ),
        record(
            "current",
            "REPORT_COLUMN",
            "PL_CURRENT",
            {"reportKey": "PL", "columnKey": "CURRENT", "columnType": "CURRENT_PERIOD", "columnLabel": "Kỳ này"},
            2,
        ),
        record(
            "activity",
            "REPORT_ROW",
            "PL_ACTIVITY",
            {
                "reportKey": "PL",
                "rowCode": "ACTIVITY",
                "rowLabel": "Hoạt động trong kỳ",
                "rowType": "DATA",
                "cells": {
                    "CURRENT": {
                        "operator": "PROFIT_LOSS_ACTIVITY",
                        "sourceKey": "DETAIL",
                        "datasetType": "GENERAL_LEDGER",
                        "accountSetKey": "ACTIVITY_SET",
                        "side": side,
                    }
                },
            },
        ),
        record(
            "activity-member",
            "ACCOUNT_SET_MEMBER",
            "ACTIVITY_MEMBER",
            {"accountSetKey": "ACTIVITY_SET", "matchType": "PREFIX", "matchValue": activity_account, "role": "INCLUDE"},
        ),
    ]
    output = build_output(
        records=records,
        report_key="PL",
        period_from=date(2026, 1, 1),
        period_to=date(2026, 1, 31),
        rows_by_offset={
            0: [
                {
                    "source_key": "DETAIL__COMPANY",
                    "base_source_key": "DETAIL",
                    "dataset_type": "GENERAL_LEDGER",
                    "canonical_data": row,
                }
                for row in [*movement_rows, closing_row]
            ]
        },
        selection={"sourceKeys": ["DETAIL__COMPANY"], "companyKeys": [], "dataSpaceKeys": []},
        include_evidence=True,
    )
    cell = output["rows"][0]["cells"][1]
    assert cell["value"] == expected
    assert cell["evidence"]["isComplete"] is True
    assert sum(item["signedValue"] for item in cell["evidence"]["components"]) == expected
    assert cell["evidence"]["components"][0]["accountNumber"] == activity_account
    assert "ACTIVITY_SET" not in cell["evidence"]["components"][0]["label"]
    assert output["warnings"] == []


def test_profit_loss_activity_falls_back_to_net_movement_when_period_is_open() -> None:
    records = [
        record("report", "REPORT", "PL", {"reportName": "Kết quả kinh doanh"}),
        record(
            "label",
            "REPORT_COLUMN",
            "PL_LABEL",
            {"reportKey": "PL", "columnKey": "LABEL", "columnType": "LABEL", "columnLabel": "Chỉ tiêu"},
            1,
        ),
        record(
            "current",
            "REPORT_COLUMN",
            "PL_CURRENT",
            {"reportKey": "PL", "columnKey": "CURRENT", "columnType": "CURRENT_PERIOD", "columnLabel": "Kỳ này"},
            2,
        ),
        record(
            "expense",
            "REPORT_ROW",
            "PL_EXPENSE",
            {
                "reportKey": "PL",
                "rowCode": "EXPENSE",
                "rowLabel": "Chi phí trong kỳ",
                "rowType": "DATA",
                "cells": {
                    "CURRENT": {
                        "operator": "PROFIT_LOSS_ACTIVITY",
                        "sourceKey": "DETAIL",
                        "datasetType": "GENERAL_LEDGER",
                        "accountSetKey": "EXPENSE_SET",
                        "side": "DEBIT",
                    }
                },
            },
        ),
        record(
            "expense-member",
            "ACCOUNT_SET_MEMBER",
            "EXPENSE_642",
            {"accountSetKey": "EXPENSE_SET", "matchType": "PREFIX", "matchValue": "642", "role": "INCLUDE"},
        ),
    ]
    output = build_output(
        records=records,
        report_key="PL",
        period_from=date(2026, 1, 1),
        period_to=date(2026, 1, 31),
        rows_by_offset={
            0: [
                {
                    "source_key": "DETAIL__COMPANY",
                    "base_source_key": "DETAIL",
                    "dataset_type": "GENERAL_LEDGER",
                    "canonical_data": {"debit_account": "6421", "credit_account": "331", "amount": 150},
                },
                {
                    "source_key": "DETAIL__COMPANY",
                    "base_source_key": "DETAIL",
                    "dataset_type": "GENERAL_LEDGER",
                    "canonical_data": {"debit_account": "331", "credit_account": "6421", "amount": 25},
                },
            ]
        },
        selection={"sourceKeys": ["DETAIL__COMPANY"], "companyKeys": [], "dataSpaceKeys": []},
    )
    assert output["rows"][0]["cells"][1]["value"] == 125


def test_cross_report_rule_uses_the_selected_target_row() -> None:
    records = [
        record("report", "REPORT", "CF", {"reportName": "Lưu chuyển tiền tệ"}),
        record(
            "label",
            "REPORT_COLUMN",
            "CF_LABEL",
            {"reportKey": "CF", "columnKey": "LABEL", "columnType": "LABEL", "columnLabel": "Chỉ tiêu"},
            1,
        ),
        record(
            "current",
            "REPORT_COLUMN",
            "CF_CURRENT",
            {"reportKey": "CF", "columnKey": "CURRENT", "columnType": "CURRENT_PERIOD", "columnLabel": "Kỳ này"},
            2,
        ),
        record(
            "profit",
            "REPORT_ROW",
            "CF_01",
            {
                "reportKey": "CF",
                "rowCode": "01",
                "rowLabel": "Lợi nhuận trước thuế",
                "rowType": "CROSS_REPORT",
                "cells": {
                    "CURRENT": {
                        "operator": "CROSS_REPORT",
                        "targetReportKey": "PL",
                        "targetRowCode": "50",
                    }
                },
            },
        ),
    ]
    calls = []

    def resolve(report_key: str, row_code: str, column: dict, rule: dict) -> float:
        calls.append((report_key, row_code, column["columnKey"]))
        return 796_815_333

    output = build_output(
        records=records,
        report_key="CF",
        period_from=date(2026, 1, 1),
        period_to=date(2026, 1, 31),
        rows_by_offset={0: []},
        selection={"sourceKeys": ["TB"], "companyKeys": [], "dataSpaceKeys": []},
        cross_report_resolver=resolve,
    )
    assert calls == [("PL", "50", "CURRENT")]
    assert output["rows"][0]["cells"][1]["value"] == 796_815_333


def test_january_2026_golden_file_keeps_three_statement_links() -> None:
    """Golden values come from the accountant-prepared 4 BCTC_Global.xlsx.

    They are test fixtures only. Production formulas remain account/source based.
    """

    def statement_records(
        report_key: str,
        inputs: dict[str, float],
        formulas: dict[str, str],
        crosses: dict[str, tuple[str, str]] | None = None,
    ) -> tuple[list[dict], list[dict]]:
        output = [
            record("report", "REPORT", report_key, {"reportName": report_key}),
            record(
                "label",
                "REPORT_COLUMN",
                f"{report_key}_LABEL",
                {"reportKey": report_key, "columnKey": "LABEL", "columnType": "LABEL", "columnLabel": "Chỉ tiêu"},
                1,
            ),
            record(
                "current",
                "REPORT_COLUMN",
                f"{report_key}_CURRENT",
                {"reportKey": report_key, "columnKey": "CURRENT", "columnType": "CURRENT_PERIOD", "columnLabel": "Kỳ này"},
                2,
            ),
        ]
        order = 10
        for code, amount in inputs.items():
            output.append(
                record(
                    f"{report_key}_{code}",
                    "REPORT_ROW",
                    f"{report_key}_{code}",
                    {
                        "reportKey": report_key,
                        "rowCode": code,
                        "rowLabel": code,
                        "rowType": "DATA",
                        "cells": {"CURRENT": approved_adjustment_cell(report_key, code)},
                    },
                    order,
                )
            )
            order += 10
        for code, expression in formulas.items():
            output.append(
                record(
                    f"{report_key}_{code}",
                    "REPORT_ROW",
                    f"{report_key}_{code}",
                    {
                        "reportKey": report_key,
                        "rowCode": code,
                        "rowLabel": code,
                        "rowType": "FORMULA",
                        "cells": {"CURRENT": {"operator": "LINE_EXPRESSION", "expression": expression}},
                    },
                    order,
                )
            )
            order += 10
        for code, (target_report, target_row) in (crosses or {}).items():
            output.append(
                record(
                    f"{report_key}_{code}",
                    "REPORT_ROW",
                    f"{report_key}_{code}",
                    {
                        "reportKey": report_key,
                        "rowCode": code,
                        "rowLabel": code,
                        "rowType": "CROSS_REPORT",
                        "cells": {
                            "CURRENT": {
                                "operator": "CROSS_REPORT",
                                "targetReportKey": target_report,
                                "targetRowCode": target_row,
                            }
                        },
                    },
                    order,
                )
            )
        return output, [
            approved_adjustment_row(report_key, code, amount)
            for code, amount in inputs.items()
        ]

    income_records, income_rows = statement_records(
        "PL",
        {
            "01": 9_700_407_176,
            "02": 1_777_942_918,
            "11": 3_785_958_100,
            "22": 862_469,
            "23": 14_352_071,
            "25": 2_831_122_786,
            "26": 499_033_246,
            "31": 4_215_562,
            "32": 260_753,
            "51": 0,
            "52": 0,
        },
        {
            "10": "01-02",
            "20": "10-11",
            "30": "20+22-23-25-26",
            "40": "31-32",
            "50": "30+40",
            "60": "50-51-52",
        },
    )
    income = build_output(
        records=income_records,
        report_key="PL",
        period_from=date(2026, 1, 1),
        period_to=date(2026, 1, 31),
        rows_by_offset={0: income_rows},
        selection={"sourceKeys": ["GOLDEN"], "companyKeys": [], "dataSpaceKeys": []},
    )
    income_values = {
        row["rowCode"]: row["cells"][1]["value"] for row in income["rows"]
    }
    assert income_values["50"] == 796_815_333
    assert income_values["60"] == 796_815_333

    cash_records, cash_rows = statement_records(
        "CF",
        {
            "02": 27_205_389,
            "05": -742_683,
            "06": 14_352_071,
            "09": -220_194_863,
            "10": -93_133_522,
            "11": 61_300_732,
            "12": -8_200_703,
            "14": -2_007_252,
            "27": 742_683,
            "34": -10_222_222,
            "60": 2_149_578_503,
            "61": 0,
        },
        {
            "08": "01+02+05+06",
            "20": "08+09+10+11+12+14",
            "30": "27",
            "40": "34",
            "50": "20+30+40",
            "70": "50+60+61",
        },
        {"01": ("PL", "50")},
    )

    def resolve_income(report_key: str, row_code: str, column: dict, rule: dict) -> float | None:
        return income_values.get(row_code) if report_key == "PL" else None

    cash_flow = build_output(
        records=cash_records,
        report_key="CF",
        period_from=date(2026, 1, 1),
        period_to=date(2026, 1, 31),
        rows_by_offset={0: cash_rows},
        selection={"sourceKeys": ["GOLDEN"], "companyKeys": [], "dataSpaceKeys": []},
        cross_report_resolver=resolve_income,
    )
    cash_values = {
        row["rowCode"]: row["cells"][1]["value"] for row in cash_flow["rows"]
    }
    assert cash_values["20"] == 575_394_502
    assert cash_values["50"] == 565_914_963
    assert cash_values["70"] == 2_715_493_466


def test_balance_sheet_exposes_a_failed_accounting_control() -> None:
    records = [
        record("report", "REPORT", "BAO_CAO_TINH_HINH_TAI_CHINH", {"reportName": "BS"}),
        record(
            "label",
            "REPORT_COLUMN",
            "BS_LABEL",
            {"reportKey": "BAO_CAO_TINH_HINH_TAI_CHINH", "columnKey": "LABEL", "columnType": "LABEL"},
            1,
        ),
        record(
            "current",
            "REPORT_COLUMN",
            "BS_CURRENT",
            {"reportKey": "BAO_CAO_TINH_HINH_TAI_CHINH", "columnKey": "CURRENT", "columnType": "CURRENT_PERIOD"},
            2,
        ),
        record(
            "assets",
            "REPORT_ROW",
            "BS_ASSETS",
            {
                "reportKey": "BAO_CAO_TINH_HINH_TAI_CHINH",
                "rowCode": "270_TONG_TS",
                "rowType": "TOTAL",
                    "cells": {"CURRENT": approved_adjustment_cell("BAO_CAO_TINH_HINH_TAI_CHINH", "270_TONG_TS")},
            },
            10,
        ),
        record(
            "sources",
            "REPORT_ROW",
            "BS_SOURCES",
            {
                "reportKey": "BAO_CAO_TINH_HINH_TAI_CHINH",
                "rowCode": "440",
                "rowType": "TOTAL",
                    "cells": {"CURRENT": approved_adjustment_cell("BAO_CAO_TINH_HINH_TAI_CHINH", "440")},
            },
            20,
        ),
    ]
    output = build_output(
        records=records,
        report_key="BAO_CAO_TINH_HINH_TAI_CHINH",
        period_from=date(2026, 1, 1),
        period_to=date(2026, 1, 31),
        rows_by_offset={0: [
            approved_adjustment_row("BAO_CAO_TINH_HINH_TAI_CHINH", "270_TONG_TS", 120),
            approved_adjustment_row("BAO_CAO_TINH_HINH_TAI_CHINH", "440", 100),
        ]},
        selection={"sourceKeys": ["TB"], "companyKeys": [], "dataSpaceKeys": []},
    )
    assert output["controlStatus"] == "NEEDS_REVIEW"
    assert output["controlChecks"][0]["checkKey"] == "BALANCE_SHEET_BALANCES"
    assert output["controlChecks"][0]["difference"] == 20
    assert any("Kiểm tra kế toán chưa đạt" in warning for warning in output["warnings"])


def test_cash_flow_exposes_a_failed_cash_reconciliation_control() -> None:
    report_key = "LUU_CHUYEN_TIEN_TE_GIAN_TIEP"
    records = [
        record("report", "REPORT", report_key, {"reportName": "CF"}),
        record(
            "label",
            "REPORT_COLUMN",
            "CF_LABEL",
            {"reportKey": report_key, "columnKey": "LABEL", "columnType": "LABEL"},
            1,
        ),
        record(
            "current",
            "REPORT_COLUMN",
            "CF_CURRENT",
            {"reportKey": report_key, "columnKey": "CURRENT", "columnType": "CURRENT_PERIOD"},
            2,
        ),
    ]
    for sort_order, row_code, value in (
        (10, "50", 100),
        (30, "61", 0),
        (40, "70", 119),
    ):
        records.append(
            record(
                f"row-{row_code}",
                "REPORT_ROW",
                f"CF_{row_code}",
                {
                    "reportKey": report_key,
                    "rowCode": row_code,
                    "rowType": "TOTAL",
                    "cells": {"CURRENT": approved_adjustment_cell(report_key, row_code)},
                },
                sort_order,
            )
        )
    records.extend(
        [
            record(
                "cash-opening",
                "REPORT_ROW",
                "CF_60",
                {
                    "reportKey": report_key,
                    "rowCode": "60",
                    "rowType": "DATA",
                    "ruleKey": "CF_CASH_RULE",
                },
                20,
            ),
            record(
                "cash-rule",
                "RULE",
                "CF_CASH_RULE",
                {
                    "reportKey": report_key,
                    "terms": [
                        {
                            "operator": "BALANCE",
                            "accountSetKey": "CASH",
                            "periodBasis": "OPENING",
                            "side": "DEBIT_BALANCE",
                        }
                    ],
                },
                50,
            ),
            record(
                "cash-member",
                "ACCOUNT_SET_MEMBER",
                "CASH_111",
                {"accountSetKey": "CASH", "matchType": "PREFIX", "matchValue": "111", "role": "INCLUDE"},
                60,
            ),
        ]
    )
    output = build_output(
        records=records,
        report_key=report_key,
        period_from=date(2026, 1, 1),
        period_to=date(2026, 1, 31),
        rows_by_offset={
            0: [
                approved_adjustment_row(report_key, "50", 100),
                approved_adjustment_row(report_key, "61", 0),
                approved_adjustment_row(report_key, "70", 119),
                {
                    "account_number": "111",
                    "is_leaf": True,
                    "opening_debit": 20,
                    "closing_debit": 130,
                }
            ]
        },
        selection={"sourceKeys": ["JOURNAL"], "companyKeys": [], "dataSpaceKeys": []},
    )
    check = next(
        item
        for item in output["controlChecks"]
        if item["checkKey"] == "CASH_FLOW_RECONCILES_TO_CASH_BALANCE"
    )
    assert output["controlStatus"] == "NEEDS_REVIEW"
    assert check["status"] == "FAIL"
    assert check["difference"] == -10


def test_profit_control_compares_report_with_account_421_movement() -> None:
    report_key = "BAO_CAO_KET_QUA_KINH_DOANH"
    bs_key = "BAO_CAO_TINH_HINH_TAI_CHINH"
    records = [
        record("report", "REPORT", report_key, {"reportName": "PL"}),
        record("bs", "REPORT", bs_key, {"reportName": "BS"}),
        record(
            "label",
            "REPORT_COLUMN",
            "PL_LABEL",
            {"reportKey": report_key, "columnKey": "LABEL", "columnType": "LABEL"},
            1,
        ),
        record(
            "current",
            "REPORT_COLUMN",
            "PL_CURRENT",
            {"reportKey": report_key, "columnKey": "CURRENT", "columnType": "CURRENT_PERIOD"},
            2,
        ),
        record(
            "profit",
            "REPORT_ROW",
            "PL_60",
            {
                "reportKey": report_key,
                "rowCode": "60",
                "rowType": "TOTAL",
                "cells": {"CURRENT": approved_adjustment_cell(report_key, "60")},
            },
            10,
        ),
        record(
            "retained-row",
            "REPORT_ROW",
            "BS_420A",
            {"reportKey": bs_key, "rowCode": "420a", "rowType": "DATA", "ruleKey": "BS_420A_RULE"},
            20,
        ),
        record(
            "retained-rule",
            "RULE",
            "BS_420A_RULE",
            {"reportKey": bs_key, "accountSetKey": "RETAINED", "operator": "BALANCE"},
            30,
        ),
        record(
            "retained-member",
            "ACCOUNT_SET_MEMBER",
            "RETAINED_421",
            {"accountSetKey": "RETAINED", "matchType": "PREFIX", "matchValue": "421", "role": "INCLUDE"},
            40,
        ),
    ]
    output = build_output(
        records=records,
        report_key=report_key,
        period_from=date(2026, 1, 1),
        period_to=date(2026, 1, 31),
        rows_by_offset={
            0: [
                approved_adjustment_row(report_key, "60", 100),
                {
                    "account_number": "4212",
                    "is_leaf": True,
                    "period_debit": 0,
                    "period_credit": 90,
                }
            ]
        },
        selection={"sourceKeys": ["TB"], "companyKeys": [], "dataSpaceKeys": []},
    )
    check = next(item for item in output["controlChecks"] if item["checkKey"] == "PROFIT_TO_RETAINED_EARNINGS")
    assert check["status"] == "FAIL"
    assert check["difference"] == 10


def test_data_check_xlsx_is_a_valid_workbook() -> None:
    output = build_output(
        records=[],
        report_key="DATA_CHECK",
        period_from=date(2026, 1, 1),
        period_to=date(2026, 6, 30),
        rows_by_offset={
            0: [
                {
                    "account_number": "111",
                    "account_name": "Tiền mặt",
                    "grade": 1,
                    "opening_debit": 10,
                    "period_debit": 5,
                    "closing_debit": 15,
                }
            ]
        },
        selection={"sourceKeys": ["SRC"], "companyKeys": [], "dataSpaceKeys": []},
    )
    payload = output_to_xlsx(output)
    workbook = load_workbook(BytesIO(payload), data_only=True)
    sheet = workbook["Báo cáo"]
    assert sheet["A1"].value == "BẢNG CÂN ĐỐI TÀI KHOẢN"
    assert sheet["A6"].value == "111"
    assert sheet["J6"].value == 0


def test_report_xlsx_skips_rows_marked_not_to_print() -> None:
    output = {
        "report": {"displayTitle": "BÁO CÁO KIỂM TRA", "hideColumnHeader": True},
        "periodFrom": "2026-01-01",
        "periodTo": "2026-01-31",
        "columns": [{"columnKey": "LABEL", "width": 28}],
        "rows": [
            {
                "rowCode": "HIDDEN",
                "rowType": "DATA",
                "hideOnPrint": True,
                "cells": [{"text": "Dòng không in"}],
            },
            {
                "rowCode": "VISIBLE",
                "rowType": "DATA",
                "cells": [{"text": "Dòng được in"}],
            },
        ],
    }
    workbook = load_workbook(BytesIO(output_to_xlsx(output)), data_only=True)
    values = [
        value
        for row in workbook["Báo cáo"].iter_rows(values_only=True)
        for value in row
    ]
    assert "Dòng không in" not in values
    assert "Dòng được in" in values


def test_required_and_optional_notes_are_rendered_with_supported_balances() -> None:
    records = [
        record(
            "report-notes",
            "REPORT",
            "B02",
            {"reportName": "Bảng cân đối", "displayTitle": "BẢNG CÂN ĐỐI"},
        ),
        record(
            "column-label",
            "REPORT_COLUMN",
            "B02_LABEL",
            {"reportKey": "B02", "columnKey": "LABEL", "columnType": "LABEL", "columnLabel": "Chỉ tiêu"},
            1,
        ),
        record(
            "column-note",
            "REPORT_COLUMN",
            "B02_NOTE",
            {"reportKey": "B02", "columnKey": "NOTE", "columnType": "NOTE_REFERENCE", "columnLabel": "Thuyết minh"},
            2,
        ),
        record(
            "row-note",
            "REPORT_ROW",
            "B02_ROW_110",
            {
                "reportKey": "B02",
                "rowCode": "110",
                "rowLabel": "Tiền",
                "rowType": "LABEL",
                "noteItemId": "note-1",
                "noteRequired": True,
                "cells": {},
            },
            1,
        ),
    ]
    note_items = [
        {
            "itemId": "note-1",
            "itemCode": "V.1",
            "itemName": "Tiền và tương đương tiền",
            "itemIndex": 1,
            "isBold": True,
            "formula": [
                {"operation_sign": 1, "operand": "DUNODK", "account_number": "111"},
                {"operation_sign": 1, "operand": "DUNOCK", "account_number": "111"},
            ],
        },
        {
            "itemId": "note-2",
            "itemCode": "V.2",
            "itemName": "Doanh thu",
            "itemIndex": 2,
            "formula": [
                {"operation_sign": 1, "operand": "PhatsinhCO", "account_number": "511"},
            ],
        },
    ]
    rows = [
        {
            "account_number": "1111",
            "is_leaf": True,
            "opening_debit": 100,
            "closing_debit": 130,
        },
        {
            "account_number": "5111",
            "is_leaf": True,
            "period_credit": 450,
        },
    ]
    assert note_required_offsets(records, "B02", note_items, ["note-2"]) == [0]
    output = build_output(
        records=records,
        report_key="B02",
        period_from=date(2026, 1, 1),
        period_to=date(2026, 12, 31),
        rows_by_offset={0: rows},
        selection={"sourceKeys": ["SRC"], "companyKeys": ["NTV"], "dataSpaceKeys": ["SPACE"]},
        note_items=note_items,
        selected_note_item_ids=["note-2"],
    )
    assert output["rows"][0]["cells"][1]["text"] == "V.1"
    assert output["notes"]["requiredItemIds"] == ["note-1"]
    assert output["notes"]["optionalItemIds"] == ["note-2"]
    assert output["notes"]["rows"][0]["openingValue"] == 100
    assert output["notes"]["rows"][0]["closingValue"] == 130
    assert output["notes"]["rows"][1]["movementValue"] == 450

    workbook = load_workbook(BytesIO(output_to_xlsx(output)), data_only=True)
    assert "Thuyết minh" in workbook.sheetnames
    note_sheet = workbook["Thuyết minh"]
    assert note_sheet["A5"].value == "V.1"
    assert note_sheet["C5"].value == 100
    assert note_sheet["E5"].value == 130


def test_note_formula_can_request_previous_year_snapshot() -> None:
    records = [
        record("report", "REPORT", "B03", {"reportName": "B03"}),
        record(
            "row",
            "REPORT_ROW",
            "B03_ROW",
            {
                "reportKey": "B03",
                "rowCode": "1",
                "rowLabel": "Dòng",
                "rowType": "LABEL",
                "noteItemId": "note-old",
                "noteRequired": True,
            },
        ),
    ]
    items = [
        {
            "itemId": "note-old",
            "formula": [
                {
                    "operation_sign": 1,
                    "operand": "DUNODK_NamTruoc",
                    "account_number": "111",
                }
            ],
        }
    ]
    assert note_required_offsets(records, "B03", items, []) == [-1, 0]


def test_balance_sheet_blocks_a_nonzero_account_assigned_to_two_debit_lines() -> None:
    report_key = "BAO_CAO_TINH_HINH_TAI_CHINH"
    records = [
        record("report", "REPORT", report_key, {"reportName": "BS"}),
        record(
            "label",
            "REPORT_COLUMN",
            "BS_LABEL",
            {"reportKey": report_key, "columnKey": "LABEL", "columnType": "LABEL"},
            1,
        ),
        record(
            "current",
            "REPORT_COLUMN",
            "BS_CURRENT",
            {
                "reportKey": report_key,
                "columnKey": "CURRENT",
                "columnType": "CURRENT_PERIOD",
                "columnLabel": "Cuối kỳ",
                "valueBasis": "ENDING",
            },
            2,
        ),
        record(
            "short-row",
            "REPORT_ROW",
            "BS_131",
            {
                "reportKey": report_key,
                "rowCode": "131",
                "rowType": "DATA",
                "ruleKey": "SHORT_RULE",
            },
            10,
        ),
        record(
            "long-row",
            "REPORT_ROW",
            "BS_211",
            {
                "reportKey": report_key,
                "rowCode": "211",
                "rowType": "DATA",
                "ruleKey": "LONG_RULE",
            },
            20,
        ),
        record(
            "short-rule",
            "RULE",
            "SHORT_RULE",
            {
                "reportKey": report_key,
                "operator": "BALANCE",
                "accountSetKey": "SHORT_RECEIVABLE",
                "side": "DEBIT",
                "periodBasis": "ENDING",
            },
        ),
        record(
            "long-rule",
            "RULE",
            "LONG_RULE",
            {
                "reportKey": report_key,
                "operator": "BALANCE",
                "accountSetKey": "LONG_RECEIVABLE",
                "side": "DEBIT",
                "periodBasis": "ENDING",
            },
        ),
        record(
            "short-member",
            "ACCOUNT_SET_MEMBER",
            "SHORT_RECEIVABLE_131",
            {
                "accountSetKey": "SHORT_RECEIVABLE",
                "companyKey": "TVG",
                "matchType": "PREFIX",
                "matchValue": "131",
                "role": "INCLUDE",
            },
        ),
        record(
            "long-member",
            "ACCOUNT_SET_MEMBER",
            "LONG_RECEIVABLE_131",
            {
                "accountSetKey": "LONG_RECEIVABLE",
                "companyKey": "TVG",
                "matchType": "PREFIX",
                "matchValue": "131",
                "role": "INCLUDE",
            },
        ),
    ]
    output = build_output(
        records=records,
        report_key=report_key,
        period_from=date(2026, 1, 1),
        period_to=date(2026, 1, 31),
        rows_by_offset={
            0: [
                {
                    "company_key": "TVG",
                    "account_number": "1311",
                    "account_name": "Phải thu khách hàng ngắn hạn",
                    "is_leaf": True,
                    "closing_debit": 125,
                    "closing_credit": 0,
                }
            ]
        },
        selection={"sourceKeys": ["TB"], "companyKeys": ["TVG"], "dataSpaceKeys": ["SPACE"]},
    )
    assert any(
        "TK 1311" in issue and "131, 211" in issue
        for issue in output["blockingIssues"]
    )


def test_opening_retained_loss_keeps_its_negative_sign() -> None:
    report_key = "BAO_CAO_TINH_HINH_TAI_CHINH"
    records = [
        record("report", "REPORT", report_key, {"reportName": "BS"}),
        record(
            "label",
            "REPORT_COLUMN",
            "BS_LABEL",
            {"reportKey": report_key, "columnKey": "LABEL", "columnType": "LABEL"},
            1,
        ),
        record(
            "current",
            "REPORT_COLUMN",
            "BS_CURRENT",
            {"reportKey": report_key, "columnKey": "CURRENT", "columnType": "CURRENT_PERIOD"},
            2,
        ),
        record(
            "retained-row",
            "REPORT_ROW",
            "BS_420A",
            {
                "reportKey": report_key,
                "rowCode": "420A",
                "rowType": "DATA",
                "ruleKey": "RETAINED_RULE",
            },
            10,
        ),
        record(
            "retained-rule",
            "RULE",
            "RETAINED_RULE",
            {
                "reportKey": report_key,
                "operator": "BALANCE",
                "accountSetKey": "RETAINED",
                "side": "CREDIT_BALANCE",
                "periodBasis": "OPENING",
                "useColumnBasis": False,
            },
        ),
        record(
            "retained-member",
            "ACCOUNT_SET_MEMBER",
            "RETAINED_421",
            {
                "accountSetKey": "RETAINED",
                "companyKey": "TVG",
                "matchType": "PREFIX",
                "matchValue": "421",
                "role": "INCLUDE",
            },
        ),
    ]
    output = build_output(
        records=records,
        report_key=report_key,
        period_from=date(2026, 1, 1),
        period_to=date(2026, 1, 31),
        rows_by_offset={
            0: [
                {
                    "company_key": "TVG",
                    "account_number": "4212",
                    "is_leaf": True,
                    "opening_debit": 1_302_574_504,
                    "opening_credit": 0,
                }
            ]
        },
        selection={"sourceKeys": ["TB"], "companyKeys": ["TVG"], "dataSpaceKeys": ["SPACE"]},
    )
    assert output["rows"][0]["cells"][1]["value"] == -1_302_574_504


def test_selected_export_adjustment_propagates_into_parent_formula() -> None:
    report_key = "BAO_CAO_KET_QUA_KINH_DOANH"
    records = [
        record("report", "REPORT", report_key, {"reportName": "KQKD"}),
        record("label", "REPORT_COLUMN", "LABEL", {"reportKey": report_key, "columnKey": "LABEL", "columnType": "LABEL"}, 1),
        record("current", "REPORT_COLUMN", "CURRENT", {"reportKey": report_key, "columnKey": "CURRENT", "columnType": "CURRENT_PERIOD"}, 2),
        record("child", "REPORT_ROW", "CHILD", {"reportKey": report_key, "rowCode": "01", "rowLabel": "Doanh thu", "rowType": "DATA", "ruleKey": "ZERO_RULE"}, 10),
        record("parent", "REPORT_ROW", "PARENT", {"reportKey": report_key, "rowCode": "10", "rowLabel": "Doanh thu thuần", "rowType": "FORMULA", "ruleKey": "SUM_RULE"}, 20),
        record("sum-rule", "RULE", "SUM_RULE", {"reportKey": report_key, "operator": "SUM_LINES", "lineRefs": ["01"]}),
        record("zero-rule", "RULE", "ZERO_RULE", {"reportKey": report_key, "operator": "LINE_EXPRESSION", "expression": "0"}),
    ]
    output = build_output(
        records=records,
        report_key=report_key,
        period_from=date(2026, 1, 1),
        period_to=date(2026, 1, 31),
        rows_by_offset={0: []},
        selection={"sourceKeys": [], "companyKeys": ["TVG"], "dataSpaceKeys": ["SPACE"]},
        include_evidence=True,
        cell_adjustments=[{
            "rowCode": "01",
            "columnKey": "CURRENT",
            "amount": 125,
            "reason": "Điều chỉnh sau khóa sổ",
            "sourceReference": "PKT-01",
        }],
    )
    assert output["rows"][0]["cells"][1]["value"] == 125
    assert output["rows"][1]["cells"][1]["value"] == 125
    components = output["rows"][0]["cells"][1]["evidence"]["components"]
    assert any(component["sourceName"] == "Bút toán điều chỉnh được chọn" for component in components)
